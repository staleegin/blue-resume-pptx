#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
单页 A4 简历 PPT 生成器（Python 版，单文件）  https://github.com/staleegin/blue-resume-pptx

用法:  python resume_pptx.py data.json out.pptx
依赖:  python-pptx、Pillow（ChatGPT 等网页 AI 的代码环境一般自带，不需要联网）
输出:  一页 A4 纵向、可编辑的 .pptx；终端打印一行 JSON：
       level 压缩档位(0=原始, 4=8 号字；单栏版式内容少时为负数，表示放大)；overflowPt>0 表示仍超出一页；
       movedToRight 为自动挪到右栏的板块（单栏版式里是被合并成「其他信息」的卡片）

==================== 给 AI 的使用说明 ====================
1. 从用户材料（旧简历、零散文字、JD）里提取信息，整理成下面的 data.json。
2. 缺姓名、求职意向、联系方式时先问用户；其余缺项一次性问完，用户说「没有」就跳过。
3. 不要编造经历、数字、证书。涉及事实的改写必须先问用户。
4. 重点短语用 **...** 标出（每条最多一处）；带单位的数字（90 人、56%、$2.5M）脚本会自动加粗。
5. 内容超过一页时（约 30 条要点）先给精简方案，用户确认后再删。
6. 运行脚本，把 .pptx 交给用户下载；若 overflowPt>0，告诉用户超出约几行并建议删哪几条。
7. 用户上传照片时，把照片放在 data.json 同目录，并在 photo 字段写文件名。

==================== data.json 格式 ====================
{
  "layout": "simple",                   // 可省略。省略=蓝色顶栏双栏；"simple"=简洁单栏（无色块、不分栏，右栏卡片排到经历下面，
                                        //   「个人信息」卡片并入姓名下方一行，照片在右上角）
  "info": ["1999 年生", "本科"],         // 可省略，仅单栏版式用：姓名下方的个人信息行；省略时取「个人信息」卡片
  "theme": "2563EB",                    // 主色，可省略（默认蓝色）；墨绿 0F766E、深蓝 1E3A8A、酒红 9F1239、深灰 334155
  "lang": "zh",                         // 可省略，自动识别中英文
  "name": "姓名",
  "title": "前端开发工程师",             // 显示为「求职意向：…」；英文简历不加前缀
  "titleLabel": "求职意向：",            // 可省略
  "contacts": [                         // type: phone/email/location/wechat/web/github/linkedin/link/portfolio/birthday/user
    {"type": "phone", "text": "138xxxx"},
    {"type": "email", "text": "a@b.com"},                         // 邮箱、网址会自动变成可点击链接
    {"type": "portfolio", "text": "作品集", "link": "https://..."} // link 只接受 http(s)/mailto/tel
  ],
  "summary": "一两句个人简介",            // 可省略
  "photo": "photo.jpg",                 // 可省略，没有照片时右上角留空
  "left": [                             // 左栏板块，数量和名称都可自定义
    {"title": "实习经历", "entries": [
      {"name": "公司 · 职位", "date": "2025.07 - 2025.12",
       "sub": "副标题（如专业、部门）", "subRight": "右侧信息（如学历）",
       "details": ["技术栈：React、TypeScript"],   // 明细行，冒号前显示为黑色
       "bullets": ["要点，**重点**会加粗"],         // 带项目符号的要点
       "text": "不带项目符号的段落"}                // 以上字段都可省略
    ]}
  ],
  "right": [                            // 右栏卡片
    {"title": "技能", "tags": ["React", "TypeScript"]},                      // 标签式
    {"title": "证书", "items": [{"name": "CET-6", "date": "2024.06", "content": "562 分"}]}  // 小项式，content 可用 \n 换行
  ]
}
"""
import base64, copy, io, json, math, os, re, sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu
from lxml import etree

# ---------- 常量（单位：pt） ----------
W, H = 595.25, 841.85
BOTTOM = H - 28
CJK, LAT, YH = 'SimHei', 'Arial', 'Microsoft YaHei'
GREY, BLACK, WHITE = '595959', '000000', 'FFFFFF'
L_X, DATE_R, BUL_X, BUL_IND, BUL_W, PARA_X, PARA_W = 25.6, 380.4, 30.1, 13.5, 359, 24.6, 356
R_X, R_W, R_IN, R_TXT_W = 403, 174, 411.9, 140
HEAD_X = 32.5   # 顶栏里姓名、求职意向、分隔线、简介的左边距
# 自动微调档位：正文字号 / 标题字号 / 行距 / 间距系数
LEVELS = [
    dict(fs=9, ts=9.9, lh=15.1, g=1),
    dict(fs=9, ts=9.9, lh=14.2, g=0.85),
    dict(fs=9, ts=9.9, lh=13.4, g=0.7),
    dict(fs=8.5, ts=9.4, lh=12.6, g=0.7),
    dict(fs=8, ts=9, lh=11.8, g=0.6),
]


def E(v):
    """pt → EMU（与 pptxgenjs 的取整方式一致）"""
    return Emu(int(math.floor(v * 12700 + 0.5)))


def jsround(v):
    """四舍五入（与 JS 的 Math.round 一致；Python 的 round 是银行家舍入）"""
    return int(math.floor(v + 0.5))


def tint(hexc, a):
    return ''.join('%02X' % jsround(255 - (255 - int(hexc[i:i + 2], 16)) * a) for i in (0, 2, 4))


def strip(t):
    return str(t or '').replace('**', '')


# ---------- 文字宽度估算 ----------
# 黑体：汉字 1em、半角 0.5em；Arial 用实际字宽表（千分之一 em，ASCII 32-126）
ARIAL = [278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278, 556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 278, 278, 584, 584, 584, 556, 1015,
         667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833, 722, 778, 667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556, 333,
         556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556, 556, 556, 333, 500, 278, 556, 500, 722, 500, 500, 500, 334, 260, 334, 584]

# 数字 + 单位（中文量词、%、万/K/M/million、percent、x 倍）；货币符号开头的金额
NUM_RE = re.compile(r'(?:超|近|逾|约)?[$¥￥€£]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\s*(?:%|％|\+)?\s*(?:万|千|百|亿|million|billion|[KMBkw](?![A-Za-z]))?(?:余)?\s*(?:人次|人|项|个月|个|次|篇|场|倍|名|元|单|秒|分钟|小时|天|件|款|家|所|条|份|位|台|套|门|周|月|年|percent|x(?![A-Za-z]))?(?:%|％|\+)?|(?:超|近|逾)?[百千万]余?人次')
SKIP_RE = (re.compile(r'^[\d.,]+$'), re.compile(r'^\d{4}\s*年$'), re.compile(r'^\d{1,2}\s*月$'))

ICONS = {
    'phone': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAAApVBMVEVMaXH////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////Mid1sAAAANnRSTlMA+O3WmTDzLwIMNu7jJC7YOD1C15MxwUUm9JL5wAFEJQuRm0PO8pD3vs0ON73rIZzkzzm/PpqWWsoUAAAACXBIWXMAAAsTAAALEwEAmpwYAAABOElEQVRo3u3aSXOCQBQE4EF2RdS4xX1fkrhknf//03JIApdM2fPIHFLVfW7ehyhgwSj1e4br1kjDGXU/5som9wNtnekCn//e0II0OvD+i+ZrXWuCwFELM8DmD7U4Gwi4yoEAAlpyoAsBkRyIIKCoP+QzpD+Lt8UmdkCM/q5DIZCgQCIE8FOfAAECBAgQIECAAAECBAgQIECAAAECBAj8CeD8yS/87LovBLYh9BmS/osQEIQAgVvAYV9fLp9Wnisgy78KceYGOOQ/jdhzAuzLytgJUC8rbSdAWlZS18Dj/zxEq7LSqwTsDBt7xQ0onBgqO1XphXX2LYTPpoYPAWvjeeqN22n62ptUfOU+l1/r7rB7+Ek6fwr+SbicHS/cUJ2aaD689ESpxZv9/GNT2WQT+BbLHyI/MH2/n3mtaBdgTd8QAAAAAElFTkSuQmCC',
    'email': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAAB3VBMVEVMaXH////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////49NrfAAAAnnRSTlMACfr5/iSAAe5vRPOR/BD9Be1CBLOZwYixJTTeI9cT+x3CA9YmEcqr+Hc+Q6rYFDjdvj1m7EzxcBjhf0Eix+sIrfANVTbboiECRQfvDPQwFrvfwOTPU5s5GuPOpOdHzdIOUfYyuFa9PKPmhCcs0fUSmLzgHOXqyytPVJc1nDobXoIgpumHsE7UtZbZv10fhsmLr03y03EPtFIzmlziwwYVdiYAAAAJcEhZcwAACxMAAAsTAQCanBgAAALxSURBVGje7ZnnWxpBEMaXfjSR3hQUItKMiC0ae01sicYSTe+999577/3+1mBk96jHnew8j0luPvFwc785luV9ZxaEpJBCCinWSmyeb5hkqcZk9Xcdwa9zsCAx417hRxkWKJjkn+cH47OsWZUq0MQChiP1/bKgEUM3YQvYUQNsgWqkhy2gRyxw/AMFjLB8I5qGLZBAD2ELvEYTs5D8pQmEKlxwfFfFspxqtHKgL1i9Ke040WMQ/EucpaEt5+jza7dnmrJVbaKLVxoqc3z/PVVja7uW31m8CtHjz6o4rmefJv1K9owW/7osjdQcGkIM+6gOF/tKZZmYOczzPWWZZTX9uYDfqaktn/+mC9NueLFc+yMaspuMtHaPRqvI8AMXWaaN28rhP9iFOXvrsw2noxFfWW9bPX/rBkz5lsh1NHkEZX828eEn2oOaFQUss76CbN3h1fCHPfj+ur7CntxxBGcMHBXP555vKFHM9P3qAE66I7Jf2q/lVljO01U0EQk8sSSG/+QkEeaD/G1LTyvOHDMoBfNHsDYg3c5SfZHSMIaTH7cIw/cP4jsq81W/QOO1243zLwiarbpHcf69bmGd3YeX+I5ApKRdG53EWQbHBbeOcbLlGr38/MUozpQZxPSm00Q5Psd5+xIfzjvwXFzzq/iEBVbTXnSZFO0kSRsW3V1zwuUpskyLtwQIJE/7ftFC7v9S6LrNJ0TikTDzaD6dezGsxcsT4DUp/gEkVIUf8sqenOFuSqDNlphwxq+SXXg8wyXCd4k2tJrLG6GUb63kzOR+WgdMO8hv3epUlj2j/aohPZSq8zLDhDq5ruq2g8YQ2P+j2DHTXAulKTN+phB+wE5vjD2vy+dbvFTn5HeqbHxXH+1B3GRb4PBTI3KISf+jU1fV21sVc576v84qpAJSgb+gAPCxZhv0wWwQzcMWOIt0sAVSvdVhSP5MSttHe1jYPyhQ0gzGT674k/sFDL+Js9mYPUj596AP2i3S34dSSCHFmonfS8GeRXnvHuoAAAAASUVORK5CYII=',
    'location': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACN1BMVEVMaXH////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////ZmOKXAAAAvHRSTlMAAQMFFwf5/fz+9xuu9ghzBvAE+x1QAu8MFvjZFe3xiUKp7AlRdTfCEuU65CDuEFfdez2AHsPSinCzNc1MLddW3HpbqNGTdLzBRYisK5J4fT9EJatyd1g5fMSGJM4K0w+Vdjhdb7uwY6QjyGaPlDITuL3hYOaEIsdlRusnzPXW+jZaO+B+X8Zkp+rLSo3QsU/0MNURVJfawF6CxaaHSYzPTi/Ult6/46VD6Its8nG0M9gUuZq+XOJhurIcNIPCva0AAAAJcEhZcwAACxMAAAsTAQCanBgAAAPFSURBVGjetVllc9tAEJVjxxCyXYcccpiZuYE2TcNMDWOZmZmZmZkZ78d1pp1OM/bu6UB+n9++N5JOewuKwgaf8rmS2eLcLKMxK7d41jX3y0fRED9Ly2KIG+w/pv21UU+YmjEREObvX4Kk5cOn0wgFJwrGpOT1SdFEBdFWvbh+/ihhwPtOQfmgcQthgqXVV0Q/cyVhRsUyfv3OGMIBeyyvfqIf4YIxh09/lYlwwrSdRz/ETLhhWsiuHxtBBGDcwarfFUOEYF/Bpu9bQQSxm+1/GCfCaGXKDxZxA0uHuv6CRUQCo+qZL4lI4Y1q/o+XM4gPVzHoJZKwqtyPL/BQc2PPpFOvd05aUyg/+mb6LTqFBgbejvpPc7oCUWId1WAGC4tzKyGCuzHmepq+P/LslgKdO1XXiyRcE62aKeVJlDmIwyuKQRkc0gaz82B2AOUvtsPvH6HrmkF6Fl5VlsPnB32pwfC9+g41OAbyh/BHdoEB31B+Cfh/ReEGBvDUNaH8sxD9Gu1cF/F9ZTBT99AMGqCIeyj9EERPphmkQxGPUTqYqqNoBgYoYidKN0J0G7VAgCIiULpZGwMzSg/U5hXtQelhEL2fZnAQbHpQeg33HTgBRUSi9G0QPYVmcBOKGETp9eAnc1I+AXgl3EL5BWDucuEGd8GAAbytgdN1MMYvhNP1fdTgA1yWduuQCycAvmBT8Uc+A1+CtTC7HWYPUw7FYeTSBzu8PuTSz6MY7ML6r1rPsmUA6xOpPXMKVk41u33pwgCM+Zla2V1GK0K/EsO88/8Sb6ND6O1HLqVPLWpIN9hshvSkQUoXfVVlFLZJtnyvU+ugIuX096rO8kLkDCrVu8CPMvpLGNrYKoe4/r4Wlka5SdxgC9sooVpUP5Jxytkh2Oub8lnHLZ/EDB4xz4vGMkT007LZJ1KVIgbneGZq9fz6y7mGgtlfefUzUvnGjks5T5LlDu/g9AqfwWLuye8IV1qtCec2UPo5ZqeOZJHp+BC7wQ2x7cRJVv1hm9gCoYpxAOzXogjCymZQKrxiCd3AdIuFii9x/KPV9cMKZdZQB9QN+uT2aMfV9J9JLuqyN2qb4zzxnLpsMcUq0lhNMzitwbLUtxjXv5igxTr2ghHTj3igzcK3DTOY0GhfHfoE1i/SbCueCa6l7Jna7dyPQAbXtdzqx3FNkAVw6alHHbdOUwOPMsZyStEYa2TLFNUF/MP5+tUjmhsoa+f90I79ihfQzjbvkPihG//pH/XxioGy9fVf/fNdipeQ+KeadIQoXkPy27CwsnKukN8qqy88IW8DWgAAAABJRU5ErkJggg==',
    'wechat': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACl1BMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////921Z3zAAAA3HRSTlMA8hox7QT+/P0C+rIGKKffRM73DR0IjgpNvibPR8L7+eHxVY80fUETZusu9jKrARkJvQOttxeiB0+2giSAqRSYQvgrFVao26Oey3nXyK5ayUiwv+DZEQ7aIxznii115mWJUNZuDPSXFrs6IAsQQw9SM7lcw4WWLM0YX+9e2KQnmwVsr6XQ0+5d6oiE5CqQuEDEPRI1V37wb9zRtUzSmTdbIkZqnWmsjbGSMNVzeJzAyoxt8y9TuljefONiqotxtJWaOJ/iYUVKhnTdzHAfS4Nnx3fGSejlJaE7s2DUc/cwgAAAAAlwSFlzAAALEwAACxMBAJqcGAAABH5JREFUaN7tWWVbWzEUTimlt0ChUGC4DtmGDZiwYWPI2AZDxgZzd3dnBnN3d3d3d3fPj1kF1ptcS9LyZU/fT3Cek/dtcpOTc04AcMIJJ/57zEncEBq9x/DIT+Nwas+x97xOt4FN0EZ6dAo0Oox9RFl2FBQifPpAP0fQnxqlglLgzhfYu1pbMqE8Iod72kHvl+4GFeE9LZmRXjdcBYnQtjMTf85USIo2gxn4Wz6BFFjYgpY/LgpSoVdXOv6hXSAlerjQ8G/yhdToPZqcP9gVMqCnPyl/Bx/IhPakAgHouC73h9wsDRfQuXplrb6AmlYQblAOGXXN8vXUrTD+gAlm82FUsx2RQBgyaHZjrDmA8q9rdPZCrMUk/IEo0/VGs38E3+rWtGU6ouF1FoFAECowpcneD1mMf+uJeicp87fG1nqr6Ay4HNEZwHDl49YfE/DpYLWXYd9Y9BtA2F1RYJggGOvN5kR8Fx237KIluLePEv9B4fnRZgyJrdAKz0Fh1ovXwltUKehFsx1i7XKv9x+i35UXJLxVEOjEwP5tXXkf4jh3h5adK43TmcYZ87p3qyxaGRXREFmSPWOo9JEeQ0n/o8Y06Punj/g3GjZjirhAL7oIvQUAzdZK8dwjJEHsHg2hoHeP1QBdWV9ph0FLPJUChRxa/TSde4UZF+XhAunkV2QN0G1zV/Jyq8PWaRQpf98UUL2bxHFvMCIwmJC/qhak/CFzdd0vF0ylcrlEULuTdLKX43gCaWQJ12/QZxf5dlCpeQqPSUb8Asm7ac5LP941kUfyi2rBDroTH6SzJe3eyu6rQQ1t6nfGNoXbis7bq0ESYshQj0iMQY7EEf2ksdn8CKJK+SegWaAkEANaItGnrWX+0TzbPgvTNr7TQ8m8RYh8UIz8b7COK7FZrOVn2jyeky8vgneT5/cB8WhyP8A6zMN2v6WKRM5Ym4DxhqzAKrBCLJPo2mCzTLZY8pFLop53FtSy1Z8BFGLHeqJpzLiN/DhtzvDUWCDnl4qhclEyHuA7mcu8m47uW99jFzM4zGs+PyTN5ST5vwKjG2RBABJVCyKk/D6TBkQcHujFEDpSwm892MQmMB272wb0FvdLAl/YBJbjt6exUNSvAmxgE6gUFoOiq7QRqNkEhLXDc/GKG2xmE+gvEBAPGduTwRomgUCcf7FEOb4WVLDwc9W4QGhT6Go/OXZplc0zASSwCGQKVqiukd6a18QbLsWcCJv9qr7kDXDRMgh0xPk15qJJW7xILEVeSs9fJGgfGkz0CxdJdJPoBa4KSHK5ZXpHlREQhglbgjF6mULoCuVXUFH3DSlruWnUPcPFPWn4TzJ0JVOqKG6CNJa+av5KUv6pqWydYf0twt/vDxgxjqi3l2tHgz71nHIletS+9n+5Ui4+0973ixZn5XusC+x/Ipk0ULZYC3bAK4xu4kz8jtPuerDe+tdLxzwkJa99mhvkbeneuvcozTpkbu88CxGPcvZgfLzLeN7M5kea5pLTnE97njvmwbnN+3q4edWyZn+gdL7ROuHEf4a/6whiPQ08kqkAAAAASUVORK5CYII=',
    'web': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACdlBMVEVMaXH////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////8rISQAAAA0XRSTlMA+y3O/Qn29wED8PX+ye9xHPoCCivpzAYdNvkfYzDPbTPigKXLcvyiNe34HiRzdnCj6zHbFxsgqyrzL+gYkya7qJVFtAdmGj7nTJk3Ro50eZ3q0PR91LWBSCnyLhmfVifFFdGmERAEWVoPamxHa/GUMnpbf+aEiUssb6eNbpLVoe5OkTR344uQFLmarFiWFijXYMdEyqDTUg5BI55Qv2IFOXxUqeyKrlF1E9wiCLw7fuWDIXjfXuSCJUkLU95nqq9XOHtcpJc9zWUNt8FAur4/OsZo1BYAAAAJcEhZcwAACxMAAAsTAQCanBgAAAVpSURBVGje7Vr5X1NHEH/BkASEkCAQQI4iQgBBRUE5xCIqiqIWsVpBRaCFel9tvWq11VZrq622VuvZ+77v+77v/Y+azOTBe7uzb19I8vn4A/NTdr47M2/ZnWNn0bSbirZcLCu7eClh6uf9nM9CVHWmLq5qp1VEftzRxSKU3Rlh5eXGqv3O3/6sbMefHzvYCPn3Is/nvVZwVwzqnwhWMXYOfz9Yzgzk2I/cxYy5L8wfo/qsQWdI1XkP2upmJmpErZ5fQ7/Lf1o/BvWTX3SCpiUwCqxmHE0IANAGA+ehaDcj5WAj6lmI4x+YQKsQuRVHPcWZ0ejvTdfVPIcnySUamIgf/aw+Pn/AtnrPa35dajNyFjCCdiGWo4/zM2zqnzQ4quRV4NRXUQbceH6HRzkzk+3o900dlWgeAtYRRtIK3P+uUU7RHLX+5bUGFS8Dq85FG5jYD/BDBlZ2hUr/LceMKpKAV8YkdBDgJCMrNU/x/anG2UXILJIZyEH8ASPPe9z+97M16NFMSrjCe5ndNVyupOQfkRu4Dz/LzKydLNP/xYB55hvoFe/JDTwMgcpTwkURyWn1tHLiZ4BdyJjqb3Qbxy2lDZzgpb8E9nQrAzUw5W6eTfr0bgc3y/m6MZrR9Cg6ShrH9jcR8TOHF16NgcNtZcCVLB7UMKWnCAZEb1qGmZhZUie5CYzdL+SXRmHOOeKQC7ROz5wcdX/AGfhWlG0AYK21AUxIvSIw16y/KU0MZZiillobuAc3UNyotCSTgbmi6FRMBUxBG2DaBBFoNepf7xQnbMRqRWWgA6YdFQFnlsHATEJyJSBTVAbehWkrCSSojdM4xY0WkfHyCGDFTElfw8RdFFT1FmBLSLl99tyAsT0WE98BbBWJFQPWojawwmKp0wF7jMQeB6xVbeAonTUN8ehJEmsD7G21gbUwsZPEnoE9dpDYblmUFLIjVRvpFf2NEFRNy1WD3Gy1gdth4nwaDNdg/9LQNJArURuYjSmXBv+Tbg+bB3LNagMl2AWgwRdC0FYawvu1V20gG2/sNFgQgmpoCAsel9pAKpZPNLgtBD1PQ1g5lasNuDHtM6mnJdxAwv9ECd/kYRqqi9sxpaM1y42Lo4U7EZcTGSrC8SBAB7tZcQl2Q/Jw/ZT6emOsr7eT2ICs8A3Rd4AdUhu4ABP3khjeBj8ksa3xSJk1Fkl/jd2kj52vdST2vUXZchKwbWoDi+UtK1dgvCgcp/iR/BK4T3VK0V9mEMigrWtsh8rAfkmzgruJfyO7iLfbu4g/reol9MbWSsjsFhdQaN4FoobPst8MaRCBV/iXmh4yYWva+9YGpsCkAnH9b/IHSTwHp+00pPAacZqpu3Yp6eTtwV5LbUA4IWJLTWvy89vUb7cpeIo/5Y5CytsyeOlNdtuam3h2Ge3PG8muoI3G7DJ7jVktmStSPsKG8Ceq1rL2lZlbNEkWkubUUrWRsjnexDXHLZ5BZpnb75+q2vsHiIRv/URx3CtWhfLyLvJAYXrE8yoeWcxrKKSPl7m9a1ph6nJVaqjIFmqXs7JHolMAHzawuvLUyafe4NIvLRLemQx0GMChpdE9c4VOa5BvidST4cK9gfcyew914cCXb846dN3WgtjIgvNn2M/RDZt1qb9gnEs9lvoA+mXk2bMhmiogc8dOFPsbx3tEAycQ+QNHO3dkRllo+EqdhotIQPCFnBuG115nqW8MtUxfMHy5ysFw08c9YPR8hoEqnO3Lg31jLJeqr7j18lzrMF3lHJ9req50X6mO5V9Brv/jjaz+R0NC8g8jr73y2vUtsVZ9Pt05t4/4eHOb7va5cS0wzy6Al3F3S3/CathLVzMyrv5+c9XV/wMznaVgIelZVAAAAABJRU5ErkJggg==',
    'birthday': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAABOFBMVEVMaXH////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////10QvYAAAAZ3RSTlMAC+wJBMYl0QfzaT/tgAjY9ZlGb8ruydMN3NImBj5HMVTP+P2+PVVmj9c3fx3Cx/rgAkXQbvQW/l6CxemrbS/EZ84PlcgiTHCjKfwy1sFfiAzVSdSBJC739lYsZEqcwESRm0NICppcxcDC7wAAAAlwSFlzAAALEwAACxMBAJqcGAAAAgFJREFUaN7tmVdXwkAQhRWE0FEUkW5X7IqIomLvvYC9l/n//0A8ARKSnRyOs0+euY/37sxHstkNZ9PSIlZPdFiKj+oAYE+Gj2oX4HBIgo9qHwCiLrqPaae9UgDzZB+Tfft3PMzMEn1MPYOgKkTzMS1sVMfDKMnHFOiqjYcuio+qsz4elik+qm6tYJriowprBTaKj8qtFfRR/GZu0RbFb2aSxyg+KpdSG79C8nFNVCsmN2m+xTWsTrndDk8b1Wf9VaATxWcAAxjAAAYwgAEMYMD/Bcy9BXNAVi7z9SJsHz8BaXp2mvsnwyBRYdPZUVxqfwBfyQAYBMk6McwvSNd7A+BJPqCjARCUD8hg5xOy5EZXrywxgAEMwACK3+tsbXUmx/vMFZElb7ySea8VcXb+WUqOZy0BsZRug802trh61LJUzND+UstCCg7oSOi9Nv+AfnN0WWSNdWd5MaB/0fia622vtbixyPpPjdndrQgQTJtf1MVjtcdx0ZyVr9RMVHcUNAJ8Do9deNQ/UrD5CusB4dnkyJoNq7N7HD7s68F6AS1zPeAtVZz4p+iVzlQv/Mic3Ufwm1K8ULOLsnX/YYup663/7bN8HL4t2tcesCYeWjzLvyaw/h8RfPkYl5Z+SYYMK1sRfIxCNoDY7wZQ2Tj8gqy+qWSRjLdrBjCAAQxgwH8C/ABsSnrVS04x3AAAAABJRU5ErkJggg==',
    'github': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACslBMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////+D/333AAAA5XRSTlMA/vcBE+1J/PkCEv0FZhvuAxp3CB9ZfKIMeLo5KfSALwm24U0hMx3nBPb7MRyRNHvi6PPMRSC58n9eZ3onu910FDzcyM0ywm/WnX6IB+8rsdXFq/iWFWKGyapsLoMqTDAZP8tULURfvmkRz64LYansKI/ScA6zVp7mhCIs9Zg223ngwerAoeSMoIHEcQ/YibKw09ojTmBqjuVTk1o+dVtoN7jwQCWnHjrZUWNQa11z0EJHvaNlRq3RDbdV+rwmrI1PNRac332C6cpt1Fib3uOlQ4vOlXaaOD3DpFLXS58XO3IGJPHrN848NQAAAAlwSFlzAAALEwAACxMBAJqcGAAABY5JREFUaN7tWudfU1cYPiGBJISQEAgYxEgAGTJkgwwZZStaHIColamytFattdbV1t26ap117929995773X+j14y7xl3nZv+fv3Q9+N7zn2e3HvOecdzAsB/xYwxPX03ymy2KvPxa7/9Glzs6J6W4f0hkG8ldcurTqcHBT2htGEOpFt+vCNFJXpk4y4obrfbotjhJ8ZNhtJmWvMtG/zj8VCuPRmmHP6RNVCJPThRGXyiMwQqM1PcmwrwN06Dyi30Rbnw2kYDZLK4cbLwdS7IagcmycCPuguyW8jdkviLcqEa038ggT/eAlXaF6L4M0OhausQwU/NhUEwhyB+9WUYFFsmlFKKg4MPTYV0gnIYLHsvkoZfg8Tm0IJPlEA+UIBEl7YIEn9cEvJED3ekP3LpZaF/uaoIgAhkA04nCfYgz6RZ3c6Y7AxJ+F2fewDW8p2TM4nUiwa4t/w79zONKHxBjW/mUcTfacQIKtHnyniHu9P7Ug0XVjq+Npt32zq2xm/yfLy09gBOKorQheKHYT9sIz98HzricpzTYSmp3vxU0scP8V0rEIRWZAzcjxFEMiRxLAy/g8Q4/NNGMxAsRSEsq3ljs/4FAjglMFRE7JR5DASdGEZe4LRNIXbfeAYCItQn+4c+JAimKsefR4C86hvKJM/PXOUEdxAgBt/W7qDUtMoJppMoZu/QbWTuZig2dZsImPneOPoSMdLIUi3f2YrDrPAEkmcI/L1s9fhJAqje7V9HLE4MY0NBFPxVbvde3N3M2rGcxxPUD273TzjBY8w9ER5yCsacVryWzmfvHqdiUCXasV4Gf4Hv2bu6i5AS9geIMp+dQDuB8rWP4gQrVfS9b1PiHf7daCWHbNuPYZ3kfNdxgu0qCB7GsE5xPlsQIqnfPsWwTnC+WJxghwoCPOnEcr4qnMClQpcxYVg3OechnKCJnaCC1io8ijszjMwEeP0Gl3DOg0SQXc9M0ELL+4sJgjJmgpdxqDGNJJwgmM2K//x9OJQ77ecT/fpqRgJiOScI5PxTjARtOFCv2/0VQZA0gwmfrIw2uP1XyYLGxoKf2CSwXRaQBIbFDARzSZwcT5a4hxzZVqQYfzeJYvG0kmArpbmbNlOh8vwuBcRb2QE7rX0sadcqwL/3DA3jae/oIF1ibHtCLnzWdiqCvts3YbbPtTQK/HnzVqDG3lwhjR5unyUgIsb75yT7XLXc0U65aueljS3LHWGCLefgd13f1AnLDYEuJiLP5xu6CEZugJhaLIA3vfAG8cPnD7WKawz5PJUzkDb/GuWi+nDq39K10hIpESOb/3MCBZPmOoi+Aq5hak045QNJ6Fd65G4hm3eKOXX1Us0FZLKTtgKTxAnQoB/Fq/kuj4DRWzq+wKRJpa5xregL9KOTnWitUQ3O8wSqA/RNdEyMYC0uePEC0lnui5+rrBj29/+H6QRiy5yRhc+2o68weOY06D/Wa7EkFZcvkltD0IUKn+0LjB75w70uLjsA3Vbw82EBaURklZ+jqILdPNn6d3cCKR3lklQ5cP4SQSd4RXiFc2jz+wKaS4hPiTMWWRXUcX4rpT/wGq8csMvpvIXwiwVqw3S+sFZZaJW8JBTAz60WVGSQ/iFv38LjYQMDP5r3JAtcZtDxz4pceWXR9f3NArcN1MmGQrG3XkC9oSinT15P3UAS1zgJ2+SdmjHrp/3+PqmVS6GEsFiB1yVnznldxj0XWau2CNzWEhPTLsmpEWY0yzw3RGcxJFcHWobV+U76NExM1GxIlF9HNSCP7qTPqkd12INKKsH0EyW8ZxdKR1NTs05hLVudrZciyAngv88iFWc+62ur10kQ9F7RAiZL2OH5x0O7gFLqvrfSVBaq0B7Cu+q4C3uhjmqVAW7ZqfavFSBLZHOPRBrB/6bW/gHUffh5vlzVdgAAAABJRU5ErkJggg==',
    'user': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAABpFBMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////9bMWvtAAAAi3RSTlMAEwWP9BXt/vzzjPgk8fvrb/r5JfL128tZbRiRuUIEa/bXmceJ1hcHFhsBU87YVxnKhSPsza5RdVY3f2AiJ1C3mOWN0LG2wD/pBu4LTv1840NnLrT3FD2AWm6QpMhmiiiz3L3ho2qODbJ0ElWdwSFkAogmDHNUeCBEaM+W2Xebvx/oyYtsld2+n+IclCRQlQAAAAlwSFlzAAALEwAACxMBAJqcGAAAAq1JREFUaN7tmGdPG0EQhs/9zsZgG1dscMHGNBd6SeiEFjqEAOm9915Jb/Ong5QvrH0+b5mViOTn653eR9q7md1ZRaGlaTcaCy55vUvBWDTcpCBj2gsBQTFiQoz/sx+AMqyWDqR4s0cDXQJnsxj5YzNQkbdPxfO7NDDAPyiaf6IeDFHfiOVvQFUiQuujVheo4/z5dW6gQLvEm784A1R0mzkFL4GSO3z5HVZagVbgEuwDNdNc/S1AL7DydL49YCDOIQixCL6w5xccLAJHG7NgF5jIMwuibAILsyDGJjjFLPjBJggyC+xsggZmgY9N4Dp+AulLdEX2R5b+m0ovtDCbYJhZ0Ca72SlFFkErx34QYRHc4NkyrZK3TMVCL3BKPrb4OSeqD7SC57yjTTddfr9Z8uF3jP94PUhzfF8WGRAuVK/hR5JHqA3RIW3cb5TvXlaEaTT4l/rrMAblFo+/0iC+iDTrF27qFLV1GvNGxPTtPLE/OFrjd7EvXNrzlomc3eez5yYs+XalRo1jgO2ecyvZlzj85z+nGys0kfTC4eNEX/Ly1xc2xnRP6GhdOeaGNktf2Rx6QLxSjNA7Tk66yjvD1EBPeKTXls3aekfCPQNTOiPCu1dU8c+c9cCJN326en7zGgjwe75KfOeoCkKoq52Grf8+CHPd4BBjvg0IzLZUXJ9fgMK5xxUEo4DEqn7+vIolWFnXy7+1Bmhc06sHJyDyRKc/1GMKvJ/KBClA5XvZmcGFK9g+UyLwADJxkUtMjsnZ5sAWrJBrNAzodMkrgn9cJARb+IJZQvAeX5AkBA34ggVCsI0vSBACkEBNUBP8bwINP99NCHL4ggwh+IgvmCQEr/EFzeSeeRU7f65k03+I3LDtZVdVO3bU/B2dqfcAL/+n/lXbeiqDUA9aJnX0+/4F9ltf91/nTXoAAAAASUVORK5CYII=',
    'linkedin': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAABblBMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////8c2ar3AAAAeXRSTlMAmvQB7vFY75c1BgXZa/bH+f74/cMHBMxqCCxv1vrGywwbXmOw8xAVGslIzmzyebgCofwZpMhm7IooCUxRmTcYvVt/wqPmIkvwDZMSVZgXfsGiIacmaUpucxG2VNo63yDFpkSHJcpozwuRuoHEQ4aqLrRS9xQ4HbGVIkpHRgAAAAlwSFlzAAALEwAACxMBAJqcGAAAAdhJREFUaN7t2tVTQlEQBnDSAGkUFGzKQMLu7u7u7o773zuMILuIDzrfeXA8+/jtzvnNMNzLPfegUsmSBavC1RqNAiqNxWTIXL/yVIGWq5Gvb6hTwFVuZIBJgdcJA8rwwBUD/HhAzwBFQEngnwP6w55A8ZLbKwq4CH6kRbM2IUCk6DOvFwHoJknDJwAYoI07AUAJ62zBAQe/kfvggFXLOvv4jyjIOoN4IIc2Kh7xgIV+RmsiLrTqdD6iEwHYH1LxrlfMvcg2XZkIi02twm7XjuaYekMjf3D+AqDRlf4csOlI2ZNhmGTxj7WbqtoTD+kdJVMTOz8CcmkjLxnmk6wgMdRQQZJo7TEYiGQ8/qu0nlIk8Bb9ukWqteOAWNZNWBUMeC3Mvs27RQEv3+wjF6wg4NtaFA2siAbOrFBgfeZL5IQB2v4+s6KYnV08vkEB4/ep4QaWV6OAofT0Nc0PQMA2mZ6jjSMQME+m40+kEQABw3R8lL4TwgAhdj156JfLAQGeGXBJx80QoJMBbXRcDwE2GXCOB3IY0C0aUEtAAhKQgAQk8E8B+TJEAr8Dwvj1x0QfNToZ4MYDLfy414Vev4afEqhC4APl8t7MTbqh3gI7cvfvLRvlnxhk4eoduYJqUCldnKoAAAAASUVORK5CYII=',
    'link': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACcFBMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////8h3Fl2AAAAz3RSTlMA8QoJaQMC/gzCAesN7y+s1AT7qGruxA5bmLr8vhdPgK1Zhi7nI2tlkzvlB9XawJEQ+B8F8jOIxVMcsVTXhxmwjNtmlBNW8I6yUNa3+iGNMPk1eAZtCzT9Q5Ueri2EWskbqexMj9Kz9rgYPOEiJwh0EnmdfqJk9JJzEbZ9giCmq07zltl3WF0+JKpsr5BxtFL3dntc4tHkzChVyokrV6SFyIpRMjd6vZ6jS/Uxwcanlxa7nDpeP2NE6WhJKnIVmzka44Gl6Ejt0w+5n2HQy4tL4xJkAAAACXBIWXMAAAsTAAALEwEAmpwYAAAFZ0lEQVRo3rVa90MTSRSeEBNCAgQCGMALLWABRJoiRUBFBA8FBVRERM/zPHvv7e4sZ29nb9d77733Ov/SbWY3knkzO7uT3byfkjff++bNzNs3b2YXIQuyuvi/toxSB67/JaOtfdVqZK+sLJmNaZk9K2wbu7v5DubJzXkuW/gLpmI9mdptnX7zo1gkf5ZZ5J+ZhsVS/5wVetcz2FjyMuPmrziGzcgTSXHye6qwOUmJr4ekR7BZuRjPLHlSsHnJS6j/EclNqP+KOMoS6r8i+Qnmx/gp6/OTk3VqZcgdCudm8R7v4y6L/v9WmPoQ0T/nCxYw05r/w9UUqDqLQYxa4U8u9AKYd1cyBN2zwD+Og9wBUSP28iOvH8BaTPD3f2iaH6HgYwDYasjv3c3jn6gHXwKQ9w07+EnCf0UGO2noXSP+ZyfI+K8IWIUiA/7gj1L+K7KHBl836GCRpP8ILabR5WL+r3sk/UfoCMjZYvRRaX70B40PCMELA5Lzo0iZzAgWSfuvhB1tcUNYY5XDDs5qLVteq2nR2XKzaYt/RB0sh/yTtYZ0khB2OnlGT9ImflEH8wH/O1qpk+5T/7/QwdpU19A2M0Qd3KSxy3ppfowXNDI2I8Cp2wL+VJAlslR12DemqtwObDbsBVEREnSwFDjzKdGGrsXqxtPr4L4IbDJEM1QMCgh1Bd6jtR9RJrNgWBSKOljCy4sLlwGK4tgQZfbkW2zsd2/bqP0sobHziXIcUyC2Cvgb2P03X1EPqb/30eAr3NDFeHfUdhLDjwt09vfp5M9kGvwNTxkJlO9V2xV9TFOlW6d++JK3Yi8S5SV2gzioGueyLV/p1SeHyN+rNPpnFTSXHUI6aficaejSrX8eJ/8/AXuTVsgNMDn8O6J3vgXUabW69ZUakR8Dg+kacC2sA3IqiL5LXPnG1m/riKYjmZ9L0QX+OeAsrZxG8ztjz6fr64huAW3RdyQKPgM6GOKkliL6bOA6R7W+TJR5mJvtlMe5j5dx6i7HqKq20gMA9wPZRNkNp+JCFA5CKfkA0f4ecw730PwnAVObem0Aj0WBtfziB79CtMO6/JsdwGCvmoV3wiEEBrS0XErr1euV83r8iL3/WU70vWx2matarKG1O4jykB5/AfuQF3N38Ihc4hU0z8dmWobfxd6P1Wib7RR2CPtIwzZezC/l86NmNrvsibY9YNraeZuRmlvQUOR4z/CjO4L67Qe4/HgO0U+jlYc1+Mb9b7rZ+0+GPzumdRVonKBmsCZau1+09ZbADkqo5tfpxjfUhxYMrFnUAby/baIH6WyIbdz0LbeoTBfdP8M7yQ8AYHvlWKMvzD2I9SSZr344lUzjw6zq0zxdDKL3XdEMtdPYy0EW0jFK89c1YIkKtw1cePIwW0lU5kS3tPdhWEwRdZDBC3NGclvW52/Rfg9A/mteUQelEs6oGYdJH8IZQk/T4BVG/JNgcYp7xDagRAhK+4//QjIjqJXmrzktNinlFz+m+UFmMYyiz2T5N3kMOgB58YxofTn8AcOwA09y2qCU//iqYVjDhH9Fxn981G3YQSMw8QUl/M/vN3GrCPcDv9c0f4rHBD979HyVd8URPz8KM3vyLjgG92EL/Aj9y5hm0bfeB5qwFX40j/Na4MTYFcHgiRwe//hU8+/2eG8+HWtOvb0hs7X372Od2Jr/irwk//JEwv+IHEwwPzpdL8cvNT9662yj/7zznr3+RyKpKKH+R95jmX2PWOVB8UmF3xT/OWfc77ozzazDeVdCvwZwnLT4vUGZ+HuGrlobvsg4rktf+atN35SMcukbCtzILrk30gLvbwtvIXul9f7d4XXlnTjtxnX/jNshK1T/A5cWLGcTiJ0aAAAAAElFTkSuQmCC',
    'portfolio': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACDVBMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////9xxQJjAAAArnRSTlMA+S7+4zg2+AwvkAv3k5Iwke70CuT9NzmuDi1LAsrrtOpDJN8y4YQNxQFdYsnt8hQa/K2qZvZqP1gIh6NChQTN3DzHbzES+iHvK7EWwF6CYwYQ3sToBc5XGR7BJt3x9YGAWykXQBVP2G7S50fsiigJjxMYf+YDqEbwslB0VeB+y9VzfemMz/PUclOWNHd8v6/TcQ+awmUcrLCgpFbQe8i4jtZfEVmrmz5Ia2iZjWeFplepAAAACXBIWXMAAAsTAAALEwEAmpwYAAADjklEQVRo3rVa918aMRQPoB6CCxAFFUFRW/feu1pn1Vqt27pX995777333s3fWKooVrwkLyHfH+/zve83l7u8e3kvCDFhKmk+vO2c22TF2GoacdgMJd8+oUAh+WelG/tD9+HrZ7u4etzYAlaH4szMFVG3ZGZhGvKdSbzymztbMBOuZGznkH9YnYOZEVSUAJRP7+jFIPTVw16tA0Ohq8xjH36Dgjlw/DSj/sEtmA86w1YW/ZRUzI3aZLp+h4IFkB1L07+uI9weZggxGmOiSEMw7yLrdxK/du0SKTSIQNKUkfTHSPph2mVaKOkZrF3q+g+I82vwESOI0SleTT9DRzQI8TGDiUTzLZXvs5j8heh9VD2Z6Vr3ax2IxswGGyjU8xZ//QTq+mWeIg9u+hucoi6iKB85nBo1/F60lr6AldBlcqSGSu6dWTNBWSy/Fa9DZDQD2fa/QSZTpFEigvX62HANE7lqtX73JWjMofNrV/9F0+Axh87PWJU/jHLEHCo/tVzlAVhjDpVfspJftXAvKCL/2vLlLt6YQ+OXei87VG+wAw3W8Au9SZBAzCHzm5ZSpXGBmEPhzy9eXRCJOWT+u8X9hVjMIfIVo+dykWDMIfJnPQaVWCJ+eQzcMg3mEJqRqY+tdjQt1cCzmH/INfjDFjD58RsVyDWwoU1yDTaibXIN6lCiXINEtEOugRlZ5RrkoyahjIhhinJEMiI6XiOXQEbEgJdoTiAjYsArNCTwu6dj1I4K+TZJjHiDUIVMA7dnr/ZR4hSZ/2UDIQIZEU1/caNWruHPiMhIbVzSeiuQEZGGP270Dva9SEakUjW6f3Sie2W2vwtp6Xou1LeTK2nlIgG7II6hSPecW950kqnK2M+rHzTJWALP4Rx/K2sh9hGfQT9zpbdVx6PvBJSq7/IYpAAMjnHoO0DV9h64wR2QQRLcoArWkGgDGzTDDIbBWQS0+3QIajAANKjpAxrEQLtOjcBJKgL3tRqgWxco0gtBBsoU2GEfbDd1Ed78awdtpw5z9Bf3miAOPE3S0qsAgzSeDmww4BnucfV4tTsBW2AuNNdJNkB5Q4wGJ3g74TUVbP/oPfy9/PhsFoNpgdMCg066vuuI0HGH2TM0g8uCByos1eQcfncuEsXgfjMhbb8RiEMnB86qFntuB+hci6WsYL0OZ/ETFDjoy7641hRDnw2jAGNyosKWFW3ShJlGHr8oeQq59S93tkBme65lnQAAAABJRU5ErkJggg==',
}
# 简洁单栏版式用的深灰图标（#595959）
ICONS_DARK = {
    'phone': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAAApVBMVEVMaXFZWVlYWFhZWVlaWlpZWVlZWVlWVlZ/f39VVVVYWFhYWFhXV1dYWFhYWFhbW1tcXFxZWVlYWFhYWFhZWVlZWVlXV1dYWFhZWVlZWVlYWFhYWFhaWlpYWFhYWFhZWVlYWFhYWFhcXFwAAABXV1dYWFhYWFhYWFhYWFhbW1tYWFhYWFhZWVlZWVlZWVlcXFxaWlpZWVlYWFhZWVlYWFhZWVlZWVna3J8PAAAANnRSTlMA7Zk2MPP4LwIM7uM9Ltg4JNb0wJL5Q5slkfKQRPe+zkLXCwEmkzHBRQ6cz+Q5vyE+mjfN673D32sSAAAACXBIWXMAAAsTAAALEwEAmpwYAAABOUlEQVRo3u3aSXOCQBQE4BFZRVTcd6PGJdFEs83//2k5JIFLpux5ZA6p6j4370MUsGCU+j3HUaur4XTXL3Nlk/uDtk5vjM+/NrQgjSa8/6L5Wtc6ILDUwhyw+UctzhACHuWABwEtObCGgEAOBBBQ1N+yLdLfRqtiEzsgQn/XvhCIUSAWAvipT4AAAQIECBAgQIAAAQIECBAgQIAAAQIE/gRw/uQXfna9FwIrH/oM8f4kBAQhQOAWsJjVN5vXSegKSLOvQpS6ARbZTyMKnQCzstJ3AtTLStsJkJSVxDXw8D8P0aSs7CoBU8PGYXED8geGylRVemGdfgv+2dTIIWBkPE/DfjtJnnaDiq/c5/Jr3R12D/+Qzu+BfxIu744XbqhmTTQfXnqi1PjZfv6yo2wy9HKL5Q9B7pm+30+kOWgXz9xUnwAAAABJRU5ErkJggg==',
    'email': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAAB3VBMVEVMaXFVVVVZWVlZWVlZWVlcXFxZWVkAAABYWFhZWVlZWVlaWlpZWVlYWFhfX19YWFhmZmZZWVk/Pz9ZWVlYWFhYWFhZWVlYWFhYWFhYWFhYWFhYWFhYWFhXV1dZWVlaWlpVVVVdXV1XV1dYWFhZWVlXV1dYWFhYWFhYWFhaWlpXV1daWlpZWVlYWFhZWVlXV1dZWVlbW1tZWVlYWFhZWVlaWlpYWFhaWlpZWVlYWFhYWFhaWlpYWFhiYmJZWVlVVVVZWVlYWFhYWFhaWlpYWFhZWVlZWVl/f39cXFxZWVlYWFhVVVVYWFhYWFhcXFxYWFhISEhYWFhYWFhfX19YWFhaWlpZWVlYWFhaWlpZWVlYWFhXV1dYWFhVVVVZWVlbW1tYWFhZWVlYWFhZWVlZWVlYWFhYWFhYWFhXV1dbW1tZWVlYWFhYWFhYWFhZWVlYWFhZWVlaWlpbW1tYWFhWVlZZWVlZWVlYWFhYWFhYWFhVVVVZWVlZWVlYWFhZWVlYWFhWVlZYWFhZWVlZWVlZWVlaWlpZWVlYWFhYWFhZWVlYWFhYWFhVVVVZWVlZWVlaWlpZWVlaWlpYWFhYWFhZWVlWVlZZWVlZWVlYWFhZWVlYWFhZWVlXV1dZWVlZWVkVGvw0AAAAnnRSTlMACfnzgCRvAe7+kUT8+hD9Be0EsTTeJUKzmcGIykOrEQMTPar4I9jdvncdPhTX+ybCONbvNkH0IqK7wFXPDd8Y5JsaMONTOQIhzuwM8XAWfwfH6witZuHbTPBFICsPvA5U5Zzqy4Lph10nl7BOtV6Wpk8cyTLU2a/ycRK/hlyLwyykTdO0H5rSUeK4VhvnR1LNM+aE9jW9PJij4NE69S/VufIAAAAJcEhZcwAACxMAAAsTAQCanBgAAALySURBVGje7ZnnWxpBEMZXEO7o0psoVUCiUTRqbImxxBp7S++999577z33twYjuwcIx53sPI9Jbj7xcHO/OZblfWcWhOSQQw451kps76lp4ahGS91jHcGv93Ag0e5Y5idYDijY4J/nB+NznFmVKhDjAMOT+n450PCh/bAF3KgGtkAdUsIWUCIOOP6BAnpYvh51wxaIoiOwBW6ikU5I/tAIQkYrHN9qXJJTRm0A+oLLt6YdJ/Eegn+AtzS0bTd9fsOWTFO2lJvo4jXaihzfv0/V2Eb3ruwsHlXS43eqeK7tKZN+VXaGFv94WRrJvGhGLHeiGhd7Q2WZ2H7MC5zi2CU1fTaA36ltKJ1/oRXTPtmxXNfHGbKb9LR2D6NWZPiBlSzThs2l8A+/xpw9VdmGM9eIr6xzrZ6/cROmvIvmOpohjrI/m/SoJ9qDZhV5LLPKSLbu/Gr48zZ8f/VMfk+e+4Izur5J5/PP1xwtZPr15X6cdFpiv/RRza+wQaCriBEJPDYkhX/pLhHm58Jty9QCzpzQakTz27A2IN2rYn2RRjuBk08Oi8MvduA7Klaqfp7G660D558XNVtNDuL8HZPiOrurR/Ed/nhRu9Z7ibN0NIluHUNkyzXahfl9CZxZppXSm3YT5dgVEuxLAjjv4mVpza/iGhZYZqzgMinGSJI6LLm75oXLVmCZ+n6KEEiB9n2fk9z/MN91V0CMxCNx5jH7K/diWI2Xxy9oUsIDSGUEP+TOJznDXa9Imy0y4TSdI7vwQ4ZLhA8RbVgwlzZCaW5byJnJ2bQOmO6R37rFqyl5RrtTS3oo1fhBlq0c57uqlx4aQ+DirULHTP3DlKbM0PV8+C43vTH2u24l32mnOidfUWXjW2doD+Im1wCP720zQEz6D7y6yPR0xOf98X+dVcgF5AJ/QQHgY81R6IPZJOqBLfAV6WALpHqrz5D89pS2D05xsH9QoKAZjB9c9ifHDRh+jLdZnztJ+fegTLqd8t+Hcsghx5qJ3wOUnkWsXwrMAAAAAElFTkSuQmCC',
    'location': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACN1BMVEVMaXEAAABVVVVmZmZYWFhZWVlISEhYWFhZWVlZWVlYWFhVVVVYWFhZWVlVVVVfX19ZWVlZWVlYWFg/Pz9ZWVlZWVl/f39ZWVlXV1dVVVVcXFxZWVlVVVVZWVlYWFhZWVlXV1dZWVlVVVVVVVVYWFhfX19ZWVlXV1dXV1dVVVVYWFhYWFhYWFhYWFhZWVlYWFhZWVlYWFhZWVlZWVlZWVlXV1dYWFhYWFhYWFhYWFhWVlZYWFhYWFhVVVVZWVlZWVlZWVlZWVlZWVlYWFhZWVlZWVlYWFhZWVlZWVlZWVlaWlpYWFhYWFhZWVlZWVlaWlpZWVlYWFhYWFhXV1dMTExZWVlZWVlZWVlZWVlYWFhaWlpaWlpbW1tZWVlYWFhcXFxZWVlYWFhYWFhXV1dYWFhZWVlYWFhaWlpYWFhaWlpZWVlYWFhaWlpYWFhXV1dZWVlYWFhYWFhZWVlZWVlYWFhZWVlZWVlYWFhZWVlZWVlYWFhaWlpbW1tZWVlYWFhXV1dYWFhZWVlYWFhZWVlZWVlYWFhaWlpaWlpYWFhaWlpZWVldXV1WVlZZWVlYWFhXV1dYWFhZWVlYWFhZWVlZWVlYWFhYWFhZWVlZWVlYWFhZWVlZWVlZWVlZWVlXV1dZWVlYWFhZWVlZWVlaWlpYWFhYWFhXV1dZWVlYWFhYWFhYWFhZWVlYWFhZWVlaWlpZWVlYWFhYWFhYWFhYWFhWVlZYWFhZWVlZWVlZWVlYWFhbW1tYWFhZWVly+JeGAAAAvHRSTlMAAQMFF/kH/fz+9xtzrgYI9vDvBPjZAu0dDBb7FVDxiSDlHgnuEORXPRLdw0Lse1GAN6l1wjrSinCzNawrD7ySeH0/fCWrxHJYOUSVdobOd9PXVnQKzdx6W5NMLTiowSTRRYhdu2+wY41gNvQi1Ul+VNpk6sDLSoLQsYc7J5fPQ05elsWm3k8w4xGMEzLU8iNxv9il6L5ci2yktBS5mkbiuGHIM+aEWo/HZeuUzL1m4frgX8Yvp/XWujQcsp/7UGoAAAAJcEhZcwAACxMAAAsTAQCanBgAAAPHSURBVGjetVllUxtRFF1IQhIIgaRoQgju7lrcvWixtmjd3d3d3d1d34/rTDudMuTet0825/O558zuvr3viqKwwbdstnimKC3BYEhIK5pxz/7wVTTEz5LSaDIPtl/TAdqox09OGQkI0/erQdLykdOxhIJNOZVS8rr8cKKCcItOXD+lmzDgYqugfFCbmTDBXOUnoh/zkDCjfwm/fms04YAtg1c/259wwZDMp//WSDhhfMGj7zIRbhgXsetnhBABGO6y6tdHEyHYzrHp+/UTQbxm+x/aiDCqmPKDWdzAXKuuv2ApkUC3eubLJ1I4r5r/M+UMMiNVDLqIJCwq9+NlPNQ03jTq0Okco5YCyo/+mH6LTqKBwTuj/tMc7mCUOEQ1mMLC8uaVEBENGPM+TT8AeXZzjs98qk8XknCNtGqmhCdRJiMO1ykGpXBINczOgtmBlL/YBr9/hO5TDtIT8KqyDD4/6EuNgO/VR6jBRpDfgT+yGwx4gPKLwf8rCjfQg6duAuXvgOibaec6ne8rg5m6iWbQC0XsRekfIHouzaAQijiC0sFUHUUz0EMRL1G6AaJbqQUCFBGC0k3aGJhQerA2r+grSrdD9BaawS6w6UHpA9x3YDsUEYrSn0P0AprBISiiEaWPgZ/MQfkE4JVwGOXngLnLjRucBAPq8LYGTtcRGD8OTtf7UYNbcFna4INcOIHwBZuKP/J2+BIchNk1MLuTcijWIZc+2OE1I5d+FsVgN9Z/DXqWLXVYn0jtmQuwcqp83peOC8SYV6iV3Xq0IvQv1s85/9fwNtpFbz/SKH1qem+h3mrVF+Y3UrroDSqjsBWy5fuQWgcVKqf/TXWW55IzGFbvAm/L6C9kaGNHnOL67ytYGuUJcYMnbKOEHlH9UMYpZ61gr29MYR233BEzOMg8L6pMEtGPTWSfSA2LGBzlmamN8evf4BoKJt7j1U9K5Rs7LuM8Sea1vIPTPXwGi7knv31caXUgkttAaeGYnTpzRabjHewG28S2E1tY9TutYguEEcYBsH+FIggLm0GJ8Iol7A3TLRYmvsQJCFfXt8fJrKGeqhs0y+3RvqjpX5Jc1CW+0zbHeWIfddlizFCk8YxmsFyDZalfEa6/Jl6LdewxA6Yf8lmbhW81ZtCu0b467AKsn67ZVjwGXEvZYrTbuX+EDA5oudXP45ogC+DEcY867pSmBh5ljHmrojFeyZYpqgv4M3P1e/o0N1BOz/mhnSsVL6CGbd4h8UOP/9P/5OsVA2X12b/6q+oVLyH7TzXpdCleQ+5Nu720jCvkN4wqLzwXzIsLAAAAAElFTkSuQmCC',
    'wechat': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACl1BMVEVMaXFYWFhYWFhYWFhZWVk/Pz9ZWVlZWVlYWFh/f39YWFhYWFhVVVVZWVlYWFhZWVlaWlpZWVlYWFhiYmJYWFhZWVlXV1dMTExfX19ZWVlYWFhXV1dZWVlZWVlZWVlZWVlZWVlaWlpYWFhYWFhZWVlZWVlaWlpdXV1aWlpYWFhYWFhWVlZZWVlZWVkAAABbW1tZWVlZWVlYWFhVVVVZWVlaWlpVVVVISEhYWFhZWVlYWFhcXFxZWVlZWVlZWVlYWFhZWVlYWFhYWFhVVVVYWFhYWFhZWVlZWVlYWFhZWVlYWFhYWFhYWFhZWVlZWVlXV1dYWFhYWFhZWVlZWVlYWFhaWlpbW1tYWFhXV1dbW1tXV1dYWFhVVVVYWFhYWFhZWVlXV1dfX19YWFhYWFhYWFhVVVVZWVlcXFxZWVlaWlpaWlpYWFhYWFhcXFxZWVlaWlpZWVlXV1dZWVlZWVlYWFhYWFhWVlZYWFhZWVlVVVVYWFhZWVlYWFhbW1tYWFhmZmZZWVlYWFhYWFhZWVlZWVlZWVlYWFhXV1dZWVlYWFhYWFhZWVlYWFhbW1tYWFhXV1dYWFhXV1dVVVVXV1dWVlZZWVlZWVlZWVlZWVlZWVlYWFhYWFhYWFhaWlpZWVlYWFhYWFhYWFhYWFhZWVlZWVlZWVlYWFhYWFhZWVlYWFhYWFhaWlpYWFhZWVlYWFhZWVlYWFhZWVlZWVlZWVlYWFhYWFhYWFhZWVlZWVlZWVlYWFhXV1dZWVlZWVlYWFhYWFhZWVlaWlpbW1tZWVlWVlZZWVlZWVlYWFhXV1dYWFhYWFhYWFhaWlpZWVlZWVlYWFhYWFhaWlpXV1dZWVlYWFhZWVlZWVlYWFhaWlpYWFhaWlpZWVlZWVn7/I9sAAAA3HRSTlMA8hox7QT+/P0C+rIGKKffRM73Db5NHQoIjs8mR8L7+eFV8TSPfUETZusuMvarARm3ooIJtk8DB629FySAqRSY+EIrFVbbqKOey3nXyciuWkiwv9ngEQ7aIxwgbg/0u5dDEFzDigy5FudSM+ZlC3UtiTpQ1paFLO9eGNjNXyebBWykr6XQ0+5d6oiE5JAqxD24QBJXNX7wb9zRtayNIp3Vc5zAabGSym143nww44xi86pTulhxlZ+LtJrSRuJhmTdbTDhqL0qGRXTdzHAfg2dLx3dJ6MblJaE7s2DU4C809wAAAAlwSFlzAAALEwAACxMBAJqcGAAABH9JREFUaN7tWWVbWzEUTimlt0ChUGC4DdsYssEY25DBkDHBBhvM3d2dGczd3d3d3d09P2YVWG9yLUnLlz19P8F5Tt63yU1OzjkBwAknnPjvMTdxU2jUI8PjAI3DqT2nXvU60AY2QRvp0c7f6DD2kWXZ0VCI8OGTAhxBf2awCkqBO1xg72ptTYfyiOzvaQd9QKYbVIT3sGRGel1/FSRCx85M/DmjICnaTGDgb3kXUmBJC1r+uGhIhZ5d6findYGU6OFCw7/BF1Kj92hy/hBXyIBegaT8HXwgE9qSCgSh47rcnHK6NFxA5+qVtfo4alpBuEE5ZNRly9dTt8L4g4aazU9QzU5EAmHIoPmNseYhyr+20dkLsRaT8PujTFcazYERfKtb05Zpj4bXWQQCfqjAuCb7GGQx/q0n6p2kzN8aW+stojPgckRnAMOVj9sITMCng9Vehn1j0W8AYXdFgemCYKw3mxPxXbTXsosW494+Svz7hOdHmzoltkIrPAdFWc9fC29RpaAXxXaItcu93r2P+lBekPBWQaAdA/u3teV9iOPcRVp2rjROZxpnzOverbJwZXREQ2RJ9sxp0kd6ECX9jxrToO+fPuLfaPrMceICPeki9FYANFsqxXOP4ASxezSYgt49VgN0ZX2lHcYv9lQKFHJo9dN07hVmXJiHC2SSX5E1QLfNXcnLbRG2ToNJ+fumgOrdJI63QxCBCYT8VbUg5Q+Zq+sDuWAqlcslgtqdpJO9EMcTyCBLuH6DPrvIt4NKzVO4TjLiF0jeTXNexvCuiTySX1QLdtCdeD+dLWn3VnZfDWpoU7+DtimcV3TeXg2SEEOqemRiDHIk9usHTM3mRxBVyj8BzTwlgRjQEok+HS3zj+LZ7luYtvGdbknmLULkg2Lkf4N1XInNYi0/M2bznHx5EbybPL8PiEeT+8nWYR62+y1NJHLG2gSMp2QFVoEVYplE1wabZazFko9cEvW8s6CWrf4MoAg71kNMY/qt58dpc4anxgI5v1QMlYuS8QDfyVz6jUx03/ruOZfKYV5z+CFpASfJ/wUY3SALgpCoWhAh5feVNCDi8EAvhtCBEn4bwQY2geHY3Ta5t7hfEvjMJrAcvz2NRaJ+FWATm0ClsBgUXaX1QM0mIKwdnolX3GAzm8AIgYB4yNieDNYwCfjj/HUS5fg6UMHCz1XjAqFNoavt2NilVTbPBJDAIpAuWKFFjfTWvCbecCLmaNj8V/Ulb4CLlkGgPc6vMRdN2uKFYinyUnr+QkH70GCiX7JQoptEL3BJQJLLLdM7qoyAMEzYEozRyxRCTym/goq6b0hZyw2j7hnW9aLhP8bQlUyporgJMlj6qvkrSflHpbF1hvVnCX9/IGBEP6LeXq4dDfq0Q8qV6BH72v/lSrn4DHvfL1qclO+xzrP/iWTAJNliLcQBrzC6ITPwO06769pG618vHfOQlLzuXq6ft6V7696jNOuOub3zIlg8ytmDifEuE3kzmxNpmktOcz7tee6YDRc07+vh5lXLmv2B0vlG64QT/xn+AolHYj3FTDdlAAAAAElFTkSuQmCC',
    'web': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACdlBMVEVMaXFZWVlZWVlaWlpYWFhVVVVZWVlYWFgAAABVVVVZWVlZWVlZWVlYWFhYWFhYWFhbW1t/f39MTExYWFhYWFhYWFhYWFhVVVVZWVlXV1dZWVlaWlpaWlpYWFhYWFhaWlpZWVlZWVlaWlpZWVlZWVlZWVlZWVlZWVlZWVlWVlZZWVlVVVVcXFxYWFhYWFhXV1dZWVlZWVlZWVlbW1tYWFhVVVVYWFhYWFhZWVlYWFhYWFhWVlZVVVVYWFhXV1dYWFhYWFhZWVlZWVlYWFhISEhYWFhaWlpaWlpYWFhZWVlZWVlZWVlZWVlYWFhXV1dYWFhZWVlYWFhYWFhYWFhYWFhXV1dYWFhYWFhZWVlXV1dYWFhaWlpbW1tZWVlbW1tZWVlYWFhVVVVZWVlZWVlfX19aWlo/Pz9YWFhVVVVXV1dZWVlZWVlYWFhYWFhYWFhYWFhZWVlYWFhYWFhYWFhYWFhZWVlWVlZYWFhZWVlWVlZaWlpYWFhZWVlZWVlYWFhZWVlZWVlYWFhYWFhYWFhZWVlYWFhZWVlZWVlZWVlZWVlYWFhYWFhZWVlYWFhcXFxaWlpYWFhaWlpZWVlYWFhZWVlYWFhZWVlaWlpaWlpbW1tXV1dYWFhZWVlYWFhZWVlmZmZZWVlYWFhYWFhZWVlaWlpXV1dZWVlZWVlZWVlYWFhaWlpZWVlZWVlZWVlZWVlYWFhZWVlXV1ddXV1YWFhZWVlYWFhYWFhcXFxZWVlYWFhYWFhZWVlYWFhZWVlZWVlcXFxfX19ZWVlZWVlbW1tYWFhZWVlXV1dZWVliYmJZWVlYWFhZWVlYWFhXV1dYWFhZWVlXV1dZWVnyZNh8AAAA0XRSTlMA+84t/Qn29wED/vD1ye9xHAIK+ivpzAY2HfkfY89tMOKAM6XLcvyi7TX4HiRzdiCr8+gqcBvrMaPbFy8YuyaTlai0RQcaZj70nerQfbUpgdRI8i6fRpk353R5TBmOJ8VWFabREBEEWQ9aamyNbtWhku5ONOORMpCLLHdrFLnxmkd/5oSUS3pbiW+nrFiWFmDKRKDXKMfTUkEOI2K/nlAFOXxUvCJJfuWDgjt4317k3iVXE6pTr1whZ+yKe1GprgsIddw4pJc9tw26vs1lQME/Ou2dZEYAAAAJcEhZcwAACxMAAAsTAQCanBgAAAVpSURBVGje7Vr5X1NHEH+RkASQEAwkgFyKEEBQAUEFK1JEKV4tiBeKAi1YtdXW1tajXvWo1qO1tda7933f933f+x81meHBe7uzb19I8vn4A/NTdr47M2/ZnWNn0bSbirZcLC29eClh6uf/nMfCVHmmMa5qp1cM/7ijgw1TdvcwKyc3Vu13/fVPeRv+fNvBRsh/Ank+79X8O2NQ/3SokrFz+Pv+MmYgxzHkLmbMfX7BGNVn9TvDqi540FY7M1ENavX8Gf5d9tOKMaif/LUTNC2BUXAp42hCEIBWGDhPRbsZKSdrUM+tOP6SCbQekVtw1FmYGY3+rnRdzXN4klyigYn40c/q4wt7bKv3fO/XpTYhZyEj6ChiAX2cl2FT/6T+USVfAae+kjLgxvM7NMqZm2xHv2/aqETtILB2MpI24v53jHIKZqr1r2kwqPgCWI0u2sDEPoAfNLCyK1T6p95nVJEEvFImoZMAJxlZaTmK708zzi5AZoHMQADxe4087y7738+WoUczKeEKH2V213C5nJJ/WG7gMfwsM7Nhskz/twPmmT+gV7wqN/AABCpPMRdFJKfV08KJnwF2EWOqv9FtHLeENrCXl/4I2DOsDFTDlLt5NunTux3cLOcHxmhG0yPoKKkc299ExM8AL7wUA4fbyoArWTyoEUpPEQyI3rQaMzGzpG5yExi7R8gvNcKcc8QhF2ilnjk5av+UM/CxKFsHwHJrA5iQukRglll/U6oYyjBFrbU2cAQ3UNyo1CSTgVmi6DRMBUxBG2DaBBFoMepf4RQnrMJqRWWgB6btEAFnlsHAXEJyHSBTVAZegWnrCCSkjdM4xY0WkfFyJ2CFTEmfwcSjFFT5HWBLSLnD9tyAseMWE98AbD2JFQLWrDaw0WKpMwB7nMROA9aiNrCDzpqGePQUibUC9qLawHKY2E1iz8AeO0hstyxKCtmRqo30iv5GGKqi5apAbp7awO0wcQENRmqwf2loOsgVqw3Mw5RLg/9Jt4fNB7latYFi7ALQ4AthaCsN4f3aqzaQjTd2GswPQ9U0hAWPS20gDcsnGtwWhp6nIaycytQG3Jj2mdTTEm4g4X+ihG/yEA01xu2Y0tGa5cbF0SKdiMuJDBWReBCkg92cuAS7QXm4flJ9vTHW14dIbEBW+IbpfcBOqQ2ch4knSAxvg2+S2NZ4pMxqi6S/zG7Sx87XShL70aJseQiwbWoDi+UtK1dwvCgcp/iR/BJ4WHVK0V9mE0i/rWtsj8rAMUmzgruJfyi7iLfZu4g/oeoldMXWSshsFxdQZN4FoobPst8MqROBb/iXmk4yYWvaa9YGpsCkfHH9L/EHSTwHB+00pPAacZCpu3Yp6eTtwV5LbUA4IWJLTWvy89vUZ7cpuJ8/5Y4iytsyeOnNdtuam3l2Ke3Pq8iuoI3G7Gp7jVktmStS3sKG8Duq1rL2rplbMEkWkmY2ULWRsjnexDXHLZ5B5pjb76+r2vt7iIRv/USxyytWhfLybviBwvSI51U8spjXUEQfL3N717TCtDWq1FCRLdQuZ2WPRPsB3m5gdeSok0+9waU/XyS8MxloO4CDa6N75gqf1hDfEqknw4V7A+9l9h7qIoEvz5x16LqtGbGRBefNtp+j6zbpUn/DOJd6LPUB9MvIs2ddNFVA5r4DKPY7jo+LBvYi8iuODuzLjLLQ8JU4DReRoOALgRuG115niW8MtUxvKHK5CmC46eUeMDrfw0AVyfZlod4xlktV19x6ea71mK5yjpc1PVe6r1XF8q8gV/7wDq/+E0NC8g8hr6386pUtsVZ9Pt05D434eG2r7va5cS0wzy6El3F3c1/CathL1zMyrv92c9XV/wNUSqVg60qaiwAAAABJRU5ErkJggg==',
    'birthday': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAABOFBMVEVMaXFYWFhZWVlVVVU/Pz9ZWVlcXFxZWVlISEhZWVlZWVlZWVlYWFhZWVlYWFhfX19ZWVlYWFhYWFhYWFhZWVlXV1dYWFhZWVliYmJZWVlaWlpYWFhVVVVXV1dYWFhYWFhYWFhZWVlYWFhZWVlYWFhXV1daWlpYWFhVVVVYWFhYWFhZWVlYWFhYWFhZWVlZWVlZWVlZWVlYWFh/f39YWFhZWVlZWVlXV1dcXFxYWFhYWFhYWFhZWVlYWFhaWlpZWVlYWFhYWFhWVlZYWFhaWlpZWVlXV1dZWVlYWFhZWVlaWlpYWFhXV1dYWFhWVlZZWVlYWFhYWFhVVVVYWFhYWFhcXFxYWFhZWVlZWVlYWFhYWFhZWVlXV1dYWFhZWVlYWFhaWlpYWFhWVlZZWVlYWFhMTExZWVlZWVlBSnDSAAAAZ3RSTlMAxmkJBCUL0QfzP+3sgNgI9cruydNGmW8N3D7SBiZUz/34vkcxPVVuD/SC0Olt/l7Fq8QClWfOHRbXN3+Px2bC+uAvRSLIKfxwo0zVSYEy1C73DMFfJIjWZJzASkObkUhEXCyaVgr2DtIqnAAAAAlwSFlzAAALEwAACxMBAJqcGAAAAgFJREFUaN7tmVdXwkAQhQWE0FEEFRAQO2LBLnax94K99/n//0A4ARKSnRyOs0+euY/37sxHstkNZ9PSIlZ3bEiKj2oXYF+Gj+oQYG9Qgo/qAABibrqPabutXADDZB+TY6cyHjZniD6m7j5QFaX5mBbnq+NhheRjCrXXxkM7xUfVVR8PYxQfVYdWsEHxUYW1Aj/FR+XRCtIUv5lbNEvxm5nkJYqPyq3Uxk+RfFwL1Yr1EZpvcQ3jqx6Pzd5K9Vl/FehE8RnAAAYwgAEMYAADGPB/AVvPwSyQlS18vgnbJ0sgTQ9Oc/94GCQqbDo7SkrtD+C7NgD6QLJKhvkF6XpqANzLB3Q2AILyAQXsfEKWPOjqlSUGMIABGEAJeJ0ulzO+nDZXREa9yXLmPVHE2eX3dXw5ZwlI9Os22Fxji7NHLetPGNofa1lUwQGdPXqvNdCr3xzdFllj3c2EGJCZNL7mUm21FhcWWebKmH39iADBAfOLunir9rgtmrPzMzUT1Z0GjQCfze4QHvWv5f2+/FxIeDa5Nu3H6hx2mw/7ejCXR8vcd3hLFSf+KXoNFKoXfmrOPiL4TSkeqdnRuXX/IYupS9X/9lk+Du8W7WsPWBMPLZ5NvPRg/V8j+PIxLi39kowaVrYi+BiFbACJygZQ3jgCgqy+qeSQjLdrBjCAAQxgwH8C/AJwq3rVb0vwPgAAAABJRU5ErkJggg==',
    'github': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACslBMVEVMaXFZWVlYWFgAAABdXV1ZWVlXV1dZWVlZWVl/f39VVVVYWFhmZmZaWlpVVVVYWFhVVVVYWFhaWlpfX19aWlpYWFhYWFhZWVlZWVlXV1dZWVlZWVlVVVVYWFhZWVlWVlZVVVVZWVlZWVlZWVlcXFxaWlpXV1c/Pz9YWFhbW1tZWVlZWVlZWVlZWVlZWVlZWVlYWFhZWVlZWVlYWFhYWFhXV1dZWVlYWFhYWFhZWVlZWVlZWVlbW1tYWFhYWFhXV1dZWVlZWVlZWVlYWFhYWFhZWVlYWFhZWVlZWVlYWFhZWVlZWVlYWFhYWFhWVlZYWFhZWVlYWFhZWVlZWVlZWVlZWVlISEhVVVVZWVlZWVlYWFhZWVlZWVlbW1taWlpaWlpbW1tZWVlYWFhZWVlaWlpaWlpYWFhYWFhZWVlaWlpYWFhZWVlcXFxZWVlZWVlVVVVaWlpbW1tZWVlYWFhYWFhZWVlYWFhZWVlYWFhYWFhZWVlYWFhWVlZYWFhYWFhYWFhYWFhYWFhYWFhZWVlZWVlYWFhYWFhZWVlYWFhYWFhYWFhYWFhYWFhZWVlYWFhYWFhZWVlZWVlYWFhXV1dYWFhaWlpZWVlZWVlZWVlZWVlYWFhXV1daWlpYWFhZWVlZWVlYWFhYWFhZWVlZWVlXV1dYWFhVVVVXV1dZWVlaWlpYWFhZWVlYWFhXV1dYWFhZWVlYWFhZWVlYWFhZWVlYWFhYWFhaWlpYWFhYWFhYWFhXV1diYmJZWVlZWVlYWFhYWFhaWlpYWFhZWVlXV1dZWVlYWFhYWFhZWVlZWVlYWFhZWVlYWFhZWVlXV1dcXFxZWVlYWFhWVlZYWFhZWVlXV1dYWFhZWVlbW1tZWVlYWFhaWlpYWFhYWFhYWFhYWFhaWlpZWVlVVVVcXFxYWFhYWFhZWVmPMFkQAAAA5XRSTlMA/vcBE+1J/PkCEv0FZhvuAxp3CB9ZfKJ4Kbo5DPSALwm24U0hMx0ENByRe+Ln9vsx6PPMRSC58n9eZ3onu910FDzc7yud1X6xlsWrYskyqvguhmzIzQcVwm+I1oMqMEwZP1TLRC2+X2kRz64LYTYPIg7qwKHkgYzEcaDYLOzScLNWnqko5oSPmNt54MH1sLKJ09ojTmBq5VOOk1o+aHVbN7jwJUCnHjrZY1FQa11z0Ky3jbycglXpym1DDd99m95P49Q9WJV2pYvDzkKaRhZHZTWtvSb6ozjRpFLXn0sXO3IGJPHrd3eGhQAAAAlwSFlzAAALEwAACxMBAJqcGAAABY5JREFUaN7tWudfU1cYPiGBJISQEAgE1EgAWQKyQYYsKYqACkVB1DJVlq2tu1qtts7WVeuse+/uvffee6/zf/SSec+469z09+uHvh/fc+7z5N5zzjueEwD+K2aM6em9UWazVZmPX/vt1+BiR/c0DO8PgXwrad9adTotKOjxpR1zIN3yExwpKtEjO9dCcbvNFsUOPzFuCpQ20+bv2OA/SoBy7asw5fCPbIZK7J6JyuATnSFQmZniPlWAv3waVG6hz8qF13YaIJPFjZeFr3NBVjswSQZ+1H2Q3ULul8RflAXVmP4FCfxxFqjS3hfFnx4KVVurCH5qFgyCOQTxqy/DoNgyoZRSHBx8aCqkE5TDYNlzkTT8GiQ2hxa8pgTy7gIkujRGkPjjk5Anergj/bZLLwv9w01FAEQgG3ABSbAPeSbb6nbGZOZIwq990QOwhe+ckkukXjTAPe3fue9pROELanwzjyD+ZiNGUIk+V8Y73M3el+q4sN3xrtm8zta6MmGn5+NlNwVwUlGENhQ/DPthy/nh+9Bhl+O8DktJLeZvk169k+/aiCDUIWPgIYwgkiGJY2H4FSTG4Z82moFgCQphWcEbm/EvEMDJgaEiYqfMZyBoxjDyAqdtMrH7xjEQEKE+2T/0JkEwVTn+fALkcd9QLnl+ZionuIMAMfi2diulplVOsIBEMXuHbidzN0OxqdtJwMz2xtFviJFOlmr5gTocZqMnkHxC4O9lq8dPEUAtbv82YnFiGBsKouCvcrv34u561o7lHJ6gfnC7f8IJHmPuifCQUzDmtOK1dD579zgVgyrRjvUy+At8wN7VXYSUsN9PlPnsBNoJlK99BCfYrqLvfZkS7/DvRis5ZNt+DOsU57uOE6xRQXAvhnWC89mCEEn99gaGdZLzxeIE61UQ4EknlvNV4QQuFbqMCcO6yTkP4QRd7AQVtFbhUdyZY2QmwOs3uJRzHiSC7AZmggZa3l9MEJQxEzyFQ41pJOEEwSxW/JcexqHcaT+f6NdXMBIQyzlBIOefYCRoxIH63O53CIKkeUz4ZGW0w+2/QhY0Nhb8xC6B7TKXJDAsZiCYSeKke7LEg+TIqiLF+OtIFIunlQQrKc3dtOkKlednKCDeyg7Yae1jSZNWAf5dZ2gYn3tHB+kSY+NncuEz1lAR9N2+CbN8riVR4M+btwI19q4KafRw+wwBETHBPyfZ56rljnbKFTsvbeze6ggTbDkH32r7uF1Ybgh0MRF5Pt/QRTByA8TUYgG864kviB8+e6hOXGPI56mcgbT51ygX1YdT/5aulZZKiRiZ/J8TKJg010H0VXANU2vCKR9IQr/SI3cLmbxTzKmrl2ouIJOdtBWYJE6ABv0oXs13eQSM3tLxBSZNKnWNa0VfYACd7ERrjWpwjidQHaBvomNiBFtwwYsXkM5yX/x8ZcWwv/8/SicQW+acDHy2HX2FwTOnwcCxPoslqbh8kdwagi5U+GxPYPTwH+51cdkB6LaCX44KSCMiq/w9RRXs5snWv7sTSOkol6TKgfPnCDrB18IrnE6b3xvQXEJ8SpyxyKqgjvNbKf2BJ3nlgF1O5y2EXyxQG6bxhbXKQqvkJaEAfla1oCKD9A95exYeD+vv/9G8L1ngMoOOf1bkyiuDru/vErhtoE42FIq99VzqDUU5ffIG6gaSuMaJXyXv1IzZAO3390qtXAolhMUKvC45c87rMu65yFq1QeC2lpiYfUlOjTCvXua5ITqLIbk60DKsznfSp2FiomZHovw6qgN5dDV9Vguqwx5UUgmmnSzhPbtQOpqa6nUKa9nqTL0UQXoA/3kWqTj3S19bvU2CoO+qFjBZ/HrPPx6aBJRS972VprJQhfYQ3tbOXdgLdVSbDHD3arV/rQAZIpt7JNII/je19g+y3fh5wGjAYgAAAABJRU5ErkJggg==',
    'user': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAABpFBMVEVMaXFdXV1mZmZZWVlYWFhZWVlVVVVZWVlZWVlZWVlZWVlZWVlcXFxZWVlZWVlYWFhYWFhYWFhZWVlYWFhZWVlYWFhZWVlYWFhYWFhZWVlVVVVZWVlZWVk/Pz9XV1dZWVlYWFhZWVlYWFhYWFhcXFxYWFhYWFhVVVVISEhYWFhZWVkAAABYWFhbW1tZWVlZWVlYWFhZWVlYWFhaWlpaWlpZWVlYWFhZWVlZWVlYWFhZWVlYWFhbW1tXV1dYWFhZWVlYWFhYWFhYWFhXV1dXV1dYWFhYWFhZWVlYWFhZWVlYWFhZWVlYWFhZWVlYWFhYWFhYWFhZWVlZWVlVVVVcXFxZWVlYWFhZWVlYWFhYWFhXV1dXV1dZWVlYWFhYWFhZWVlaWlpYWFhYWFhZWVliYmJYWFhYWFhYWFhZWVlVVVV/f39aWlpZWVlYWFhYWFhZWVlYWFhcXFxYWFhZWVlZWVlYWFhaWlpYWFhVVVVaWlpYWFhZWVlZWVlXV1dZWVlZWVlYWFhZWVlXV1daWlpZWVlZWVlYWFhZWVlZWVlYWFhbW1tZWVkbvEK5AAAAi3RSTlMAEwWP9O0V/vzzjPgk+SXy8ev7+m9Zy23b9RiRuQRXU9jOQmsW15kbB8f2ARcZidbKt41gIuXA0LHptu4nQ04//XzjPSMu92eFtOwUUYBWN3/NrgYLdZhQbpBaIJ1zVGRVaM94DZabydkSAkS/ld3oviGfi2ykd4oMH7PiyHQo3LK9Jmbho8FqjogcUD0cNwAAAAlwSFlzAAALEwAACxMBAJqcGAAAAq1JREFUaN7tmGdTIkEQhpe8CyJIFEFEEUSUoHemM+dwnvGMl3POOXo59Z8+q+6LA8syoafKq+L5ulvvUzW73TM9ikJL03401jjs9w83xqLBJgUZ024ICIphE2L8n70IlGG1dCDFmwMa6BI5k8HIHxmAilz9Ip7fooEBnmXR/BM+MES9IpY/D1UJC62PWl2g7vDn13mBAu0jb/7mAFDRauYUXANKPvDld1hpBVqOS7AH1PRy9bcIvcDK0/l2gYEkhyDEIrjFnp9zsAgc3cyCfWAiziyIsgkszIIYm+AUs+Ahm6CRWWBjEzQwC9xsAtfxE0hfogPZH1n6byq90IJsghVmQbfsZqcUWQSdHPtBmEXwk2fLtEreMhULvcAp+dji4ZyoXtIKLvCONq10+etmyYffEf7j9TLN8X1MZED4Wr2GL0keoeZFh7Qdj1G+d0wRpt7gX1qvwxiUmwOeSoP4JtKsn/ulU9TWXswbEdPz98T+4OhM3sG+cGmPW3oSNrfbluixxNuVGjWOAfbLzsXUUvbwn3+Ur6/QRPJDh4+zS6m3zy7aGdMDoaN15ejfmi19ZXbrB/FKMUzvOLntKu8MfRsLwdUZeyZjn1kNLmz06YwIN85Rxb92+oATf/5s9fy2KRDg3XiV+K41FYRQJ7sMW/95EOa7wSHG/AkQGGyuuD5PAYUXdysI1gCJSf38cRVLMDGql399CtD4plcPTkDkt05/8GEK/G/KBGlA5UHZmcGFK5g7XSIIADJJkUtMjsnZ7sAWTJBrtALotMgrgn+8IgSL+IJBQnAbX5AiBA34giFCMIcvyBICkEBNUBP8bwINP99LCBL4ggIhuIkv2CYE9/AFbeSe+Rk7v79k03+C3LBtZVdV0zbU/Gmdqfc+Xv5j/au20XQBoR60Qvro9/0LAwZf9wGaJzQAAAAASUVORK5CYII=',
    'linkedin': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAABblBMVEVMaXFZWVlYWFhVVVVYWFhYWFhZWVlYWFhZWVlWVlYAAABmZmZZWVlZWVlZWVlYWFhZWVlYWFhYWFhYWFhZWVlISEg/Pz9ZWVlYWFhfX19ZWVlYWFhZWVlYWFhYWFhYWFhaWlpWVlZZWVlZWVlVVVVVVVVfX19YWFhVVVVYWFhZWVlYWFhZWVlZWVlYWFhYWFh/f39ZWVlYWFhbW1tXV1dYWFhaWlpYWFhZWVlYWFhZWVlYWFhZWVlZWVlaWlpYWFhYWFhYWFhVVVVZWVlYWFhXV1diYmJZWVlVVVVZWVlZWVlZWVlaWlpYWFhYWFhZWVlYWFhZWVlYWFhVVVVaWlpYWFhZWVlcXFxYWFhZWVlYWFhYWFhZWVlYWFhYWFhYWFhZWVlYWFhaWlpZWVlYWFhcXFxYWFhZWVlXV1dZWVlYWFhaWlpZWVlYWFhZWVlYWFhaWlpbW1tYWFhZWVlXV1dYWFhZWVlYWFhXV1dZWVku2URHAAAAeXRSTlMAmvQG7vFY75c1AQXZ+f79+MNrx/YHBMuwEF4a88lI8mMszmwMGwjMFfpqxm/WeaEC/LgZIG4ic35UotppSlWHymgJts9DDd8YxaYlEYHEkaq6LhJE94YhpLTsihRRmTfIf1Io5gtLvSZbk2bCmKMXTDjB8DqnsZUdtFrL+QAAAAlwSFlzAAALEwAACxMBAJqcGAAAAdhJREFUaN7t2tdPAkEQBnCkqSgdVBAQ7CLYe++9995773r/vSGAzCA8YL59MOw8fjPZX0K44/YWlUqWLFg5271qBVRqu02TvH7VmAKtmla+vqZcAVdxAQNsCryGGVCPBxYY4MEDOgYoAkoCWQ7oBg9d1VfBsChguTaaFnWZhAAtRT95hQhAu0oaJQKAY9pYEQAEWKcXDpj5jbwEDhgNrDOD/4hqWecOD+hpo3IdD9jpZ9Qp4kLzJ/IOrQjAchKP28Ji7kWm+6pIWG3bEXa7Nh815oyq5Q/OfwDU2rrMAZOWlCUWlpGsObr2lvs08pA+Ehi4mcwIyKWNvFiYT7LCyFBDJUmafENgoCXp8V9lCNUhgc+m31sknwUHNKbchLlhwIcz9TZvCQV8pdlHbhpBQNp6Fw30iwbOjVBg7vFXZIUBhutLh6I4rPs83kABE9vx4QaW+1HAQ2J6nOZrIOCNTL/QxiwI6CHTzXuk4QIBZ3T8gr4TwgCl7HoK0S+XGQI8M2CejjsgwBMDdum4DgJ0M2ARD+gZcCAayJGABCQgAQlIIEsB+TJEAn8DyvDrv4o+arQyIIgHbvlxbw16fS8/JVCVgg+Ui6eSN+maCjvsyN0z3Vcg/8QgC1ff/RVqUE/ocDsAAAAASUVORK5CYII=',
    'link': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACcFBMVEVMaXFYWFhMTExVVVVZWVlVVVV/f39ZWVlVVVVZWVkAAABYWFhiYmJYWFhWVlZYWFhZWVk/Pz9YWFhYWFhZWVlZWVlZWVlbW1tZWVlYWFhZWVlZWVlYWFhaWlpZWVlYWFhYWFhYWFhZWVlYWFhYWFhYWFhZWVlYWFhaWlpfX19mZmZZWVlZWVlYWFhXV1dYWFhaWlpZWVlaWlpISEhYWFhYWFhZWVlZWVlYWFhZWVlbW1tYWFhYWFhYWFhbW1tYWFhZWVlYWFhZWVlYWFhVVVVZWVlYWFhXV1dZWVlYWFhYWFhaWlpdXV1YWFhcXFxYWFhVVVVcXFxZWVlWVlZYWFhaWlpYWFhZWVlZWVlZWVlZWVlZWVlaWlpYWFhYWFhXV1dVVVVXV1dYWFhaWlpZWVlYWFhZWVlZWVlZWVlYWFhXV1dZWVlZWVlbW1tYWFhZWVlYWFhVVVVZWVlZWVlaWlpZWVlYWFhZWVlYWFhZWVlYWFhYWFhVVVVYWFhYWFhYWFhZWVlaWlpZWVlYWFhaWlpYWFhZWVlYWFhZWVlYWFhXV1dZWVlaWlpZWVlfX19ZWVlZWVlZWVlYWFhaWlpcXFxZWVlZWVlYWFhaWlpZWVlYWFhYWFhXV1dZWVlVVVVXV1dYWFhYWFhYWFhYWFhZWVlZWVlZWVlYWFhYWFhYWFhYWFhWVlZZWVlYWFhaWlpZWVlaWlpZWVlZWVlYWFhZWVlYWFhYWFhZWVlYWFhZWVlbW1tYWFhcXFxZWVlYWFhZWVlYWFhYWFhZWVlZWVlYWFhVVVVZWVlYWFhYWFhXV1dYWFhZWVlZWVlZWVlZWVlZWVmH6k3iAAAAz3RSTlMA8QoJaQMC/gzCAe8N6y+s1ATuxPuoag66vvxbmE+AF61Zhi7V2uXAOxAFkfjyI2sf5zMHZZPFU4ixHFSH1xmwjNu3jR75bUN4NP0wE5UhVgYLlDWyZvrwjlDWri3JhFobIPQinXN+omSCXZK2J059lhimqxHzqtmvWJBxEvd2XGx3tOw+0nuz4rh0qVKPCPY84XlMJOTRzFUoyitXiQ9Ju5zpaJdeP5sa44EyckhjOUSl6J/tpIXTirkqURZhN8ieS3q9MRWjwcY6p/XQy4soUNhhAAAACXBIWXMAAAsTAAALEwEAmpwYAAAFaElEQVRo3rVa90MTSRSeEBNCIAQwEE5KgCAIAgeiKCDFU7AACoogIti7nmc5u556p+fZe9frvffee51/6TazG8m8mZ3dyW7eT8mb733zZubtmzezi5AFWVX8X1tGnQNXf5TR1r5yFbJXVpTMwrTMmhm2jd3dfAfz5OYily38BROxnkwcsk6/9jEski9rLPJPT8ViqX7KCr3rCWws+Zlx89fvxmbk2aQ4+T3l2JykxNdD0iPYrJyIZ5Y8Kdi85CfU/4jkJdR/RRw1CfVfkdwE82P8uPX5ycresiLkDoXzsnmP9z6XRf8/KUx7iOhd+hILmG7N/+EKClSRzSBGrPAnF3oBzLshGYLuWuAfx0FegahRe/mR1w9gLSb4e18wzY9Q8GkAbDXk927k8T+qB58HkPcMO3hLwn9F+gZp6H0j/mcmyPivCFiFIgP+4JtS/iuyiQbfMOhgtqT/CC2h0WVi/q8GJP1HaDXI2WL0Lml+9IDGB4TghQHJ+VGkRmYEs6X9V8KOtrgurLHKYAfHtJbtP1e16Gy5ObTFP6IO5kL+aVpDOkkI65w8o+doE7+og2WA/3et1En3qf+/72ZtKqpomxmiDm7S2AVdND/Gy5sYm1Hg1C0BfxrIEtmqOuwbU5UeBjY714CoCAk6mA+ceY1oQ6djdePpdXCfADYZohkqBgWEugJ/0doXKZOZMCwKRR3M4+XFhQsARXFsiDJ78m029ofOXNJ+ltDYZUQ5jikQWwX8Dez+m6uoe9TfW2nwUW7oYrwxajuV4ccFOvv7FPJnGg1+h6eMBMr7qu2cDqap1K1TP7zMW7EviPIcu0FsU43z2Jbn9eqTHeTvKRr9oQpazA4hnTR8wDT069Y/58n/V8HepBVynUwOf53onX8DdWqtbn2lRuSfwGCKBrwK64CseqLvF1e+sfXbSaLpTubnUnSIfw44Risn0/zO2PPpxUqiW05bdKyOgveCDno4qaWIPhu49lOtbxNlPuZmO+Vx7uBlnMojMaryC/QAwP1ADlEOwak4FIWDUEreQ7Sfx5zDPTT/AcDUpl4bwGNR4Cq/+MHfEe2wLv9aBzBYo2bhdXAIgU4tLdfRevV65aAeP2Lvf+YSfRebXRarFmdp7RWi3KHHX8A+5MXcHTwi53gFzWexmZbhd7H3Y1XaZjuJHcJW0nCGF/Pz+fyomc0um6JtnzJt7bzNSM0tqCdyvGf40R1B/fYGXH68lOgn08rNGvzS+h/d7P0nw58T07oSNE5QM1gjrV0v2npLYAclVPMPdONv6kMLBtYs6gDe3zbSg3Q2xDZe/pVbVKaL7p/hneQ3AHC4dKzRF+YexAaSzFc/nEqm6WFW9WmeLgHR+4dohtpp7JEgC+keofkrG7BEhdsGLjx5mAskKrOiW9orMCwmiTrI4IU5I3ktF3O3a787If9pr6iDOgln1IzDpA/hDKEnafAcI/6psDjFA2IbUCIEpf3HHyOZEdRK81ddE5vU8Ysf0/wgsxhG0Xuy/Jc9Bh2AvLhXtL4c/oBh2IEnObVPyn98yjCsYcI/KuM/3uU27KAJmPiCEv7n9pq4VYT7gd9rmj/FY4KfPXr+xLviiJ8fhZk9eQMcg3uzBX6E/mVMs+lb7z2N2Ao/WsR5LXB87Iqg73gWj398mvl3e7w3n46zW37Zmdna9e3uQWzNf0W+ln95IuF/RLYlmB9dq5bjl5ofvXW20X/eec9e/yORVJRQ/yPvscy+Ryz3oPik3m+Kf78z7nfdmWbW4aAroV8DOA5Y/N6gRvw9Q3+tDV9k7NOlL33Xpm9KRrj0DQVuZJfcHW2B97eFt5G90nrv/vDJskGcev2Gf8atkBWq/wGmwCxnJZp1EQAAAABJRU5ErkJggg==',
    'portfolio': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACDVBMVEVMaXFZWVlYWFhZWVlYWFhbW1tZWVlZWVlVVVVWVlZYWFhcXFxYWFhZWVlYWFhaWlpZWVlYWFhYWFhMTExZWVlYWFhZWVlYWFhZWVlbW1t/f39YWFhaWlpYWFhZWVlYWFhZWVlXV1dcXFxZWVlXV1diYmJZWVlYWFhYWFhWVlZYWFhZWVlZWVlYWFgAAABZWVlYWFhZWVlYWFhYWFhZWVlaWlpZWVlZWVlZWVlYWFhfX19ZWVlYWFhYWFg/Pz9VVVVYWFhYWFhZWVlfX19ZWVlYWFhXV1dYWFhmZmZaWlpVVVVZWVlcXFxZWVlYWFhYWFhcXFxZWVlYWFhVVVVZWVlYWFhZWVlYWFhbW1tZWVlYWFhXV1dYWFhYWFhZWVlYWFhZWVlXV1dZWVlXV1dVVVVYWFhaWlpYWFhYWFhYWFhVVVVYWFhYWFhZWVlZWVlaWlpYWFhVVVVYWFhZWVlYWFhYWFhVVVVYWFhZWVlVVVVdXV1ZWVlZWVlZWVlZWVlYWFhYWFhZWVlaWlpZWVlYWFhYWFhZWVlXV1dZWVlXV1dZWVlYWFhZWVlYWFhZWVlYWFhZWVlYWFhZWVlZWVlYWFhYWFhbW1tZWVlYWFhZWVlZWVlYWFhZWVlYWFhYWFhYWFhZWVlaWlpYWFhaWlpYWFhZWVlYWFhZWVlZWVlYWFhYWFhYWFhZWVlYWFhZWVl0buXFAAAArnRSTlMA+S7+4zg2+AwvkAv3kpMwke70CuT9OTeuDgJLLcq06+pDJN9dDcViyTLy7RSEAeEa/K2q9mZqP1iHCKPvKwQewIKxEF7eV8QFYxLoIc5ChRbNxwbcMTz6GW/dJsHx9YGAKVtAFRdP2G7SD9VzfstV6QnPfZY0GHyMAxPz1HJTr3G/d9PsippG53RHfyjmj7Ko4PBQrLAcoGXC0KR7VlmbqxFIPrjIX47WaGuNZ5k9Mdc9AAAACXBIWXMAAAsTAAALEwEAmpwYAAADjklEQVRo3rVa9V8bMRRPe0BLi7Wl0AIthbLh7j5kMIZvuAy3ubu7u7v78jeug0IZ5ZK8pPn+eJ/vfb+53OXdy3tBiAkzSfMR9XvdJivGVtOQw2Yo+fQBBQrJPyrd2B/K93dv7eLqCWMLWB06Z2qWiLolNQ3TkO9M4pXfWleFmXAwZQuH/M7ODMyMoOI4oHxOYysGob0P9modGAqlMo99+IM6zIGH+xj1j27CfFAM21n003MxN7qT6fqNOiyAzBia/hmFcHuYIcRojI0iDcF8k6xfR/zatUuk0CACSVNG0h8j6Ydpl2mhpGewtqjrnyfOr8FHDCdGp3g1/RSFaBDiYwYTieYjKt9nEfkL0fuoejLTte7XOhCNmQ02UKi7LP76cdT1yzxFHhz2N7hBXURRPnIENWr4vWgtfQHrQpfJkRoquXV6zQSlsfxWvA6R0Qxk2/8GqUyRRhcerNfHRGiYyB2r9Rv2Q2MOnd+9+i+aDY85dH7KqvxhlCPmUPm5FSoPwBpzqPySlfyqintBEfmnly+38MYcGr/Ue9mheoMdaLCGX+hNggRiDpnfv5QqjQvEHAp/fvHqgkjMIfM/L+4vxGIOka8zei4XC8YcIn/WY1CJJeKPx8At02AOoWmZ+thqR1NSDTyL+adcg99sAZMfv1CBXAMb2ibXYCPaLNegByXKNUhEO+QamJFVrkE+6hfKiBimKEMkI6LjJXIJZEQMeIXmBDIiBrxBIwK/ezpG7aiQb5PEiNcIlcs0cHv2al8lTpH5XzYQIpAR0fQXN2oVGv6MiIzc3iWtLwIZEWn440bvYL+JZEQqVaNL9ycbVmb7o5CW0ry7r4tcSasQCdgFCQxFuifc8qZHTFXGJl79oAnGEngG5/irWQuxV/kMmpgrvdUKj74TUKq+wGOQDjC4xqHvAFXbm+EGx0AGSXCDDlhDoh5s0AYzGAZnEdDu0y2owQDQoKYdaBAL7Tr1AiepGNzXGoRuXaDIKQQZ6GbADgdgu6lD8OZfF2g7dYejv3jbBHHgaZKWngIYZPN0YIMBz3COq8ervQ7YAnOhrUeyAcobYTS4x9sJryln+0df4e/lx2eyGEwJnBaoddL1XSeEjjvMPqAZHBc8UGHpJOfwl7OQKGrvmglp+9lAHDo5uUe12HMxQOdaLGUF63U4i56jwEFf9t61phj6eBgFGBOT5ba0aJMmzDT04lnJU8itfwGRUUBmQPW0NAAAAABJRU5ErkJggg==',
}
# 透明 1x1 图片：盖在联系方式胶囊上承载超链接，整颗胶囊可点，文字不变色、不加下划线
CLEAR = 'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII='

# ---------- 简洁单栏版式（layout: "simple"）常量 ----------
S_X0, S_X1 = 51, W - 51
S_CW, S_TOP, S_BOTTOM = S_X1 - S_X0, 40, H - 36
LEVELS_S = [
    dict(fs=10.5, ts=11.5, lh=18, g=1.3),
    dict(fs=10, ts=11, lh=17, g=1.2),
    dict(fs=10, ts=11, lh=16, g=1.1),
    dict(fs=9.5, ts=10.5, lh=15.5, g=1.05),
] + LEVELS
S_BASE = 4   # LEVELS_S[S_BASE] 就是原始版式（level 0），更宽松的档位 level 为负数
INFO_RE = re.compile(r'个人信息|基本信息|基本资料|个人资料|基本情况|personal', re.I)
SPLIT_RE = re.compile(r'\n|\s+[·•｜|]\s+|[｜|]')


def split_parts(t):
    return [x.strip() for x in SPLIT_RE.split(str(t or '')) if x.strip()]


MOVABLE = re.compile(r'教育|学历|证书|资格|荣誉|奖|语言|兴趣|爱好|社团|技能|education|certif|licen|award|honou?r|language|interest|affiliation|skill|activit', re.I)


class Builder:
    def __init__(self, data, base_dir):
        self.D = data
        self.base = base_dir
        self.theme = str(data.get('theme') or '2563EB').replace('#', '').upper()
        self.card = tint(self.theme, 0.1)
        alltxt = json.dumps(data, ensure_ascii=False, separators=(',', ':'))
        if data.get('lang'):
            self.EN = data['lang'] == 'en'
        else:
            self.EN = len(re.findall('[一-鿿]', alltxt)) < len(alltxt) * 0.03
        self.BODY = LAT if self.EN else CJK
        self.BB = not self.EN   # 正文是否加粗（中文版式全加粗，英文只加粗标题和重点）
        self.photo = None
        self.S = None

    # ---- 宽度与断行 ----
    def cw(self, ch, size, arial, bold):
        c = ord(ch)
        if c >= 0x80:
            return size
        if not arial:
            return size * 0.5
        w = ARIAL[c - 32] if 32 <= c <= 126 else 556
        return size * w / 1000 * (1.08 if bold else 1)

    def textW(self, t, size, arial=None, bold=None):
        arial = self.EN if arial is None else arial
        bold = self.BB if bold is None else bold
        return sum(self.cw(ch, size, arial, bold) for ch in strip(t))

    def wrap(self, t, size, width, arial=None, bold=None):
        """断行：英文按单词、中文按字，模拟 PowerPoint 的换行"""
        arial = self.EN if arial is None else arial
        bold = self.BB if bold is None else bold
        toks = re.findall(r'[\x21-\x7E]+|\s+|\S', strip(t))
        lim = width * (0.97 if arial else 1) + 0.5
        lines, cur, w = [], '', 0.0
        for tok in toks:
            tw = self.textW(tok, size, arial, bold)
            if tok.isspace():
                if cur:
                    cur += tok
                    w += tw
                continue
            if w + tw <= lim:
                cur += tok
                w += tw
                continue
            if cur.strip():
                lines.append(cur.rstrip())
                cur, w = '', 0.0
            if tw <= lim:
                cur, w = tok, tw
                continue
            for ch in tok:
                c = self.cw(ch, size, arial, bold)
                if w + c > lim and cur:
                    lines.append(cur)
                    cur, w = '', 0.0
                cur += ch
                w += c
        if cur.strip() or not lines:
            lines.append(cur.rstrip())
        return lines

    def nLines(self, t, size, width, arial=None, bold=None):
        return len(self.wrap(t, size, width, arial, bold))

    def tagLines(self, tags, size, width):
        """标签排行：每个标签不拆开，用「 · 」连接；单个标签超宽时才按普通文字断行"""
        SEP, lines, cur = ' · ', [], ''
        for t in [str(x).strip() for x in tags if str(x).strip()]:
            cand = cur + SEP + t if cur else t
            if self.textW(cand, size, True, self.BB) <= width:
                cur = cand
                continue
            if cur:
                lines.append(cur)
            if self.textW(t, size, True, self.BB) <= width:
                cur = t
            else:
                w = self.wrap(t, size, width, True, self.BB)
                lines.extend(w[:-1])
                cur = w[-1]
        if cur:
            lines.append(cur)
        return lines or ['']

    # ---- 富文本 ----
    def rich(self, text, color, auto):
        """**加粗** → 黑色加粗；数字成果自动加粗。返回 run 列表"""
        runs = []
        for seg in re.split(r'(\*\*.+?\*\*)', str(text or '')):
            if not seg:
                continue
            if seg.startswith('**') and seg.endswith('**') and len(seg) >= 4:
                runs.append(dict(text=seg[2:-2], color=BLACK, bold=True))
                continue
            if not auto:
                runs.append(dict(text=seg, color=color, bold=self.BB))
                continue
            last = 0
            for m in NUM_RE.finditer(seg):
                k = m.group(0).strip()
                if not m.group(0) or any(r.match(k) for r in SKIP_RE):
                    continue
                if m.start() > last:
                    runs.append(dict(text=seg[last:m.start()], color=color, bold=self.BB))
                runs.append(dict(text=m.group(0), color=BLACK, bold=True))
                last = m.end()
            if last < len(seg):
                runs.append(dict(text=seg[last:], color=color, bold=self.BB))
        return runs

    def mixRuns(self, t, color):
        """拆分中英文，英文用 Arial"""
        out, cur, lat = [], '', None
        for ch in t:
            is_lat = lat if ch == ' ' else ord(ch) < 0x80
            if lat is not None and is_lat != lat:
                out.append(dict(text=cur, color=color, font=LAT if lat else CJK))
                cur = ''
            cur += ch
            lat = is_lat
        if cur:
            out.append(dict(text=cur, color=color, font=LAT if lat else CJK))
        return out

    # ---- 底层绘制 ----
    @staticmethod
    def _font(run, face):
        rPr = run._r.get_or_add_rPr()
        for tag in ('a:latin', 'a:ea', 'a:cs'):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set('typeface', face)

    def add_text(self, x, y, w, h, paras, size, font, color=BLACK, bold=True, align='left', valign='top',
                 lh=None, name=None, bullet=False, wrap=True):
        """paras: 段落列表，每段是 run 字典列表（text/color/bold/font 可覆盖默认值）"""
        tb = self.S.shapes.add_textbox(E(x), E(y), E(w), E(h))
        if name:
            tb.name = str(name)
        tf = tb.text_frame
        tf.word_wrap = bool(wrap)
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE if valign == 'middle' else MSO_ANCHOR.TOP
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.RIGHT if align == 'right' else PP_ALIGN.LEFT
            if lh:   # 行距写成 spcPts，四舍五入到 0.01pt（与 pptxgenjs 一致）
                pPr = p._p.get_or_add_pPr()
                ln = etree.SubElement(pPr, qn('a:lnSpc'))
                etree.SubElement(ln, qn('a:spcPts')).set('val', str(int(math.floor(lh * 100 + 0.5))))
            if bullet:
                pPr = p._p.get_or_add_pPr()
                pPr.set('marL', '171450')
                pPr.set('indent', '-171450')
                etree.SubElement(pPr, qn('a:buSzPct')).set('val', '100000')
                etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
                etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
            for r in runs:
                if not r.get('text'):
                    continue
                run = p.add_run()
                run.text = r['text']
                f = run.font
                f.size = E(size)
                f.bold = r.get('bold', bold)
                f.color.rgb = RGBColor.from_string(r.get('color', color))
                self._font(run, r.get('font', font))
        return tb

    def T(self, x, top, w, size, text, color=BLACK, font=None, bold=True, align='left', lh=None, lines=1,
          runs=None, name=None, boxY=None, boxH=None, valign='top', wrap=True):
        """单个文本框：top 为文字顶边位置，自动换算成文本框位置"""
        if not self.S:
            return
        l = lh or size * 1.25
        y = boxY if boxY is not None else top - size * 0.12 - (l - size) / 2
        h = boxH if boxH is not None else l * lines
        rs = runs if runs is not None else [dict(text=str(text or ''))]
        self.add_text(x, y, w, h, [rs], size, font or self.BODY, color, bold, align, valign, l,
                      name or strip(text)[:20], wrap=wrap)

    def RR(self, x, y, w, h, color, r, tr=None, name=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        if not self.S:
            return
        sp = self.S.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = min(0.5, r / min(w, h))
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(color)
        if tr:
            clr = sp.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            etree.SubElement(clr, qn('a:alpha')).set('val', str(int((100 - tr) * 1000)))
        sp.line.fill.background()
        # 去掉主题样式引用（否则 LibreOffice / WPS 会套用主题阴影），与 JS 版生成的形状一致
        style = sp._element.find(qn('p:style'))
        if style is not None:
            sp._element.remove(style)
        if name:
            sp.name = name
        # 去掉形状默认文字样式的影响
        return sp

    def bar(self, x, y):
        self.RR(x, y, 8.5, 2.85, self.theme, 1.4)

    def pic(self, b64_or_bytes, x, y, w, h, name=None, link=None):
        data = base64.b64decode(b64_or_bytes) if isinstance(b64_or_bytes, str) else b64_or_bytes
        p = self.S.shapes.add_picture(io.BytesIO(data), E(x), E(y), E(w), E(h))
        if name:
            p.name = name
        if link:
            p.click_action.hyperlink.address = link
        return p

    # ---- 链接安全：只放行 http/https/mailto/tel ----
    @staticmethod
    def safe_url(u):
        u = str(u or '').strip()
        if not re.match(r'^(https?://|mailto:|tel:)', u, re.I) or re.search(r'[\s<>"]', u):
            return None
        return u   # python-pptx 写 rels 时会自动做 XML 转义

    def contact_link(self, c):
        if c.get('link'):
            return self.safe_url(c['link'])
        t = str(c.get('text') or '').strip()
        typ = c.get('type')
        if typ == 'email' and re.match(r'^[^\s@]+@[^\s@]+\.[^\s@]+$', t):
            return self.safe_url('mailto:' + t)
        if typ in ('web', 'github', 'linkedin', 'link', 'portfolio') and re.match(r'^(https?://)?[\w-]+(\.[\w-]+)+(/\S*)?$', t):
            return self.safe_url(t if re.match(r'^https?:', t) else 'https://' + t)
        return None

    # ---- 头像：裁成正方形，用 PowerPoint 自带的椭圆形状显示（「更改图片」后仍是圆形） ----
    def load_photo(self):
        p = self.D.get('photo')
        if not p:
            return
        path = os.path.join(self.base, p)
        if not os.path.exists(path):
            print('照片不存在：' + path, file=sys.stderr)
            return
        from PIL import Image, ImageOps
        im = ImageOps.exif_transpose(Image.open(path)).convert('RGB')
        s = min(im.size)
        left = (im.width - s) // 2
        top = (im.height - s) // 4      # 偏上取景，尽量保住头顶
        im = im.crop((left, top, left + s, top + s)).resize((600, 600), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=90)
        self.photo = buf.getvalue()

    # ---------- 排版（S 为空时只计算高度） ----------
    def layout(self, cfg):
        D, S = self.D, self.S
        fs, ts, lh, g = cfg['fs'], cfg['ts'], cfg['lh'], cfg['g']
        T, RR, bar = self.T, self.RR, self.bar

        # ===== 头部 =====
        summary = D.get('summary')
        sumLines = self.nLines(summary, 8.1, 521) if summary else 0
        headBottom = 128.9 + (sumLines - 1) * 12 + 22.9 if summary else 112
        RR(20.2, 20.3, 555, headBottom - 20.3, self.theme, 10, name='顶部背景')
        T(HEAD_X, 33.4, 300, 17.1, D.get('name'), color=WHITE, font=LAT if self.EN else YH, name='姓名')
        if D.get('title'):
            label = D.get('titleLabel')
            if label is None:
                label = '' if self.EN else '求职意向：'
            T(HEAD_X, 58.0, 450, 9, label + D['title'], color=WHITE, name='求职意向')
        px = 30
        for i, c in enumerate(D.get('contacts') or []):
            text = str(c.get('text') or '')
            lat = not re.search(r'[^\x00-\x7F]', text)
            icon = ICONS.get(c.get('type'))
            tx = 18.1 if icon else 8
            w = tx + self.textW(text, 8, lat, True) * 1.06 + 10   # 留余量，避免字体差异把文字挤成两行
            RR(px, 81.5, w, 17, WHITE, 8.5, tr=70, name='联系方式底框%d' % (i + 1))
            if S and icon:
                self.pic(icon, px + 6.3, 81.5 + (17 - 8.5) / 2, 8.5, 8.5, name='图标%d' % (i + 1))
            T(px + tx, 0, w - tx, 8 if lat else 8.1, text, color=WHITE, font=LAT if lat else CJK,
              boxY=81.5, boxH=17, valign='middle', wrap=False, name='联系方式%d' % (i + 1))
            url = self.contact_link(c)
            if S and url:
                self.pic(CLEAR, px, 81.5, w, 17, name='链接%d' % (i + 1), link=url)
            px += w + 4
        if summary:
            RR(HEAD_X, 112.8, 568.5 - HEAD_X, 1, WHITE, 0, tr=70, name='分隔线', shape=MSO_SHAPE.RECTANGLE)
            T(HEAD_X, 128.9, 526, 8.1, summary, color=WHITE, lh=12, lines=sumLines, bold=self.BB, name='个人简介')
        if S and self.photo:
            p = self.pic(self.photo, 493.2, 35.0, 62.4, 62.4, name='照片')
            try:
                p.auto_shape_type = MSO_SHAPE.OVAL
            except AttributeError:   # 老版本 python-pptx：直接改 XML
                p._element.spPr.find(qn('a:prstGeom')).set('prst', 'ellipse')

        # ===== 左栏 =====
        top0 = headBottom + 18
        y = top0
        auto = D.get('autoBold', True) is not False
        for sec in D.get('left') or []:
            bar(L_X, y + 2.9)
            T(36, y, 200, ts, sec.get('title'), name=sec.get('title'))
            y += 21 * max(g, 0.8)
            for ei, e in enumerate(sec.get('entries') or []):
                # 没有标题、时间和副标题的条目（纯列表/段落）：直接接在板块标题下面，不留空行
                if not e.get('name') and not e.get('date') and not e.get('sub') and not e.get('subRight'):
                    y -= 21 * max(g, 0.8) if ei == 0 else 11.7 * g + 18.3 * max(g, 0.75) - 4
                dw = self.textW(e['date'], fs, self.EN, True) + 4 if e.get('date') else 0
                nameW = DATE_R - L_X - dw - 6
                nameN = self.nLines(e['name'], ts, nameW, self.EN, True) if e.get('name') else 1
                if e.get('name'):
                    T(L_X, y, nameW, ts, e['name'], name=(e['name'] or '条目')[:20], lh=ts * 1.25, lines=nameN)
                if e.get('date'):
                    T(DATE_R - 120, y - 1.5, 120, fs, e['date'], align='right', name='时间')
                if e.get('subRight'):
                    T(DATE_R - 120, y + 9.3, 120, fs, e['subRight'], align='right', name='右侧副信息')
                ny = y + (nameN - 1) * ts * 1.25   # 标题多行时整体下移
                if e.get('sub'):
                    subW = DATE_R - L_X - self.textW(e['subRight'], fs, self.EN, True) - 10 if e.get('subRight') else DATE_R - L_X
                    subN = self.nLines(e['sub'], fs, subW)
                    sy = ny + 12.5 * max(g, 0.85)
                    T(L_X, sy, subW, fs, e['sub'], color=GREY, bold=self.BB, lh=fs * 1.3, lines=subN, name='副标题')
                    cur = sy + (subN - 1) * fs * 1.3 + 16.7 * max(g, 0.75)
                else:
                    cur = ny + 18.3 * max(g, 0.75)
                any_ = False
                # details：成绩/主修课程/技术栈等明细行，冒号前的标签用黑色
                for dl in e.get('details') or []:
                    n = self.nLines(dl, fs, PARA_W)
                    mm = re.match(r'^([^：:]{1,12}[：:]\s?)([\s\S]*)$', str(dl))
                    runs = ([dict(text=mm.group(1), color=BLACK, bold=True)] + self.rich(mm.group(2), GREY, auto)) if mm else self.rich(dl, GREY, auto)
                    T(PARA_X, cur, PARA_W, fs, dl, runs=runs, color=GREY, bold=self.BB, lh=lh, lines=n, name='明细')
                    cur += n * lh
                    any_ = True
                bullets = e.get('bullets') or []
                if bullets:
                    total = sum(self.nLines(b, fs, BUL_W - BUL_IND) for b in bullets)
                    if S:
                        self.add_text(BUL_X, cur - fs * 0.12 - (lh - fs) / 2, BUL_W, lh * total,
                                      [self.rich(b, GREY, auto) for b in bullets], fs, self.BODY, GREY, self.BB,
                                      lh=lh, name='列表_' + str(e.get('name') or '')[:16], bullet=True)
                    cur += total * lh
                    any_ = True
                if e.get('text'):
                    n = self.nLines(e['text'], fs, PARA_W)
                    T(PARA_X, cur, PARA_W, fs, e['text'], runs=self.rich(e['text'], GREY, auto), color=GREY,
                      bold=self.BB, lh=lh, lines=n, name='描述')
                    cur += n * lh
                    any_ = True
                if not any_:
                    cur -= lh * 0.3
                y = cur + 11.7 * g
            y += 1.6 * g
        leftBottom = y - 13.3 * g

        # ===== 右栏卡片 =====
        ry = top0 - 1.2
        for card in D.get('right') or []:
            c0 = ry
            it = c0 + 28.3 * max(g, 0.9)
            lastTop = it
            # tags：标签式卡片，没有小标题和三个点
            items = [dict(name='', tags=card['tags'])] if card.get('tags') is not None else (card.get('items') or [])
            ops = []
            for item in items:
                if item.get('tags') is not None:
                    lines = self.tagLines(item['tags'], 8, R_TXT_W)
                else:
                    lines = [l for raw in str(item.get('content') or '').split('\n') for l in self.wrap(raw, 8, R_TXT_W, True, self.BB)]
                dW = self.textW(item['date'], 7.5, True, True) + 6 if item.get('date') else 0
                nameN = self.nLines(item['name'], ts, 150 - dW, self.EN, True) if item.get('name') else 0
                shift = max(nameN - 1, 0) * ts * 1.25
                t0 = it + 25.3 + shift if item.get('name') else it
                ops.append((it, t0, item, lines, nameN, shift, dW))
                lastTop = t0 + (len(lines) - 1) * 9.6
                it = lastTop + 18.3 * g
            bottom = lastTop + 15.3 if items else c0 + 30
            if S:
                RR(R_X, c0, R_W, bottom - c0, self.card, 6, name=str(card.get('title') or '') + '卡片')
                bar(410.8, c0 + 12.7)
                T(422.9, c0 + 10.6, 140, ts, card.get('title'), name=card.get('title'))
                for it_, t0, item, lines, nameN, shift, dW in ops:
                    if item.get('name'):
                        T(R_IN, it_, 150 - dW, ts, item['name'], name=item['name'], lh=ts * 1.25, lines=nameN)
                        if item.get('date'):
                            T(R_X + R_W - 8 - 70, it_ + 0.8, 70, 7.5, item['date'], color=GREY, align='right', font=LAT, name='时间')
                        for bx in (412.0, 423.0, 434.2):
                            bar(bx, it_ + 15.9 + shift)
                    self.add_text(R_IN, t0 - 8 * 0.12 - 0.8, R_TXT_W + 15, 9.6 * len(lines),
                                  [self.mixRuns(l, GREY) for l in lines], 8, CJK, GREY, self.BB, lh=9.6,
                                  name=str(item.get('name') or '') + '内容')
            ry = bottom + 12
        rightBottom = ry - 12
        return max(leftBottom, rightBottom)

    # ---------- 简洁单栏版式 ----------
    def simple_parts(self, merge):
        """个人信息并入页眉，其余右栏卡片排在左栏板块后面；merge=True 时把卡片合成「其他信息」"""
        D, EN = self.D, self.EN
        own = isinstance(D.get('info'), list)
        info = [str(x) for x in D['info'] if str(x).strip()] if own else []
        secs = [dict(title=s.get('title'), kind='entries', entries=s.get('entries') or []) for s in D.get('left') or []]
        for c in D.get('right') or []:
            if not own and INFO_RE.search(str(c.get('title') or '')):
                if c.get('tags') is not None:
                    info += [str(x) for x in c['tags'] if str(x).strip()]
                else:
                    for it in c.get('items') or []:
                        parts = split_parts(it.get('content'))
                        if it.get('name') and not INFO_RE.search(str(it['name'])) and parts:
                            parts[0] = str(it['name']) + (': ' if EN else '：') + parts[0]
                        info += parts
                continue
            if c.get('tags') is not None:
                secs.append(dict(title=c.get('title'), kind='tags', tags=c['tags'], card=True))
            else:
                secs.append(dict(title=c.get('title'), kind='rows', items=c.get('items') or [], card=True))
        if merge:
            cards = [x for x in secs if x.get('card')]
            if len(cards) >= 2:
                sep = '; ' if EN else '；'
                items = []
                for c in cards:
                    if c['kind'] == 'tags':
                        content = (', ' if EN else '、').join(str(x) for x in c['tags'] if str(x).strip())
                    else:
                        out = []
                        for it in c['items']:
                            if not (str(it.get('name') or '').strip() or str(it.get('content') or '').strip()):
                                continue
                            body = (', ' if EN else '，').join(split_parts(it.get('content')))
                            head = str(it.get('name') or '') + (((' (' + str(it['date']) + ')') if EN else ('（' + str(it['date']) + '）')) if it.get('date') else '')
                            out.append(head + (': ' if EN else '：') + body if head and body else head or body)
                        content = sep.join(out)
                    items.append(dict(name=c['title'], content=content))
                secs = [x for x in secs if not x.get('card')] + [dict(title='Additional Information' if EN else '其他信息', kind='rows', items=items)]
        return info, secs

    def layout_simple(self, cfg, merge):
        D, S, EN, BB = self.D, self.S, self.EN, self.BB
        fs, ts, lh, g = cfg['fs'], cfg['ts'], cfg['lh'], cfg['g']
        k = fs / 9                      # 条目内部间距随字号缩放
        gi = min(max(g, 0.75), 1)       # 条目内部只收紧、不放大
        auto = D.get('autoBold', True) is not False
        LINE = tint(self.theme, 0.35)
        T = self.T
        def RR(x, y, w, h, color, r, name=None):
            self.RR(x, y, w, h, color, r, name=name, shape=MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE)
        info, secs = self.simple_parts(merge)
        has_photo, PH = bool(self.photo), 72
        HW = S_CW - (PH + 12 if has_photo else 0)

        # ===== 页眉 =====
        y = S_TOP
        T(S_X0, y, HW, 22, D.get('name'), font=LAT if EN else YH, name='姓名')
        y += 22 * 1.25 + 6
        if D.get('title'):
            label = D.get('titleLabel')
            if label is None:
                label = '' if EN else '求职意向：'
            T(S_X0, y, HW, 10.5, label + str(D['title']), color='333333', bold=BB, name='求职意向')
            y += 10.5 * 1.25 + 7
        contacts = [c for c in D.get('contacts') or [] if str(c.get('text') or '').strip()]
        if contacts:
            cx, row = S_X0, y
            for i, c in enumerate(contacts):
                text = str(c['text'])
                lat = not re.search(r'[^\x00-\x7F]', text)
                icon = ICONS_DARK.get(c.get('type'))
                iw = 13 if icon else 0
                tw = self.textW(text, 9, lat, BB) * 1.06 + 2
                w = iw + tw
                if cx > S_X0 and cx + w > S_X0 + HW:
                    cx, row = S_X0, row + 16
                if S and icon:
                    self.pic(icon, cx, row + 0.6, 9, 9, name='图标%d' % (i + 1))
                T(cx + iw, 0, tw, 9, text, color='333333', bold=BB, font=LAT if lat else CJK,
                  boxY=row - 2, boxH=14, valign='middle', wrap=False, name='联系方式%d' % (i + 1))
                url = self.contact_link(c)
                if S and url:
                    self.pic(CLEAR, cx, row - 2, w, 14, name='链接%d' % (i + 1), link=url)
                cx += w + 16
            y = row + 16
        if info:
            t = '  |  '.join(info)
            n = self.nLines(t, 9, HW)
            T(S_X0, y, HW, 9, t, color=GREY, bold=BB, lh=14, lines=n, name='个人信息')
            y += n * 14 + 2
        if S and has_photo:
            p = self.pic(self.photo, S_X1 - PH, S_TOP - 4, PH, PH, name='照片')
            try:
                p.auto_shape_type = MSO_SHAPE.OVAL
            except AttributeError:
                p._element.spPr.find(qn('a:prstGeom')).set('prst', 'ellipse')
        if has_photo:
            y = max(y, S_TOP - 4 + PH + 4)
        if D.get('summary'):
            y += 6
            n = self.nLines(D['summary'], fs, S_CW)
            T(S_X0, y, S_CW, fs, D['summary'], runs=self.rich(D['summary'], GREY, auto), color=GREY, bold=BB, lh=lh, lines=n, name='个人简介')
            y += n * lh
        y += 10 * g

        # ===== 板块 =====
        gap = 6 * g
        for sec in secs:
            th = ts + 1
            RR(S_X0, y - th * 0.05, 3, th * 1.1, self.theme, 1.2)
            T(S_X0 + 8, y, S_CW - 8, th, sec.get('title'), name=sec.get('title'))
            RR(S_X0, y + th * 1.25 + 2, S_CW, 0.75, LINE, 0, name='分隔线')
            y += th * 1.25 + 2 + 7 * max(g, 0.8)
            if sec['kind'] == 'entries':
                ents = sec['entries']
                for ei, e in enumerate(ents):
                    bare = not e.get('name') and not e.get('date') and not e.get('sub') and not e.get('subRight')
                    dw = self.textW(e['date'], fs, EN, True) + 4 if e.get('date') else 0
                    nameW = S_CW - dw - 8
                    nameN = self.nLines(e['name'], ts, nameW, EN, True) if e.get('name') else 1
                    if bare:
                        cur = y
                    else:
                        if e.get('name'):
                            T(S_X0, y, nameW, ts, e['name'], name=str(e['name'])[:20], lh=ts * 1.25, lines=nameN)
                        if e.get('date'):
                            T(S_X1 - 140, y - 1.5 * k, 140, fs, e['date'], align='right', name='时间')
                        if e.get('subRight'):
                            T(S_X1 - 140, y + 9.3 * k, 140, fs, e['subRight'], align='right', name='右侧副信息')
                        ny = y + (nameN - 1) * ts * 1.25
                        if e.get('sub'):
                            subW = S_CW - self.textW(e['subRight'], fs, EN, True) - 10 if e.get('subRight') else S_CW
                            subN = self.nLines(e['sub'], fs, subW)
                            sy = ny + 12.5 * k * gi
                            T(S_X0, sy, subW, fs, e['sub'], color=GREY, bold=BB, lh=fs * 1.3, lines=subN, name='副标题')
                            cur = sy + (subN - 1) * fs * 1.3 + 16.7 * k * gi
                        else:
                            cur = ny + 18.3 * k * gi
                    for dl in e.get('details') or []:
                        n = self.nLines(dl, fs, S_CW)
                        mm = re.match(r'^([^：:]{1,12}[：:]\s?)([\s\S]*)$', str(dl))
                        runs = ([dict(text=mm.group(1), color=BLACK, bold=True)] + self.rich(mm.group(2), GREY, auto)) if mm else self.rich(dl, GREY, auto)
                        T(S_X0 - 1, cur, S_CW, fs, dl, runs=runs, color=GREY, bold=BB, lh=lh, lines=n, name='明细')
                        cur += n * lh
                    bullets = e.get('bullets') or []
                    if bullets:
                        BX, BW = S_X0 + 4.5, S_CW - 4.5
                        total = sum(self.nLines(b, fs, BW - BUL_IND) for b in bullets)
                        if S:
                            self.add_text(BX, cur - fs * 0.12 - (lh - fs) / 2, BW, lh * total,
                                          [self.rich(b, GREY, auto) for b in bullets], fs, self.BODY, GREY, BB,
                                          lh=lh, name='列表_' + str(e.get('name') or '')[:16], bullet=True)
                        cur += total * lh
                    if e.get('text'):
                        n = self.nLines(e['text'], fs, S_CW)
                        T(S_X0 - 1, cur, S_CW, fs, e['text'], runs=self.rich(e['text'], GREY, auto), color=GREY, bold=BB, lh=lh, lines=n, name='描述')
                        cur += n * lh
                    y = cur + (7 * k * g if ei < len(ents) - 1 else 0)
            elif sec['kind'] == 'rows':
                items = [it for it in sec['items'] if str(it.get('name') or '').strip() or str(it.get('content') or '').strip()]
                named = any(str(it.get('name') or '').strip() for it in items)
                nameW = min(max(self.textW(it.get('name') or '', fs, EN, True) for it in items) + 16, S_CW * 0.34) if named else 0
                dW = max([0] + [self.textW(it['date'], fs, EN, True) + 12 for it in items if it.get('date')])
                cX, cW = S_X0 + nameW, S_CW - nameW - dW
                for i, it in enumerate(items):
                    content = ('; ' if EN else '；').join(split_parts(it.get('content')))
                    nN = self.nLines(it['name'], fs, nameW - 8, EN, True) if named and it.get('name') else 1
                    cN = self.nLines(content, fs, cW) if content else 1
                    if named and it.get('name'):
                        T(S_X0, y, nameW - 8, fs, it['name'], lh=lh, lines=nN, name=str(it['name'])[:20])
                    if content:
                        T(cX, y, cW, fs, content, runs=self.rich(content, GREY, auto), color=GREY, bold=BB, lh=lh, lines=cN, name='内容')
                    if it.get('date'):
                        T(S_X1 - 120, y, 120, fs, it['date'], align='right', lh=lh, name='时间')
                    y += max(nN, cN) * lh + (2 * g if i < len(items) - 1 else 0)
            else:
                t = (', ' if EN else '、').join(str(x) for x in sec.get('tags') or [] if str(x).strip())
                n = self.nLines(t, fs, S_CW)
                T(S_X0 - 1, y, S_CW, fs, t, runs=self.rich(t, GREY, auto), color=GREY, bold=BB, lh=lh, lines=n, name='标签')
                y += n * lh
            y += gap + 4 * g
        return y - gap - 4 * g + lh * 0.25

    # ---------- 选档 + 左右栏平衡 + 输出 ----------
    def fit(self):
        for i, c in enumerate(LEVELS):
            if self.layout(c) <= BOTTOM:
                return i, 0.0
        return len(LEVELS) - 1, self.layout(LEVELS[-1]) - BOTTOM

    @staticmethod
    def to_card(sec):
        items = []
        for e in sec.get('entries') or []:
            parts = ['' if e.get('name') else e.get('date'),
                     ' | '.join(x for x in (e.get('sub'), e.get('subRight')) if x)]
            parts += list(e.get('details') or []) + [e.get('text')] + list(e.get('bullets') or [])
            items.append(dict(name=e.get('name') or '', date=e.get('date') if e.get('name') else None,
                              content='\n'.join(p for p in parts if p)))
        return dict(title=sec.get('title'), items=items)

    def fit_simple(self, merge):
        for i, c in enumerate(LEVELS_S):
            if self.layout_simple(c, merge) <= S_BOTTOM:
                return i, 0.0
        return len(LEVELS_S) - 1, self.layout_simple(LEVELS_S[-1], merge) - S_BOTTOM

    def build(self, out):
        self.load_photo()
        simple = self.D.get('layout') == 'simple'
        merge = False
        if simple:
            lv, overflow = self.fit_simple(False)
            # 比原始版式还挤时，试试把右栏卡片合并成「其他信息」，能放得更宽松就采用
            if lv > S_BASE and self.D.get('balance', True) is not False:
                lv2, ov2 = self.fit_simple(True)
                if lv2 < lv or (lv2 == lv and ov2 < overflow - 1):
                    lv, overflow, merge = lv2, ov2, True
            cfg = LEVELS_S[lv]
            level = lv - S_BASE
            moved = [x['title'] for x in self.simple_parts(False)[1] if x.get('card')] if merge else []
            return self._emit(out, cfg, level, overflow, moved, [], S_BOTTOM, lambda: self.layout_simple(cfg, merge))
        level, overflow = self.fit()
        # 左右栏平衡：原始版式放不下时，把左栏的短板块（教育、证书、荣誉、语言等）挪成右栏卡片
        moved = []
        if level > 0 and self.D.get('balance', True) is not False:
            size = lambda s: len(json.dumps(s, ensure_ascii=False, separators=(',', ':')))
            left = self.D.get('left') or []
            cands = sorted([s for s in left if s.get('side') == 'right' or (s.get('side') != 'left' and MOVABLE.search(str(s.get('title') or '')))], key=size)
            for sec in cands:
                prev = self.D
                self.D = dict(prev, left=[s for s in prev.get('left') or [] if s is not sec],
                              right=list(prev.get('right') or []) + [self.to_card(sec)])
                lv, ov = self.fit()
                if lv < level or (lv == level and ov < overflow - 1):
                    level, overflow = lv, ov
                    moved.append(sec.get('title'))
                else:
                    self.D = prev
                if level == 0:
                    break
        cfg = LEVELS[level]
        warns = []
        if self.D.get('summary') and self.nLines(self.D['summary'], 8.1, 521) > 3:
            warns.append('个人简介超过 3 行，顶栏会被撑高，建议压缩到 2～3 句')
        return self._emit(out, cfg, level, overflow, moved, warns, BOTTOM, lambda: self.layout(cfg))

    def _emit(self, out, cfg, level, overflow, moved, warns, page_bottom, draw):
        prs = Presentation()
        prs.slide_width, prs.slide_height = E(W), E(H)
        self.S = prs.slides.add_slide(prs.slide_layouts[6])
        bottom = draw()
        prs.core_properties.title = str(self.D.get('name') or '') + ' 个人简历'
        prs.save(out)
        self.S = None
        res = dict(out=out, level=level, fontSize=cfg['fs'], lineHeight=cfg['lh'], bottom=jsround(bottom),
                   pageBottom=page_bottom, overflowPt=jsround(overflow), lang='en' if self.EN else 'zh', movedToRight=moved)
        print(json.dumps(res, ensure_ascii=False))
        for w in warns:
            print('提示：' + w)
        if overflow > 0:
            print('内容超出一页约 %dpt（约 %d 行），需要删减。' % (jsround(overflow), math.ceil(overflow / cfg['lh'])))
        return res


def main(argv):
    if len(argv) < 2:
        print('用法: python resume_pptx.py data.json [out.pptx]')
        return 2
    src = argv[1]
    out = argv[2] if len(argv) > 2 else 'resume.pptx'
    with open(src, encoding='utf-8') as f:
        data = json.load(f)
    Builder(data, os.path.dirname(os.path.abspath(src))).build(out)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
