#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
单页 A4 简历 PPT 生成器（Python 版，单文件）  https://github.com/staleegin/blue-resume-pptx

用法:  python resume_pptx.py data.json out.pptx
依赖:  python-pptx、Pillow（ChatGPT 等网页 AI 的代码环境一般自带，不需要联网）
输出:  一页 A4 纵向、可编辑的 .pptx；终端打印一行 JSON：
       level 压缩档位(0=原始, 4=8 号字)；overflowPt>0 表示仍超出一页；movedToRight 为自动挪到右栏的板块

==================== 给 AI 的使用说明 ====================
1. 从用户材料（旧简历、零散文字、JD）里提取信息，整理成下面的 data.json。
2. 缺姓名、求职意向、联系方式时先问用户；其余缺项一次性问完，用户说「没有」就跳过。
3. 不要编造经历、数字、证书。涉及事实的改写必须先问用户。
4. 重点短语用 **...** 标出（每条最多一处）；带单位的数字（90 人、56%、$2.5M）脚本会自动加粗。
5. 内容超过一页时（约 30 条要点）先给精简方案，用户确认后再删。
6. 运行脚本，把 .pptx 交给用户下载；若 overflowPt>0，告诉用户超出约几行并建议删哪几条。
7. 用户上传照片时，把照片放在 data.json 同目录，并在 photo 字段写文件名。

==================== data.json 格式 ====================
{
  "theme": "2563EB",                    // 主色，可省略（默认蓝色）；墨绿 0F766E、深蓝 1E3A8A、酒红 9F1239、深灰 334155
  "lang": "zh",                         // 可省略，自动识别中英文
  "name": "姓名",
  "title": "前端开发工程师",             // 显示为「求职意向：…」；英文简历不加前缀
  "titleLabel": "求职意向：",            // 可省略
  "contacts": [                         // type: phone/email/location/wechat/web/github/linkedin/link/portfolio/birthday/user
    {"type": "phone", "text": "138xxxx"},
    {"type": "email", "text": "a@b.com"},                         // 邮箱、网址会自动变成可点击链接
    {"type": "portfolio", "text": "作品集", "link": "https://..."} // link 只接受 http(s)/mailto/tel
  ],
  "summary": "一两句个人简介",            // 可省略
  "photo": "photo.jpg",                 // 可省略，没有照片时右上角留空
  "left": [                             // 左栏板块，数量和名称都可自定义
    {"title": "实习经历", "entries": [
      {"name": "公司 · 职位", "date": "2025.07 - 2025.12",
       "sub": "副标题（如专业、部门）", "subRight": "右侧信息（如学历）",
       "details": ["技术栈：React、TypeScript"],   // 明细行，冒号前显示为黑色
       "bullets": ["要点，**重点**会加粗"],         // 带项目符号的要点
       "text": "不带项目符号的段落"}                // 以上字段都可省略
    ]}
  ],
  "right": [                            // 右栏卡片
    {"title": "技能", "tags": ["React", "TypeScript"]},                      // 标签式
    {"title": "证书", "items": [{"name": "CET-6", "date": "2024.06", "content": "562 分"}]}  // 小项式，content 可用 \n 换行
  ]
}
"""
import base64, copy, io, json, math, os, re, sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu
from lxml import etree

# ---------- 常量（单位：pt） ----------
W, H = 595.25, 841.85
BOTTOM = H - 28
CJK, LAT, YH = 'SimHei', 'Arial', 'Microsoft YaHei'
GREY, BLACK, WHITE = '595959', '000000', 'FFFFFF'
L_X, DATE_R, BUL_X, BUL_IND, BUL_W, PARA_X, PARA_W = 25.6, 380.4, 30.1, 13.5, 359, 24.6, 356
R_X, R_W, R_IN, R_TXT_W = 403, 174, 411.9, 140
HEAD_X = 32.5   # 顶栏里姓名、求职意向、分隔线、简介的左边距
# 自动微调档位：正文字号 / 标题字号 / 行距 / 间距系数
LEVELS = [
    dict(fs=9, ts=9.9, lh=15.1, g=1),
    dict(fs=9, ts=9.9, lh=14.2, g=0.85),
    dict(fs=9, ts=9.9, lh=13.4, g=0.7),
    dict(fs=8.5, ts=9.4, lh=12.6, g=0.7),
    dict(fs=8, ts=9, lh=11.8, g=0.6),
]


def E(v):
    """pt → EMU（与 pptxgenjs 的取整方式一致）"""
    return Emu(int(math.floor(v * 12700 + 0.5)))


def jsround(v):
    """四舍五入（与 JS 的 Math.round 一致；Python 的 round 是银行家舍入）"""
    return int(math.floor(v + 0.5))


def tint(hexc, a):
    return ''.join('%02X' % jsround(255 - (255 - int(hexc[i:i + 2], 16)) * a) for i in (0, 2, 4))


def strip(t):
    return str(t or '').replace('**', '')


# ---------- 文字宽度估算 ----------
# 黑体：汉字 1em、半角 0.5em；Arial 用实际字宽表（千分之一 em，ASCII 32-126）
ARIAL = [278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278, 556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 278, 278, 584, 584, 584, 556, 1015,
         667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833, 722, 778, 667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556, 333,
         556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556, 556, 556, 333, 500, 278, 556, 500, 722, 500, 500, 500, 334, 260, 334, 584]

# 数字 + 单位（中文量词、%、万/K/M/million、percent、x 倍）；货币符号开头的金额
NUM_RE = re.compile(r'(?:超|近|逾|约)?[$¥￥€£]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\s*(?:%|％|\+)?\s*(?:万|千|百|亿|million|billion|[KMBkw](?![A-Za-z]))?(?:余)?\s*(?:人次|人|项|个月|个|次|篇|场|倍|名|元|单|秒|分钟|小时|天|件|款|家|所|条|份|位|台|套|门|周|月|年|percent|x(?![A-Za-z]))?(?:%|％|\+)?|(?:超|近|逾)?[百千万]余?人次')
SKIP_RE = (re.compile(r'^[\d.,]+$'), re.compile(r'^\d{4}\s*年$'), re.compile(r'^\d{1,2}\s*月$'))

ICONS = {
    'phone': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAAApVBMVEVMaXH////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////Mid1sAAAANnRSTlMA+O3WmTDzLwIMNu7jJC7YOD1C15MxwUUm9JL5wAFEJQuRm0PO8pD3vs0ON73rIZzkzzm/PpqWWsoUAAAACXBIWXMAAAsTAAALEwEAmpwYAAABOElEQVRo3u3aSXOCQBQE4EF2RdS4xX1fkrhknf//03JIApdM2fPIHFLVfW7ehyhgwSj1e4br1kjDGXU/5som9wNtnekCn//e0II0OvD+i+ZrXWuCwFELM8DmD7U4Gwi4yoEAAlpyoAsBkRyIIKCoP+QzpD+Lt8UmdkCM/q5DIZCgQCIE8FOfAAECBAgQIECAAAECBAgQIECAAAECBAj8CeD8yS/87LovBLYh9BmS/osQEIQAgVvAYV9fLp9Wnisgy78KceYGOOQ/jdhzAuzLytgJUC8rbSdAWlZS18Dj/zxEq7LSqwTsDBt7xQ0onBgqO1XphXX2LYTPpoYPAWvjeeqN22n62ptUfOU+l1/r7rB7+Ek6fwr+SbicHS/cUJ2aaD689ESpxZv9/GNT2WQT+BbLHyI/MH2/n3mtaBdgTd8QAAAAAElFTkSuQmCC',
    'email': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAAB3VBMVEVMaXH////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////49NrfAAAAnnRSTlMACfr5/iSAAe5vRPOR/BD9Be1CBLOZwYixJTTeI9cT+x3CA9YmEcqr+Hc+Q6rYFDjdvj1m7EzxcBjhf0Eix+sIrfANVTbboiECRQfvDPQwFrvfwOTPU5s5GuPOpOdHzdIOUfYyuFa9PKPmhCcs0fUSmLzgHOXqyytPVJc1nDobXoIgpumHsE7UtZbZv10fhsmLr03y03EPtFIzmlziwwYVdiYAAAAJcEhZcwAACxMAAAsTAQCanBgAAALxSURBVGje7ZnnWxpBEMaXfjSR3hQUItKMiC0ae01sicYSTe+999577/3+1mBk96jHnew8j0luPvFwc785luV9ZxaEpJBCCinWSmyeb5hkqcZk9Xcdwa9zsCAx417hRxkWKJjkn+cH47OsWZUq0MQChiP1/bKgEUM3YQvYUQNsgWqkhy2gRyxw/AMFjLB8I5qGLZBAD2ELvEYTs5D8pQmEKlxwfFfFspxqtHKgL1i9Ke040WMQ/EucpaEt5+jza7dnmrJVbaKLVxoqc3z/PVVja7uW31m8CtHjz6o4rmefJv1K9owW/7osjdQcGkIM+6gOF/tKZZmYOczzPWWZZTX9uYDfqaktn/+mC9NueLFc+yMaspuMtHaPRqvI8AMXWaaN28rhP9iFOXvrsw2noxFfWW9bPX/rBkz5lsh1NHkEZX828eEn2oOaFQUss76CbN3h1fCHPfj+ur7CntxxBGcMHBXP555vKFHM9P3qAE66I7Jf2q/lVljO01U0EQk8sSSG/+QkEeaD/G1LTyvOHDMoBfNHsDYg3c5SfZHSMIaTH7cIw/cP4jsq81W/QOO1243zLwiarbpHcf69bmGd3YeX+I5ApKRdG53EWQbHBbeOcbLlGr38/MUozpQZxPSm00Q5Psd5+xIfzjvwXFzzq/iEBVbTXnSZFO0kSRsW3V1zwuUpskyLtwQIJE/7ftFC7v9S6LrNJ0TikTDzaD6dezGsxcsT4DUp/gEkVIUf8sqenOFuSqDNlphwxq+SXXg8wyXCd4k2tJrLG6GUb63kzOR+WgdMO8hv3epUlj2j/aohPZSq8zLDhDq5ruq2g8YQ2P+j2DHTXAulKTN+phB+wE5vjD2vy+dbvFTn5HeqbHxXH+1B3GRb4PBTI3KISf+jU1fV21sVc576v84qpAJSgb+gAPCxZhv0wWwQzcMWOIt0sAVSvdVhSP5MSttHe1jYPyhQ0gzGT674k/sFDL+Js9mYPUj596AP2i3S34dSSCHFmonfS8GeRXnvHuoAAAAASUVORK5CYII=',
    'location': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACN1BMVEVMaXH////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////ZmOKXAAAAvHRSTlMAAQMFFwf5/fz+9xuu9ghzBvAE+x1QAu8MFvjZFe3xiUKp7AlRdTfCEuU65CDuEFfdez2AHsPSinCzNc1MLddW3HpbqNGTdLzBRYisK5J4fT9EJatyd1g5fMSGJM4K0w+Vdjhdb7uwY6QjyGaPlDITuL3hYOaEIsdlRusnzPXW+jZaO+B+X8Zkp+rLSo3QsU/0MNURVJfawF6CxaaHSYzPTi/Ult6/46VD6Its8nG0M9gUuZq+XOJhurIcNIPCva0AAAAJcEhZcwAACxMAAAsTAQCanBgAAAPFSURBVGjetVllc9tAEJVjxxCyXYcccpiZuYE2TcNMDWOZmZmZmZkZ78d1pp1OM/bu6UB+n9++N5JOewuKwgaf8rmS2eLcLKMxK7d41jX3y0fRED9Ly2KIG+w/pv21UU+YmjEREObvX4Kk5cOn0wgFJwrGpOT1SdFEBdFWvbh+/ihhwPtOQfmgcQthgqXVV0Q/cyVhRsUyfv3OGMIBeyyvfqIf4YIxh09/lYlwwrSdRz/ETLhhWsiuHxtBBGDcwarfFUOEYF/Bpu9bQQSxm+1/GCfCaGXKDxZxA0uHuv6CRUQCo+qZL4lI4Y1q/o+XM4gPVzHoJZKwqtyPL/BQc2PPpFOvd05aUyg/+mb6LTqFBgbejvpPc7oCUWId1WAGC4tzKyGCuzHmepq+P/LslgKdO1XXiyRcE62aKeVJlDmIwyuKQRkc0gaz82B2AOUvtsPvH6HrmkF6Fl5VlsPnB32pwfC9+g41OAbyh/BHdoEB31B+Cfh/ReEGBvDUNaH8sxD9Gu1cF/F9ZTBT99AMGqCIeyj9EERPphmkQxGPUTqYqqNoBgYoYidKN0J0G7VAgCIiULpZGwMzSg/U5hXtQelhEL2fZnAQbHpQeg33HTgBRUSi9G0QPYVmcBOKGETp9eAnc1I+AXgl3EL5BWDucuEGd8GAAbytgdN1MMYvhNP1fdTgA1yWduuQCycAvmBT8Uc+A1+CtTC7HWYPUw7FYeTSBzu8PuTSz6MY7ML6r1rPsmUA6xOpPXMKVk41u33pwgCM+Zla2V1GK0K/EsO88/8Sb6ND6O1HLqVPLWpIN9hshvSkQUoXfVVlFLZJtnyvU+ugIuX096rO8kLkDCrVu8CPMvpLGNrYKoe4/r4Wlka5SdxgC9sooVpUP5Jxytkh2Oub8lnHLZ/EDB4xz4vGMkT007LZJ1KVIgbneGZq9fz6y7mGgtlfefUzUvnGjks5T5LlDu/g9AqfwWLuye8IV1qtCec2UPo5ZqeOZJHp+BC7wQ2x7cRJVv1hm9gCoYpxAOzXogjCymZQKrxiCd3AdIuFii9x/KPV9cMKZdZQB9QN+uT2aMfV9J9JLuqyN2qb4zzxnLpsMcUq0lhNMzitwbLUtxjXv5igxTr2ghHTj3igzcK3DTOY0GhfHfoE1i/SbCueCa6l7Jna7dyPQAbXtdzqx3FNkAVw6alHHbdOUwOPMsZyStEYa2TLFNUF/MP5+tUjmhsoa+f90I79ihfQzjbvkPihG//pH/XxioGy9fVf/fNdipeQ+KeadIQoXkPy27CwsnKukN8qqy88IW8DWgAAAABJRU5ErkJggg==',
    'wechat': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACl1BMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////921Z3zAAAA3HRSTlMA8hox7QT+/P0C+rIGKKffRM73DR0IjgpNvibPR8L7+eHxVY80fUETZusu9jKrARkJvQOttxeiB0+2giSAqRSYQvgrFVao26Oey3nXyK5ayUiwv+DZEQ7aIxznii115mWJUNZuDPSXFrs6IAsQQw9SM7lcw4WWLM0YX+9e2KQnmwVsr6XQ0+5d6oiE5CqQuEDEPRI1V37wb9zRtUzSmTdbIkZqnWmsjbGSMNVzeJzAyoxt8y9TuljefONiqotxtJWaOJ/iYUVKhnTdzHAfS4Nnx3fGSejlJaE7s2DUc/cwgAAAAAlwSFlzAAALEwAACxMBAJqcGAAABH5JREFUaN7tWWVbWzEUTimlt0ChUGC4DtmGDZiwYWPI2AZDxgZzd3dnBnN3d3d3d3fPj1kF1ptcS9LyZU/fT3Cek/dtcpOTc04AcMIJJ/57zEncEBq9x/DIT+Nwas+x97xOt4FN0EZ6dAo0Oox9RFl2FBQifPpAP0fQnxqlglLgzhfYu1pbMqE8Iod72kHvl+4GFeE9LZmRXjdcBYnQtjMTf85USIo2gxn4Wz6BFFjYgpY/LgpSoVdXOv6hXSAlerjQ8G/yhdToPZqcP9gVMqCnPyl/Bx/IhPakAgHouC73h9wsDRfQuXplrb6AmlYQblAOGXXN8vXUrTD+gAlm82FUsx2RQBgyaHZjrDmA8q9rdPZCrMUk/IEo0/VGs38E3+rWtGU6ouF1FoFAECowpcneD1mMf+uJeicp87fG1nqr6Ay4HNEZwHDl49YfE/DpYLWXYd9Y9BtA2F1RYJggGOvN5kR8Fx237KIluLePEv9B4fnRZgyJrdAKz0Fh1ovXwltUKehFsx1i7XKv9x+i35UXJLxVEOjEwP5tXXkf4jh3h5adK43TmcYZ87p3qyxaGRXREFmSPWOo9JEeQ0n/o8Y06Punj/g3GjZjirhAL7oIvQUAzdZK8dwjJEHsHg2hoHeP1QBdWV9ph0FLPJUChRxa/TSde4UZF+XhAunkV2QN0G1zV/Jyq8PWaRQpf98UUL2bxHFvMCIwmJC/qhak/CFzdd0vF0ylcrlEULuTdLKX43gCaWQJ12/QZxf5dlCpeQqPSUb8Asm7ac5LP941kUfyi2rBDroTH6SzJe3eyu6rQQ1t6nfGNoXbis7bq0ESYshQj0iMQY7EEf2ksdn8CKJK+SegWaAkEANaItGnrWX+0TzbPgvTNr7TQ8m8RYh8UIz8b7COK7FZrOVn2jyeky8vgneT5/cB8WhyP8A6zMN2v6WKRM5Ym4DxhqzAKrBCLJPo2mCzTLZY8pFLop53FtSy1Z8BFGLHeqJpzLiN/DhtzvDUWCDnl4qhclEyHuA7mcu8m47uW99jFzM4zGs+PyTN5ST5vwKjG2RBABJVCyKk/D6TBkQcHujFEDpSwm892MQmMB272wb0FvdLAl/YBJbjt6exUNSvAmxgE6gUFoOiq7QRqNkEhLXDc/GKG2xmE+gvEBAPGduTwRomgUCcf7FEOb4WVLDwc9W4QGhT6Go/OXZplc0zASSwCGQKVqiukd6a18QbLsWcCJv9qr7kDXDRMgh0xPk15qJJW7xILEVeSs9fJGgfGkz0CxdJdJPoBa4KSHK5ZXpHlREQhglbgjF6mULoCuVXUFH3DSlruWnUPcPFPWn4TzJ0JVOqKG6CNJa+av5KUv6pqWydYf0twt/vDxgxjqi3l2tHgz71nHIletS+9n+5Ui4+0973ixZn5XusC+x/Ipk0ULZYC3bAK4xu4kz8jtPuerDe+tdLxzwkJa99mhvkbeneuvcozTpkbu88CxGPcvZgfLzLeN7M5kea5pLTnE97njvmwbnN+3q4edWyZn+gdL7ROuHEf4a/6whiPQ08kqkAAAAASUVORK5CYII=',
    'web': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACdlBMVEVMaXH////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////8rISQAAAA0XRSTlMA+y3O/Qn29wED8PX+ye9xHPoCCivpzAYdNvkfYzDPbTPigKXLcvyiNe34HiRzdnCj6zHbFxsgqyrzL+gYkya7qJVFtAdmGj7nTJk3Ro50eZ3q0PR91LWBSCnyLhmfVifFFdGmERAEWVoPamxHa/GUMnpbf+aEiUssb6eNbpLVoe5OkTR344uQFLmarFiWFijXYMdEyqDTUg5BI55Qv2IFOXxUqeyKrlF1E9wiCLw7fuWDIXjfXuSCJUkLU95nqq9XOHtcpJc9zWUNt8FAur4/OsZo1BYAAAAJcEhZcwAACxMAAAsTAQCanBgAAAVpSURBVGje7Vr5X1NHEH/BkASEkCAQQI4iQgBBRUE5xCIqiqIWsVpBRaCFel9tvWq11VZrq622VuvZ+77v+77v/Y+azOTBe7uzb19I8vn4A/NTdr47M2/ZnWNn0bSbirZcLCu7eClh6uf9nM9CVHWmLq5qp1VEftzRxSKU3Rlh5eXGqv3O3/6sbMefHzvYCPn3Is/nvVZwVwzqnwhWMXYOfz9Yzgzk2I/cxYy5L8wfo/qsQWdI1XkP2upmJmpErZ5fQ7/Lf1o/BvWTX3SCpiUwCqxmHE0IANAGA+ehaDcj5WAj6lmI4x+YQKsQuRVHPcWZ0ejvTdfVPIcnySUamIgf/aw+Pn/AtnrPa35dajNyFjCCdiGWo4/zM2zqnzQ4quRV4NRXUQbceH6HRzkzk+3o900dlWgeAtYRRtIK3P+uUU7RHLX+5bUGFS8Dq85FG5jYD/BDBlZ2hUr/LceMKpKAV8YkdBDgJCMrNU/x/anG2UXILJIZyEH8ASPPe9z+97M16NFMSrjCe5ndNVyupOQfkRu4Dz/LzKydLNP/xYB55hvoFe/JDTwMgcpTwkURyWn1tHLiZ4BdyJjqb3Qbxy2lDZzgpb8E9nQrAzUw5W6eTfr0bgc3y/m6MZrR9Cg6ShrH9jcR8TOHF16NgcNtZcCVLB7UMKWnCAZEb1qGmZhZUie5CYzdL+SXRmHOOeKQC7ROz5wcdX/AGfhWlG0AYK21AUxIvSIw16y/KU0MZZiillobuAc3UNyotCSTgbmi6FRMBUxBG2DaBBFoNepf7xQnbMRqRWWgA6YdFQFnlsHATEJyJSBTVAbehWkrCSSojdM4xY0WkfHyCGDFTElfw8RdFFT1FmBLSLl99tyAsT0WE98BbBWJFQPWojawwmKp0wF7jMQeB6xVbeAonTUN8ehJEmsD7G21gbUwsZPEnoE9dpDYblmUFLIjVRvpFf2NEFRNy1WD3Gy1gdth4nwaDNdg/9LQNJArURuYjSmXBv+Tbg+bB3LNagMl2AWgwRdC0FYawvu1V20gG2/sNFgQgmpoCAsel9pAKpZPNLgtBD1PQ1g5lasNuDHtM6mnJdxAwv9ECd/kYRqqi9sxpaM1y42Lo4U7EZcTGSrC8SBAB7tZcQl2Q/Jw/ZT6emOsr7eT2ICs8A3Rd4AdUhu4ABP3khjeBj8ksa3xSJk1Fkl/jd2kj52vdST2vUXZchKwbWoDi+UtK1dgvCgcp/iR/BK4T3VK0V9mEMigrWtsh8rAfkmzgruJfyO7iLfbu4g/reol9MbWSsjsFhdQaN4FoobPst8MaRCBV/iXmh4yYWva+9YGpsCkAnH9b/IHSTwHp+00pPAacZqpu3Yp6eTtwV5LbUA4IWJLTWvy89vUb7cpeIo/5Y5CytsyeOlNdtuam3h2Ge3PG8muoI3G7DJ7jVktmStSPsKG8Ceq1rL2lZlbNEkWkubUUrWRsjnexDXHLZ5BZpnb75+q2vsHiIRv/URx3CtWhfLyLvJAYXrE8yoeWcxrKKSPl7m9a1ph6nJVaqjIFmqXs7JHolMAHzawuvLUyafe4NIvLRLemQx0GMChpdE9c4VOa5BvidST4cK9gfcyew914cCXb846dN3WgtjIgvNn2M/RDZt1qb9gnEs9lvoA+mXk2bMhmiogc8dOFPsbx3tEAycQ+QNHO3dkRllo+EqdhotIQPCFnBuG115nqW8MtUxfMHy5ysFw08c9YPR8hoEqnO3Lg31jLJeqr7j18lzrMF3lHJ9req50X6mO5V9Brv/jjaz+R0NC8g8jr73y2vUtsVZ9Pt05t4/4eHOb7va5cS0wzy6Al3F3S3/CathLVzMyrv5+c9XV/wMznaVgIelZVAAAAABJRU5ErkJggg==',
    'birthday': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAABOFBMVEVMaXH////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////10QvYAAAAZ3RSTlMAC+wJBMYl0QfzaT/tgAjY9ZlGb8ruydMN3NImBj5HMVTP+P2+PVVmj9c3fx3Cx/rgAkXQbvQW/l6CxemrbS/EZ84PlcgiTHCjKfwy1sFfiAzVSdSBJC739lYsZEqcwESRm0NICppcxcDC7wAAAAlwSFlzAAALEwAACxMBAJqcGAAAAgFJREFUaN7tmVdXwkAQhRWE0FEUkW5X7IqIomLvvYC9l/n//0A8ARKSnRyOs0+euY/37sxHstkNZ9PSIlZPdFiKj+oAYE+Gj2oX4HBIgo9qHwCiLrqPaae9UgDzZB+Tfft3PMzMEn1MPYOgKkTzMS1sVMfDKMnHFOiqjYcuio+qsz4elik+qm6tYJriowprBTaKj8qtFfRR/GZu0RbFb2aSxyg+KpdSG79C8nFNVCsmN2m+xTWsTrndDk8b1Wf9VaATxWcAAxjAAAYwgAEMYMD/Bcy9BXNAVi7z9SJsHz8BaXp2mvsnwyBRYdPZUVxqfwBfyQAYBMk6McwvSNd7A+BJPqCjARCUD8hg5xOy5EZXrywxgAEMwACK3+tsbXUmx/vMFZElb7ySea8VcXb+WUqOZy0BsZRug802trh61LJUzND+UstCCg7oSOi9Nv+AfnN0WWSNdWd5MaB/0fia622vtbixyPpPjdndrQgQTJtf1MVjtcdx0ZyVr9RMVHcUNAJ8Do9deNQ/UrD5CusB4dnkyJoNq7N7HD7s68F6AS1zPeAtVZz4p+iVzlQv/Mic3Ufwm1K8ULOLsnX/YYup663/7bN8HL4t2tcesCYeWjzLvyaw/h8RfPkYl5Z+SYYMK1sRfIxCNoDY7wZQ2Tj8gqy+qWSRjLdrBjCAAQxgwH8C/ABsSnrVS04x3AAAAABJRU5ErkJggg==',
    'github': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACslBMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////+D/333AAAA5XRSTlMA/vcBE+1J/PkCEv0FZhvuAxp3CB9ZfKIMeLo5KfSALwm24U0hMx3nBPb7MRyRNHvi6PPMRSC58n9eZ3onu910FDzcyM0ywm/WnX6IB+8rsdXFq/iWFWKGyapsLoMqTDAZP8tULURfvmkRz64LYansKI/ScA6zVp7mhCIs9Zg223ngwerAoeSMoIHEcQ/YibKw09ojTmBqjuVTk1o+dVtoN7jwQCWnHjrZUWNQa11z0EJHvaNlRq3RDbdV+rwmrI1PNRac332C6cpt1Fib3uOlQ4vOlXaaOD3DpFLXS58XO3IGJPHrN848NQAAAAlwSFlzAAALEwAACxMBAJqcGAAABY5JREFUaN7tWudfU1cYPiGBJISQEAgYxEgAGTJkgwwZZStaHIColamytFattdbV1t26ap117929995773X+j14y7xl3nZv+fv3Q9+N7zn2e3HvOecdzAsB/xYwxPX03ymy2KvPxa7/9Glzs6J6W4f0hkG8ldcurTqcHBT2htGEOpFt+vCNFJXpk4y4obrfbotjhJ8ZNhtJmWvMtG/zj8VCuPRmmHP6RNVCJPThRGXyiMwQqM1PcmwrwN06Dyi30Rbnw2kYDZLK4cbLwdS7IagcmycCPuguyW8jdkviLcqEa038ggT/eAlXaF6L4M0OhausQwU/NhUEwhyB+9WUYFFsmlFKKg4MPTYV0gnIYLHsvkoZfg8Tm0IJPlEA+UIBEl7YIEn9cEvJED3ekP3LpZaF/uaoIgAhkA04nCfYgz6RZ3c6Y7AxJ+F2fewDW8p2TM4nUiwa4t/w79zONKHxBjW/mUcTfacQIKtHnyniHu9P7Ug0XVjq+Npt32zq2xm/yfLy09gBOKorQheKHYT9sIz98HzricpzTYSmp3vxU0scP8V0rEIRWZAzcjxFEMiRxLAy/g8Q4/NNGMxAsRSEsq3ljs/4FAjglMFRE7JR5DASdGEZe4LRNIXbfeAYCItQn+4c+JAimKsefR4C86hvKJM/PXOUEdxAgBt/W7qDUtMoJppMoZu/QbWTuZig2dZsImPneOPoSMdLIUi3f2YrDrPAEkmcI/L1s9fhJAqje7V9HLE4MY0NBFPxVbvde3N3M2rGcxxPUD273TzjBY8w9ER5yCsacVryWzmfvHqdiUCXasV4Gf4Hv2bu6i5AS9geIMp+dQDuB8rWP4gQrVfS9b1PiHf7daCWHbNuPYZ3kfNdxgu0qCB7GsE5xPlsQIqnfPsWwTnC+WJxghwoCPOnEcr4qnMClQpcxYVg3OechnKCJnaCC1io8ijszjMwEeP0Gl3DOg0SQXc9M0ELL+4sJgjJmgpdxqDGNJJwgmM2K//x9OJQ77ecT/fpqRgJiOScI5PxTjARtOFCv2/0VQZA0gwmfrIw2uP1XyYLGxoKf2CSwXRaQBIbFDARzSZwcT5a4hxzZVqQYfzeJYvG0kmArpbmbNlOh8vwuBcRb2QE7rX0sadcqwL/3DA3jae/oIF1ibHtCLnzWdiqCvts3YbbPtTQK/HnzVqDG3lwhjR5unyUgIsb75yT7XLXc0U65aueljS3LHWGCLefgd13f1AnLDYEuJiLP5xu6CEZugJhaLIA3vfAG8cPnD7WKawz5PJUzkDb/GuWi+nDq39K10hIpESOb/3MCBZPmOoi+Aq5hak045QNJ6Fd65G4hm3eKOXX1Us0FZLKTtgKTxAnQoB/Fq/kuj4DRWzq+wKRJpa5xregL9KOTnWitUQ3O8wSqA/RNdEyMYC0uePEC0lnui5+rrBj29/+H6QRiy5yRhc+2o68weOY06D/Wa7EkFZcvkltD0IUKn+0LjB75w70uLjsA3Vbw82EBaURklZ+jqILdPNn6d3cCKR3lklQ5cP4SQSd4RXiFc2jz+wKaS4hPiTMWWRXUcX4rpT/wGq8csMvpvIXwiwVqw3S+sFZZaJW8JBTAz60WVGSQ/iFv38LjYQMDP5r3JAtcZtDxz4pceWXR9f3NArcN1MmGQrG3XkC9oSinT15P3UAS1zgJ2+SdmjHrp/3+PqmVS6GEsFiB1yVnznldxj0XWau2CNzWEhPTLsmpEWY0yzw3RGcxJFcHWobV+U76NExM1GxIlF9HNSCP7qTPqkd12INKKsH0EyW8ZxdKR1NTs05hLVudrZciyAngv88iFWc+62ur10kQ9F7RAiZL2OH5x0O7gFLqvrfSVBaq0B7Cu+q4C3uhjmqVAW7ZqfavFSBLZHOPRBrB/6bW/gHUffh5vlzVdgAAAABJRU5ErkJggg==',
    'user': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAABpFBMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////9bMWvtAAAAi3RSTlMAEwWP9BXt/vzzjPgk8fvrb/r5JfL128tZbRiRuUIEa/bXmceJ1hcHFhsBU87YVxnKhSPsza5RdVY3f2AiJ1C3mOWN0LG2wD/pBu4LTv1840NnLrT3FD2AWm6QpMhmiiiz3L3ho2qODbJ0ElWdwSFkAogmDHNUeCBEaM+W2Xebvx/oyYtsld2+n+IclCRQlQAAAAlwSFlzAAALEwAACxMBAJqcGAAAAq1JREFUaN7tmGdPG0EQhs/9zsZgG1dscMHGNBd6SeiEFjqEAOm9915Jb/Ong5QvrH0+b5mViOTn653eR9q7md1ZRaGlaTcaCy55vUvBWDTcpCBj2gsBQTFiQoz/sx+AMqyWDqR4s0cDXQJnsxj5YzNQkbdPxfO7NDDAPyiaf6IeDFHfiOVvQFUiQuujVheo4/z5dW6gQLvEm784A1R0mzkFL4GSO3z5HVZagVbgEuwDNdNc/S1AL7DydL49YCDOIQixCL6w5xccLAJHG7NgF5jIMwuibAILsyDGJjjFLPjBJggyC+xsggZmgY9N4Dp+AulLdEX2R5b+m0ovtDCbYJhZ0Ca72SlFFkErx34QYRHc4NkyrZK3TMVCL3BKPrb4OSeqD7SC57yjTTddfr9Z8uF3jP94PUhzfF8WGRAuVK/hR5JHqA3RIW3cb5TvXlaEaTT4l/rrMAblFo+/0iC+iDTrF27qFLV1GvNGxPTtPLE/OFrjd7EvXNrzlomc3eez5yYs+XalRo1jgO2ecyvZlzj85z+nGys0kfTC4eNEX/Ly1xc2xnRP6GhdOeaGNktf2Rx6QLxSjNA7Tk66yjvD1EBPeKTXls3aekfCPQNTOiPCu1dU8c+c9cCJN326en7zGgjwe75KfOeoCkKoq52Grf8+CHPd4BBjvg0IzLZUXJ9fgMK5xxUEo4DEqn7+vIolWFnXy7+1Bmhc06sHJyDyRKc/1GMKvJ/KBClA5XvZmcGFK9g+UyLwADJxkUtMjsnZ5sAWrJBrNAzodMkrgn9cJARb+IJZQvAeX5AkBA34ggVCsI0vSBACkEBNUBP8bwINP99NCHL4ggwh+IgvmCQEr/EFzeSeeRU7f65k03+I3LDtZVdVO3bU/B2dqfcAL/+n/lXbeiqDUA9aJnX0+/4F9ltf91/nTXoAAAAASUVORK5CYII=',
    'linkedin': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAABblBMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////8c2ar3AAAAeXRSTlMAmvQB7vFY75c1BgXZa/bH+f74/cMHBMxqCCxv1vrGywwbXmOw8xAVGslIzmzyebgCofwZpMhm7IooCUxRmTcYvVt/wqPmIkvwDZMSVZgXfsGiIacmaUpucxG2VNo63yDFpkSHJcpozwuRuoHEQ4aqLrRS9xQ4HbGVIkpHRgAAAAlwSFlzAAALEwAACxMBAJqcGAAAAdhJREFUaN7t2tVTQlEQBnDSAGkUFGzKQMLu7u7u7o773zuMILuIDzrfeXA8+/jtzvnNMNzLPfegUsmSBavC1RqNAiqNxWTIXL/yVIGWq5Gvb6hTwFVuZIBJgdcJA8rwwBUD/HhAzwBFQEngnwP6w55A8ZLbKwq4CH6kRbM2IUCk6DOvFwHoJknDJwAYoI07AUAJ62zBAQe/kfvggFXLOvv4jyjIOoN4IIc2Kh7xgIV+RmsiLrTqdD6iEwHYH1LxrlfMvcg2XZkIi02twm7XjuaYekMjf3D+AqDRlf4csOlI2ZNhmGTxj7WbqtoTD+kdJVMTOz8CcmkjLxnmk6wgMdRQQZJo7TEYiGQ8/qu0nlIk8Bb9ukWqteOAWNZNWBUMeC3Mvs27RQEv3+wjF6wg4NtaFA2siAbOrFBgfeZL5IQB2v4+s6KYnV08vkEB4/ep4QaWV6OAofT0Nc0PQMA2mZ6jjSMQME+m40+kEQABw3R8lL4TwgAhdj156JfLAQGeGXBJx80QoJMBbXRcDwE2GXCOB3IY0C0aUEtAAhKQgAQk8E8B+TJEAr8Dwvj1x0QfNToZ4MYDLfy414Vev4afEqhC4APl8t7MTbqh3gI7cvfvLRvlnxhk4eoduYJqUCldnKoAAAAASUVORK5CYII=',
    'link': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACcFBMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////8h3Fl2AAAAz3RSTlMA8QoJaQMC/gzCAesN7y+s1AT7qGruxA5bmLr8vhdPgK1Zhi7nI2tlkzvlB9XawJEQ+B8F8jOIxVMcsVTXhxmwjNtmlBNW8I6yUNa3+iGNMPk1eAZtCzT9Q5Ueri2EWskbqexMj9Kz9rgYPOEiJwh0EnmdfqJk9JJzEbZ9giCmq07zltl3WF0+JKpsr5BxtFL3dntc4tHkzChVyokrV6SFyIpRMjd6vZ6jS/Uxwcanlxa7nDpeP2NE6WhJKnIVmzka44Gl6Ejt0w+5n2HQy4tL4xJkAAAACXBIWXMAAAsTAAALEwEAmpwYAAAFZ0lEQVRo3rVa90MTSRSeEBNCAgQCGMALLWABRJoiRUBFBA8FBVRERM/zPHvv7e4sZ29nb9d77733Ov/SbWY3knkzO7uT3byfkjff++bNzNs3b2YXIQuyuvi/toxSB67/JaOtfdVqZK+sLJmNaZk9K2wbu7v5DubJzXkuW/gLpmI9mdptnX7zo1gkf5ZZ5J+ZhsVS/5wVetcz2FjyMuPmrziGzcgTSXHye6qwOUmJr4ekR7BZuRjPLHlSsHnJS6j/EclNqP+KOMoS6r8i+Qnmx/gp6/OTk3VqZcgdCudm8R7v4y6L/v9WmPoQ0T/nCxYw05r/w9UUqDqLQYxa4U8u9AKYd1cyBN2zwD+Og9wBUSP28iOvH8BaTPD3f2iaH6HgYwDYasjv3c3jn6gHXwKQ9w07+EnCf0UGO2noXSP+ZyfI+K8IWIUiA/7gj1L+K7KHBl836GCRpP8ILabR5WL+r3sk/UfoCMjZYvRRaX70B40PCMELA5Lzo0iZzAgWSfuvhB1tcUNYY5XDDs5qLVteq2nR2XKzaYt/RB0sh/yTtYZ0khB2OnlGT9ImflEH8wH/O1qpk+5T/7/QwdpU19A2M0Qd3KSxy3ppfowXNDI2I8Cp2wL+VJAlslR12DemqtwObDbsBVEREnSwFDjzKdGGrsXqxtPr4L4IbDJEM1QMCgh1Bd6jtR9RJrNgWBSKOljCy4sLlwGK4tgQZfbkW2zsd2/bqP0sobHziXIcUyC2Cvgb2P03X1EPqb/30eAr3NDFeHfUdhLDjwt09vfp5M9kGvwNTxkJlO9V2xV9TFOlW6d++JK3Yi8S5SV2gzioGueyLV/p1SeHyN+rNPpnFTSXHUI6aficaejSrX8eJ/8/AXuTVsgNMDn8O6J3vgXUabW69ZUakR8Dg+kacC2sA3IqiL5LXPnG1m/riKYjmZ9L0QX+OeAsrZxG8ztjz6fr64huAW3RdyQKPgM6GOKkliL6bOA6R7W+TJR5mJvtlMe5j5dx6i7HqKq20gMA9wPZRNkNp+JCFA5CKfkA0f4ecw730PwnAVObem0Aj0WBtfziB79CtMO6/JsdwGCvmoV3wiEEBrS0XErr1euV83r8iL3/WU70vWx2matarKG1O4jykB5/AfuQF3N38Ihc4hU0z8dmWobfxd6P1Wib7RR2CPtIwzZezC/l86NmNrvsibY9YNraeZuRmlvQUOR4z/CjO4L67Qe4/HgO0U+jlYc1+Mb9b7rZ+0+GPzumdRVonKBmsCZau1+09ZbADkqo5tfpxjfUhxYMrFnUAby/baIH6WyIbdz0LbeoTBfdP8M7yQ8AYHvlWKMvzD2I9SSZr344lUzjw6zq0zxdDKL3XdEMtdPYy0EW0jFK89c1YIkKtw1cePIwW0lU5kS3tPdhWEwRdZDBC3NGclvW52/Rfg9A/mteUQelEs6oGYdJH8IZQk/T4BVG/JNgcYp7xDagRAhK+4//QjIjqJXmrzktNinlFz+m+UFmMYyiz2T5N3kMOgB58YxofTn8AcOwA09y2qCU//iqYVjDhH9Fxn981G3YQSMw8QUl/M/vN3GrCPcDv9c0f4rHBD979HyVd8URPz8KM3vyLjgG92EL/Aj9y5hm0bfeB5qwFX40j/Na4MTYFcHgiRwe//hU8+/2eG8+HWtOvb0hs7X372Od2Jr/irwk//JEwv+IHEwwPzpdL8cvNT9662yj/7zznr3+RyKpKKH+R95jmX2PWOVB8UmF3xT/OWfc77ozzazDeVdCvwZwnLT4vUGZ+HuGrlobvsg4rktf+atN35SMcukbCtzILrk30gLvbwtvIXul9f7d4XXlnTjtxnX/jNshK1T/A5cWLGcTiJ0aAAAAAElFTkSuQmCC',
    'portfolio': 'iVBORw0KGgoAAAANSUhEUgAAAGAAAABgCAMAAADVRocKAAACDVBMVEVMaXH///////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////////9xxQJjAAAArnRSTlMA+S7+4zg2+AwvkAv3k5Iwke70CuT9NzmuDi1LAsrrtOpDJN8y4YQNxQFdYsnt8hQa/K2qZvZqP1gIh6NChQTN3DzHbzES+iHvK7EWwF6CYwYQ3sToBc5XGR7BJt3x9YGAWykXQBVP2G7S50fsiigJjxMYf+YDqEbwslB0VeB+y9VzfemMz/PUclOWNHd8v6/TcQ+awmUcrLCgpFbQe8i4jtZfEVmrmz5Ia2iZjWeFplepAAAACXBIWXMAAAsTAAALEwEAmpwYAAADjklEQVRo3rVa918aMRQPoB6CCxAFFUFRW/feu1pn1Vqt27pX995777333s3fWKooVrwkLyHfH+/zve83l7u8e3kvCDFhKmk+vO2c22TF2GoacdgMJd8+oUAh+WelG/tD9+HrZ7u4etzYAlaH4szMFVG3ZGZhGvKdSbzymztbMBOuZGznkH9YnYOZEVSUAJRP7+jFIPTVw16tA0Ohq8xjH36Dgjlw/DSj/sEtmA86w1YW/ZRUzI3aZLp+h4IFkB1L07+uI9weZggxGmOiSEMw7yLrdxK/du0SKTSIQNKUkfTHSPph2mVaKOkZrF3q+g+I82vwESOI0SleTT9DRzQI8TGDiUTzLZXvs5j8heh9VD2Z6Vr3ax2IxswGGyjU8xZ//QTq+mWeIg9u+hucoi6iKB85nBo1/F60lr6AldBlcqSGSu6dWTNBWSy/Fa9DZDQD2fa/QSZTpFEigvX62HANE7lqtX73JWjMofNrV/9F0+Axh87PWJU/jHLEHCo/tVzlAVhjDpVfspJftXAvKCL/2vLlLt6YQ+OXei87VG+wAw3W8Au9SZBAzCHzm5ZSpXGBmEPhzy9eXRCJOWT+u8X9hVjMIfIVo+dykWDMIfJnPQaVWCJ+eQzcMg3mEJqRqY+tdjQt1cCzmH/INfjDFjD58RsVyDWwoU1yDTaibXIN6lCiXINEtEOugRlZ5RrkoyahjIhhinJEMiI6XiOXQEbEgJdoTiAjYsArNCTwu6dj1I4K+TZJjHiDUIVMA7dnr/ZR4hSZ/2UDIQIZEU1/caNWruHPiMhIbVzSeiuQEZGGP270Dva9SEakUjW6f3Sie2W2vwtp6Xou1LeTK2nlIgG7II6hSPecW950kqnK2M+rHzTJWALP4Rx/K2sh9hGfQT9zpbdVx6PvBJSq7/IYpAAMjnHoO0DV9h64wR2QQRLcoArWkGgDGzTDDIbBWQS0+3QIajAANKjpAxrEQLtOjcBJKgL3tRqgWxco0gtBBsoU2GEfbDd1Ed78awdtpw5z9Bf3miAOPE3S0qsAgzSeDmww4BnucfV4tTsBW2AuNNdJNkB5Q4wGJ3g74TUVbP/oPfy9/PhsFoNpgdMCg066vuuI0HGH2TM0g8uCByos1eQcfncuEsXgfjMhbb8RiEMnB86qFntuB+hci6WsYL0OZ/ETFDjoy7641hRDnw2jAGNyosKWFW3ShJlGHr8oeQq59S93tkBme65lnQAAAABJRU5ErkJggg==',
}
# 透明 1x1 图片：盖在联系方式胶囊上承载超链接，整颗胶囊可点，文字不变色、不加下划线
CLEAR = 'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII='

MOVABLE = re.compile(r'教育|学历|证书|资格|荣誉|奖|语言|兴趣|爱好|社团|技能|education|certif|licen|award|honou?r|language|interest|affiliation|skill|activit', re.I)


class Builder:
    def __init__(self, data, base_dir):
        self.D = data
        self.base = base_dir
        self.theme = str(data.get('theme') or '2563EB').replace('#', '').upper()
        self.card = tint(self.theme, 0.1)
        alltxt = json.dumps(data, ensure_ascii=False, separators=(',', ':'))
        if data.get('lang'):
            self.EN = data['lang'] == 'en'
        else:
            self.EN = len(re.findall('[一-鿿]', alltxt)) < len(alltxt) * 0.03
        self.BODY = LAT if self.EN else CJK
        self.BB = not self.EN   # 正文是否加粗（中文版式全加粗，英文只加粗标题和重点）
        self.photo = None
        self.S = None

    # ---- 宽度与断行 ----
    def cw(self, ch, size, arial, bold):
        c = ord(ch)
        if c >= 0x80:
            return size
        if not arial:
            return size * 0.5
        w = ARIAL[c - 32] if 32 <= c <= 126 else 556
        return size * w / 1000 * (1.08 if bold else 1)

    def textW(self, t, size, arial=None, bold=None):
        arial = self.EN if arial is None else arial
        bold = self.BB if bold is None else bold
        return sum(self.cw(ch, size, arial, bold) for ch in strip(t))

    def wrap(self, t, size, width, arial=None, bold=None):
        """断行：英文按单词、中文按字，模拟 PowerPoint 的换行"""
        arial = self.EN if arial is None else arial
        bold = self.BB if bold is None else bold
        toks = re.findall(r'[\x21-\x7E]+|\s+|\S', strip(t))
        lim = width * (0.97 if arial else 1) + 0.5
        lines, cur, w = [], '', 0.0
        for tok in toks:
            tw = self.textW(tok, size, arial, bold)
            if tok.isspace():
                if cur:
                    cur += tok
                    w += tw
                continue
            if w + tw <= lim:
                cur += tok
                w += tw
                continue
            if cur.strip():
                lines.append(cur.rstrip())
                cur, w = '', 0.0
            if tw <= lim:
                cur, w = tok, tw
                continue
            for ch in tok:
                c = self.cw(ch, size, arial, bold)
                if w + c > lim and cur:
                    lines.append(cur)
                    cur, w = '', 0.0
                cur += ch
                w += c
        if cur.strip() or not lines:
            lines.append(cur.rstrip())
        return lines

    def nLines(self, t, size, width, arial=None, bold=None):
        return len(self.wrap(t, size, width, arial, bold))

    def tagLines(self, tags, size, width):
        """标签排行：每个标签不拆开，用「 · 」连接；单个标签超宽时才按普通文字断行"""
        SEP, lines, cur = ' · ', [], ''
        for t in [str(x).strip() for x in tags if str(x).strip()]:
            cand = cur + SEP + t if cur else t
            if self.textW(cand, size, True, self.BB) <= width:
                cur = cand
                continue
            if cur:
                lines.append(cur)
            if self.textW(t, size, True, self.BB) <= width:
                cur = t
            else:
                w = self.wrap(t, size, width, True, self.BB)
                lines.extend(w[:-1])
                cur = w[-1]
        if cur:
            lines.append(cur)
        return lines or ['']

    # ---- 富文本 ----
    def rich(self, text, color, auto):
        """**加粗** → 黑色加粗；数字成果自动加粗。返回 run 列表"""
        runs = []
        for seg in re.split(r'(\*\*.+?\*\*)', str(text or '')):
            if not seg:
                continue
            if seg.startswith('**') and seg.endswith('**') and len(seg) >= 4:
                runs.append(dict(text=seg[2:-2], color=BLACK, bold=True))
                continue
            if not auto:
                runs.append(dict(text=seg, color=color, bold=self.BB))
                continue
            last = 0
            for m in NUM_RE.finditer(seg):
                k = m.group(0).strip()
                if not m.group(0) or any(r.match(k) for r in SKIP_RE):
                    continue
                if m.start() > last:
                    runs.append(dict(text=seg[last:m.start()], color=color, bold=self.BB))
                runs.append(dict(text=m.group(0), color=BLACK, bold=True))
                last = m.end()
            if last < len(seg):
                runs.append(dict(text=seg[last:], color=color, bold=self.BB))
        return runs

    def mixRuns(self, t, color):
        """拆分中英文，英文用 Arial"""
        out, cur, lat = [], '', None
        for ch in t:
            is_lat = lat if ch == ' ' else ord(ch) < 0x80
            if lat is not None and is_lat != lat:
                out.append(dict(text=cur, color=color, font=LAT if lat else CJK))
                cur = ''
            cur += ch
            lat = is_lat
        if cur:
            out.append(dict(text=cur, color=color, font=LAT if lat else CJK))
        return out

    # ---- 底层绘制 ----
    @staticmethod
    def _font(run, face):
        rPr = run._r.get_or_add_rPr()
        for tag in ('a:latin', 'a:ea', 'a:cs'):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set('typeface', face)

    def add_text(self, x, y, w, h, paras, size, font, color=BLACK, bold=True, align='left', valign='top',
                 lh=None, name=None, bullet=False):
        """paras: 段落列表，每段是 run 字典列表（text/color/bold/font 可覆盖默认值）"""
        tb = self.S.shapes.add_textbox(E(x), E(y), E(w), E(h))
        if name:
            tb.name = str(name)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE if valign == 'middle' else MSO_ANCHOR.TOP
        for i, runs in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.RIGHT if align == 'right' else PP_ALIGN.LEFT
            if lh:   # 行距写成 spcPts，四舍五入到 0.01pt（与 pptxgenjs 一致）
                pPr = p._p.get_or_add_pPr()
                ln = etree.SubElement(pPr, qn('a:lnSpc'))
                etree.SubElement(ln, qn('a:spcPts')).set('val', str(int(math.floor(lh * 100 + 0.5))))
            if bullet:
                pPr = p._p.get_or_add_pPr()
                pPr.set('marL', '171450')
                pPr.set('indent', '-171450')
                etree.SubElement(pPr, qn('a:buSzPct')).set('val', '100000')
                etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
                etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
            for r in runs:
                if not r.get('text'):
                    continue
                run = p.add_run()
                run.text = r['text']
                f = run.font
                f.size = E(size)
                f.bold = r.get('bold', bold)
                f.color.rgb = RGBColor.from_string(r.get('color', color))
                self._font(run, r.get('font', font))
        return tb

    def T(self, x, top, w, size, text, color=BLACK, font=None, bold=True, align='left', lh=None, lines=1,
          runs=None, name=None, boxY=None, boxH=None, valign='top'):
        """单个文本框：top 为文字顶边位置，自动换算成文本框位置"""
        if not self.S:
            return
        l = lh or size * 1.25
        y = boxY if boxY is not None else top - size * 0.12 - (l - size) / 2
        h = boxH if boxH is not None else l * lines
        rs = runs if runs is not None else [dict(text=str(text or ''))]
        self.add_text(x, y, w, h, [rs], size, font or self.BODY, color, bold, align, valign, l,
                      name or strip(text)[:20])

    def RR(self, x, y, w, h, color, r, tr=None, name=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        if not self.S:
            return
        sp = self.S.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = min(0.5, r / min(w, h))
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(color)
        if tr:
            clr = sp.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            etree.SubElement(clr, qn('a:alpha')).set('val', str(int((100 - tr) * 1000)))
        sp.line.fill.background()
        # 去掉主题样式引用（否则 LibreOffice / WPS 会套用主题阴影），与 JS 版生成的形状一致
        style = sp._element.find(qn('p:style'))
        if style is not None:
            sp._element.remove(style)
        if name:
            sp.name = name
        # 去掉形状默认文字样式的影响
        return sp

    def bar(self, x, y):
        self.RR(x, y, 8.5, 2.85, self.theme, 1.4)

    def pic(self, b64_or_bytes, x, y, w, h, name=None, link=None):
        data = base64.b64decode(b64_or_bytes) if isinstance(b64_or_bytes, str) else b64_or_bytes
        p = self.S.shapes.add_picture(io.BytesIO(data), E(x), E(y), E(w), E(h))
        if name:
            p.name = name
        if link:
            p.click_action.hyperlink.address = link
        return p

    # ---- 链接安全：只放行 http/https/mailto/tel ----
    @staticmethod
    def safe_url(u):
        u = str(u or '').strip()
        if not re.match(r'^(https?://|mailto:|tel:)', u, re.I) or re.search(r'[\s<>"]', u):
            return None
        return u   # python-pptx 写 rels 时会自动做 XML 转义

    def contact_link(self, c):
        if c.get('link'):
            return self.safe_url(c['link'])
        t = str(c.get('text') or '').strip()
        typ = c.get('type')
        if typ == 'email' and re.match(r'^[^\s@]+@[^\s@]+\.[^\s@]+$', t):
            return self.safe_url('mailto:' + t)
        if typ in ('web', 'github', 'linkedin', 'link', 'portfolio') and re.match(r'^(https?://)?[\w-]+(\.[\w-]+)+(/\S*)?$', t):
            return self.safe_url(t if re.match(r'^https?:', t) else 'https://' + t)
        return None

    # ---- 头像：裁成正方形，用 PowerPoint 自带的椭圆形状显示（「更改图片」后仍是圆形） ----
    def load_photo(self):
        p = self.D.get('photo')
        if not p:
            return
        path = os.path.join(self.base, p)
        if not os.path.exists(path):
            print('照片不存在：' + path, file=sys.stderr)
            return
        from PIL import Image, ImageOps
        im = ImageOps.exif_transpose(Image.open(path)).convert('RGB')
        s = min(im.size)
        left = (im.width - s) // 2
        top = (im.height - s) // 4      # 偏上取景，尽量保住头顶
        im = im.crop((left, top, left + s, top + s)).resize((600, 600), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=90)
        self.photo = buf.getvalue()

    # ---------- 排版（S 为空时只计算高度） ----------
    def layout(self, cfg):
        D, S = self.D, self.S
        fs, ts, lh, g = cfg['fs'], cfg['ts'], cfg['lh'], cfg['g']
        T, RR, bar = self.T, self.RR, self.bar

        # ===== 头部 =====
        summary = D.get('summary')
        sumLines = self.nLines(summary, 8.1, 521) if summary else 0
        headBottom = 128.9 + (sumLines - 1) * 12 + 22.9 if summary else 112
        RR(20.2, 20.3, 555, headBottom - 20.3, self.theme, 10, name='顶部背景')
        T(HEAD_X, 33.4, 300, 17.1, D.get('name'), color=WHITE, font=LAT if self.EN else YH, name='姓名')
        if D.get('title'):
            label = D.get('titleLabel')
            if label is None:
                label = '' if self.EN else '求职意向：'
            T(HEAD_X, 58.0, 450, 9, label + D['title'], color=WHITE, name='求职意向')
        px = 30
        for i, c in enumerate(D.get('contacts') or []):
            text = str(c.get('text') or '')
            lat = not re.search(r'[^\x00-\x7F]', text)
            icon = ICONS.get(c.get('type'))
            tx = 18.1 if icon else 8
            w = tx + self.textW(text, 8, lat, True) + 8.5
            RR(px, 81.5, w, 17, WHITE, 8.5, tr=70, name='联系方式底框%d' % (i + 1))
            if S and icon:
                self.pic(icon, px + 6.3, 81.5 + (17 - 8.5) / 2, 8.5, 8.5, name='图标%d' % (i + 1))
            T(px + tx, 0, w - tx, 8 if lat else 8.1, text, color=WHITE, font=LAT if lat else CJK,
              boxY=81.5, boxH=17, valign='middle', name='联系方式%d' % (i + 1))
            url = self.contact_link(c)
            if S and url:
                self.pic(CLEAR, px, 81.5, w, 17, name='链接%d' % (i + 1), link=url)
            px += w + 4
        if summary:
            RR(HEAD_X, 112.8, 568.5 - HEAD_X, 1, WHITE, 0, tr=70, name='分隔线', shape=MSO_SHAPE.RECTANGLE)
            T(HEAD_X, 128.9, 526, 8.1, summary, color=WHITE, lh=12, lines=sumLines, bold=self.BB, name='个人简介')
        if S and self.photo:
            p = self.pic(self.photo, 493.2, 35.0, 62.4, 62.4, name='照片')
            try:
                p.auto_shape_type = MSO_SHAPE.OVAL
            except AttributeError:   # 老版本 python-pptx：直接改 XML
                p._element.spPr.find(qn('a:prstGeom')).set('prst', 'ellipse')

        # ===== 左栏 =====
        top0 = headBottom + 18
        y = top0
        auto = D.get('autoBold', True) is not False
        for sec in D.get('left') or []:
            bar(L_X, y + 2.9)
            T(36, y, 200, ts, sec.get('title'), name=sec.get('title'))
            y += 21 * max(g, 0.8)
            for ei, e in enumerate(sec.get('entries') or []):
                # 没有标题、时间和副标题的条目（纯列表/段落）：直接接在板块标题下面，不留空行
                if not e.get('name') and not e.get('date') and not e.get('sub') and not e.get('subRight'):
                    y -= 21 * max(g, 0.8) if ei == 0 else 11.7 * g + 18.3 * max(g, 0.75) - 4
                dw = self.textW(e['date'], fs, self.EN, True) + 4 if e.get('date') else 0
                nameW = DATE_R - L_X - dw - 6
                nameN = self.nLines(e['name'], ts, nameW, self.EN, True) if e.get('name') else 1
                if e.get('name'):
                    T(L_X, y, nameW, ts, e['name'], name=(e['name'] or '条目')[:20], lh=ts * 1.25, lines=nameN)
                if e.get('date'):
                    T(DATE_R - 120, y - 1.5, 120, fs, e['date'], align='right', name='时间')
                if e.get('subRight'):
                    T(DATE_R - 120, y + 9.3, 120, fs, e['subRight'], align='right', name='右侧副信息')
                ny = y + (nameN - 1) * ts * 1.25   # 标题多行时整体下移
                if e.get('sub'):
                    subW = DATE_R - L_X - self.textW(e['subRight'], fs, self.EN, True) - 10 if e.get('subRight') else DATE_R - L_X
                    subN = self.nLines(e['sub'], fs, subW)
                    sy = ny + 12.5 * max(g, 0.85)
                    T(L_X, sy, subW, fs, e['sub'], color=GREY, bold=self.BB, lh=fs * 1.3, lines=subN, name='副标题')
                    cur = sy + (subN - 1) * fs * 1.3 + 16.7 * max(g, 0.75)
                else:
                    cur = ny + 18.3 * max(g, 0.75)
                any_ = False
                # details：成绩/主修课程/技术栈等明细行，冒号前的标签用黑色
                for dl in e.get('details') or []:
                    n = self.nLines(dl, fs, PARA_W)
                    mm = re.match(r'^([^：:]{1,12}[：:]\s?)([\s\S]*)$', str(dl))
                    runs = ([dict(text=mm.group(1), color=BLACK, bold=True)] + self.rich(mm.group(2), GREY, auto)) if mm else self.rich(dl, GREY, auto)
                    T(PARA_X, cur, PARA_W, fs, dl, runs=runs, color=GREY, bold=self.BB, lh=lh, lines=n, name='明细')
                    cur += n * lh
                    any_ = True
                bullets = e.get('bullets') or []
                if bullets:
                    total = sum(self.nLines(b, fs, BUL_W - BUL_IND) for b in bullets)
                    if S:
                        self.add_text(BUL_X, cur - fs * 0.12 - (lh - fs) / 2, BUL_W, lh * total,
                                      [self.rich(b, GREY, auto) for b in bullets], fs, self.BODY, GREY, self.BB,
                                      lh=lh, name='列表_' + str(e.get('name') or '')[:16], bullet=True)
                    cur += total * lh
                    any_ = True
                if e.get('text'):
                    n = self.nLines(e['text'], fs, PARA_W)
                    T(PARA_X, cur, PARA_W, fs, e['text'], runs=self.rich(e['text'], GREY, auto), color=GREY,
                      bold=self.BB, lh=lh, lines=n, name='描述')
                    cur += n * lh
                    any_ = True
                if not any_:
                    cur -= lh * 0.3
                y = cur + 11.7 * g
            y += 1.6 * g
        leftBottom = y - 13.3 * g

        # ===== 右栏卡片 =====
        ry = top0 - 1.2
        for card in D.get('right') or []:
            c0 = ry
            it = c0 + 28.3 * max(g, 0.9)
            lastTop = it
            # tags：标签式卡片，没有小标题和三个点
            items = [dict(name='', tags=card['tags'])] if card.get('tags') else (card.get('items') or [])
            ops = []
            for item in items:
                if item.get('tags'):
                    lines = self.tagLines(item['tags'], 8, R_TXT_W)
                else:
                    lines = [l for raw in str(item.get('content') or '').split('\n') for l in self.wrap(raw, 8, R_TXT_W, True, self.BB)]
                dW = self.textW(item['date'], 7.5, True, True) + 6 if item.get('date') else 0
                nameN = self.nLines(item['name'], ts, 150 - dW, self.EN, True) if item.get('name') else 0
                shift = max(nameN - 1, 0) * ts * 1.25
                t0 = it + 25.3 + shift if item.get('name') else it
                ops.append((it, t0, item, lines, nameN, shift, dW))
                lastTop = t0 + (len(lines) - 1) * 9.6
                it = lastTop + 18.3 * g
            bottom = lastTop + 15.3 if items else c0 + 30
            if S:
                RR(R_X, c0, R_W, bottom - c0, self.card, 6, name=str(card.get('title') or '') + '卡片')
                bar(410.8, c0 + 12.7)
                T(422.9, c0 + 10.6, 140, ts, card.get('title'), name=card.get('title'))
                for it_, t0, item, lines, nameN, shift, dW in ops:
                    if item.get('name'):
                        T(R_IN, it_, 150 - dW, ts, item['name'], name=item['name'], lh=ts * 1.25, lines=nameN)
                        if item.get('date'):
                            T(R_X + R_W - 8 - 70, it_ + 0.8, 70, 7.5, item['date'], color=GREY, align='right', font=LAT, name='时间')
                        for bx in (412.0, 423.0, 434.2):
                            bar(bx, it_ + 15.9 + shift)
                    self.add_text(R_IN, t0 - 8 * 0.12 - 0.8, R_TXT_W + 15, 9.6 * len(lines),
                                  [self.mixRuns(l, GREY) for l in lines], 8, CJK, GREY, self.BB, lh=9.6,
                                  name=str(item.get('name') or '') + '内容')
            ry = bottom + 12
        rightBottom = ry - 12
        return max(leftBottom, rightBottom)

    # ---------- 选档 + 左右栏平衡 + 输出 ----------
    def fit(self):
        for i, c in enumerate(LEVELS):
            if self.layout(c) <= BOTTOM:
                return i, 0.0
        return len(LEVELS) - 1, self.layout(LEVELS[-1]) - BOTTOM

    @staticmethod
    def to_card(sec):
        items = []
        for e in sec.get('entries') or []:
            parts = ['' if e.get('name') else e.get('date'),
                     ' | '.join(x for x in (e.get('sub'), e.get('subRight')) if x)]
            parts += list(e.get('details') or []) + [e.get('text')] + list(e.get('bullets') or [])
            items.append(dict(name=e.get('name') or '', date=e.get('date') if e.get('name') else None,
                              content='\n'.join(p for p in parts if p)))
        return dict(title=sec.get('title'), items=items)

    def build(self, out):
        self.load_photo()
        level, overflow = self.fit()
        # 左右栏平衡：原始版式放不下时，把左栏的短板块（教育、证书、荣誉、语言等）挪成右栏卡片
        moved = []
        if level > 0 and self.D.get('balance', True) is not False:
            size = lambda s: len(json.dumps(s, ensure_ascii=False, separators=(',', ':')))
            left = self.D.get('left') or []
            cands = sorted([s for s in left if s.get('side') == 'right' or (s.get('side') != 'left' and MOVABLE.search(str(s.get('title') or '')))], key=size)
            for sec in cands:
                prev = self.D
                self.D = dict(prev, left=[s for s in prev.get('left') or [] if s is not sec],
                              right=list(prev.get('right') or []) + [self.to_card(sec)])
                lv, ov = self.fit()
                if lv < level or (lv == level and ov < overflow - 1):
                    level, overflow = lv, ov
                    moved.append(sec.get('title'))
                else:
                    self.D = prev
                if level == 0:
                    break
        cfg = LEVELS[level]
        warns = []
        if self.D.get('summary') and self.nLines(self.D['summary'], 8.1, 521) > 3:
            warns.append('个人简介超过 3 行，顶栏会被撑高，建议压缩到 2～3 句')

        prs = Presentation()
        prs.slide_width, prs.slide_height = E(W), E(H)
        self.S = prs.slides.add_slide(prs.slide_layouts[6])
        bottom = self.layout(cfg)
        prs.core_properties.title = str(self.D.get('name') or '') + ' 个人简历'
        prs.save(out)
        self.S = None
        res = dict(out=out, level=level, fontSize=cfg['fs'], lineHeight=cfg['lh'], bottom=jsround(bottom),
                   pageBottom=BOTTOM, overflowPt=jsround(overflow), lang='en' if self.EN else 'zh', movedToRight=moved)
        print(json.dumps(res, ensure_ascii=False))
        for w in warns:
            print('提示：' + w)
        if overflow > 0:
            print('内容超出一页约 %dpt（约 %d 行），需要删减。' % (jsround(overflow), math.ceil(overflow / cfg['lh'])))
        return res


def main(argv):
    if len(argv) < 2:
        print('用法: python resume_pptx.py data.json [out.pptx]')
        return 2
    src = argv[1]
    out = argv[2] if len(argv) > 2 else 'resume.pptx'
    with open(src, encoding='utf-8') as f:
        data = json.load(f)
    Builder(data, os.path.dirname(os.path.abspath(src))).build(out)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
