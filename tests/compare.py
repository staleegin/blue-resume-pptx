"""对比 JS 版和 Python 版生成的 pptx：逐个形状比较位置、大小、文字、加粗、颜色、字体、项目符号。
用法: python compare.py a.pptx b.pptx    退出码 0 = 一致"""
import sys
from pptx import Presentation
from pptx.oxml.ns import qn


def sig(path):
    s = Presentation(path).slides[0]
    out = []
    for sh in s.shapes:
        kind = 'pic' if sh.shape_type == 13 else ('txt' if sh.has_text_frame and sh.text_frame.text.strip() else 'shape')
        item = dict(kind=kind, box=(sh.left, sh.top, sh.width, sh.height))
        if kind == 'pic':
            item['link'] = sh.click_action.hyperlink.address
        if kind == 'txt':
            paras = []
            for p in sh.text_frame.paragraphs:
                pPr = p._p.find(qn('a:pPr'))
                bul = pPr is not None and pPr.find(qn('a:buChar')) is not None
                ln = pPr.find(qn('a:lnSpc')) if pPr is not None else None
                runs = []
                for r in p.runs:
                    rPr = r._r.find(qn('a:rPr'))
                    ea = rPr.find(qn('a:ea')) if rPr is not None else None
                    clr = rPr.find('.//' + qn('a:srgbClr')) if rPr is not None else None
                    runs.append((r.text, rPr.get('b') in ('1', 'true') if rPr is not None else None,
                                 clr.get('val') if clr is not None else None,
                                 ea.get('typeface') if ea is not None else None,
                                 rPr.get('sz') if rPr is not None else None))
                # 合并相邻同格式的 run（两边切分方式可能不同）
                merged = []
                for r in runs:
                    if merged and merged[-1][1:] == r[1:]:
                        merged[-1] = (merged[-1][0] + r[0],) + r[1:]
                    else:
                        merged.append(r)
                paras.append(dict(bul=bul, align=p.alignment if p.alignment not in (None, 1) else 'left', ln=ln.find(qn('a:spcPts')).get('val') if ln is not None else None, runs=merged))
            item['paras'] = paras
            item['anchor'] = sh.text_frame._txBody.find(qn('a:bodyPr')).get('anchor')
        if kind == 'shape':
            sp = sh._element.find('.//' + qn('a:solidFill'))
            clr = sp.find(qn('a:srgbClr')) if sp is not None else None
            alpha = clr.find(qn('a:alpha')) if clr is not None else None
            item['fill'] = (clr.get('val') if clr is not None else None, alpha.get('val') if alpha is not None else None)
        out.append(item)
    return out


def main(a, b):
    A, B = sig(a), sig(b)
    errs = []
    if len(A) != len(B):
        errs.append('形状数量不同: %d vs %d' % (len(A), len(B)))
    for i, (x, y) in enumerate(zip(A, B)):
        for k in set(x) | set(y):
            if k == 'box':
                if any(abs(p - q) > 2 for p, q in zip(x['box'], y['box'])):   # 允许 2 EMU 取整误差
                    errs.append('#%d box %s vs %s' % (i, x['box'], y['box']))
            elif k == 'anchor':
                if (x.get(k) or 't') != (y.get(k) or 't'):
                    errs.append('#%d anchor %s vs %s' % (i, x.get(k), y.get(k)))
            elif x.get(k) != y.get(k):
                errs.append('#%d %s:\n   %s\n   %s' % (i, k, x.get(k), y.get(k)))
    for e in errs[:30]:
        print(e)
    print('形状数 %d / %d，差异 %d 处' % (len(A), len(B), len(errs)))
    return 1 if errs else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1], sys.argv[2]))
